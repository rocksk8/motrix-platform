"""網路架構規劃書 Excel／PDF 匯出建構邏輯，供 routers/network_plans.py 呼叫。

章節/欄位結構完全比照 NETWORK-PLAN-MODULE-DESIGN.md §3.2 調閱的實際業務範本
（小林機械案 9 個分頁）＋前端 network-plan-form.html 的 DATA_TABS 定義，
兩邊 key/欄位順序需保持一致——改其中一邊記得同步改另一邊。
"""
import io
import os
import subprocess
import tempfile
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from helpers import _get_edge_path
from network_plan_topology import build_topology_svg, build_topology_text_summary_html

_COMPANY  = "允碩整合集創股份有限公司"
_COMPANY2 = "MOTRIX Synergy Integration Corp."

# (data_json key, 章節/分頁標題, 欄位 [(key, 中文表頭), ...]) —— 與
# network-plan-form.html::DATA_TABS 一一對應。
#
# ⚠️ 分頁標題（第二個欄位）之後新增/修改時絕對不能包含全形斜線「／」（U+FF0F）：
# 2026-08-26 實測發現，含「／」的分頁名稱（原本的「WAN／對外線路」「IP／Port
# 群組」）會讓 Microsoft Excel 判定 workbook.xml 損毀、跳出修復對話框、把該
# 分頁之後的內容整個吞掉重組成「復原_工作表1」——即使 openpyxl／python
# zipfile／XML well-formedness 檢查全部通過也一樣（這是 Excel 自己額外的
# 分頁名稱驗證規則，並非標準 OOXML schema 或一般 XML 工具會擋下的問題）。
# 已改用半形空格／連字號替代（"WAN 對外線路"／"IP-Port 群組"）。欄位表頭
# （每個 tuple 的第二個字串，如 "角色／用途"）是儲存格文字內容，不受此限制，
# 可以繼續用全形斜線。
SECTIONS = [
    ("wanLines", "WAN 對外線路", [
        ("isp", "ISP"), ("lineType", "線路類型"), ("bandwidthUp", "上行頻寬"),
        ("bandwidthDown", "下行頻寬"), ("publicIp", "固定 IP"), ("subnet", "遮罩"),
        ("gateway", "Gateway"), ("dns1", "DNS1"), ("dns2", "DNS2"),
        ("contractNo", "合約編號"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("devices", "設備清單", [
        ("seq", "項次"), ("name", "設備名稱"), ("category", "類別"), ("model", "型號"),
        ("location", "位置"), ("mgmtIp", "管理 IP"), ("mgmtVlan", "管理 VLAN"),
        ("role", "角色／用途"), ("mac", "MAC"), ("phase", "階段"),
        ("status", "狀態"), ("note", "備註"),
    ]),
    ("vlans", "VLAN 規劃", [
        ("vlanId", "VLAN ID"), ("name", "名稱"), ("nameZh", "中文名稱"),
        ("purposeTag", "Purpose"), ("cidr", "網段 CIDR"), ("gateway", "Gateway"),
        ("dhcpMode", "DHCP 模式"), ("dhcpStart", "DHCP 起始"), ("dhcpEnd", "DHCP 結束"),
        ("dns", "DNS"), ("usage", "用途"), ("location", "使用單位／樓層"),
        ("phase", "階段"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("ipAllocations", "IP 位址配置", [
        ("seq", "項次"), ("device", "主機／設備"), ("ip", "IP"), ("tag", "標籤名稱"),
        ("openPorts", "開放埠"), ("vlan", "VLAN"), ("mask", "遮罩"), ("gateway", "Gateway"),
        ("assignType", "配發方式"), ("mac", "MAC"), ("osFw", "OS／韌體"),
        ("usage", "用途"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("portProfiles", "Port Profile 定義", [
        ("name", "Profile 名稱"), ("nativeVlan", "Native VLAN"), ("nativeVlanName", "Native VLAN 名稱"),
        ("taggedVlans", "Tagged VLAN"), ("poeMode", "PoE 模式"), ("usage", "用途說明"),
        ("applyTo", "套用對象"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("switchPorts", "交換器 Port 對應", [
        ("device", "設備"), ("location", "位置"), ("mgmtIp", "管理 IP"), ("portNo", "埠號"),
        ("portType", "埠類型"), ("endpoint", "連接對象／端點"), ("portProfile", "Port Profile"),
        ("nativeVlanAuto", "Native VLAN"), ("taggedVlanAuto", "Tagged VLAN"), ("poe", "PoE"),
        ("cableLabel", "配線標籤"), ("endpointLocation", "端點位置"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("firewallRules", "防火牆規則", [
        ("priority", "優先權"), ("name", "規則名稱"), ("ruleType", "類型"), ("action", "動作"),
        ("protocol", "協定"), ("source", "來源"), ("destination", "目的"),
        ("destPort", "目的 Port"), ("enabled", "啟用"), ("note", "備註"),
    ]),
    ("ipPortGroups", "IP-Port 群組", [
        ("category", "類別"), ("groupName", "群組名稱"), ("members", "成員（IP／網段／Port）"),
        ("usage", "用途"), ("referencedBy", "被哪些規則引用"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("wifiSsids", "無線 SSID", [
        ("seq", "編號"), ("ssid", "SSID"), ("profile", "Profile 名稱"), ("vlan", "VLAN"),
        ("security", "加密方式"), ("auth", "密碼／認證"), ("band", "頻段"), ("apGroup", "AP 群組"),
        ("downLimit", "下載限速"), ("upLimit", "上傳限速"), ("clientIsolation", "用戶端隔離"),
        ("roaming11r", "漫遊(11r)"), ("status", "狀態"), ("note", "備註"),
    ]),
    ("cabling", "線路與幹線", [
        ("seq", "編號"), ("segment", "路段"), ("from", "起點設備／埠"), ("to", "終點設備／埠"),
        ("media", "介質"), ("spec", "規格"), ("lengthEst", "概估長度"), ("label", "標籤編號"),
        ("routePath", "路由路徑"), ("status", "狀態"), ("note", "備註"),
    ]),
]


def _cell_text(v):
    """Excel 值/HTML 顯示共用的欄位轉換。空值一律回傳 None（不是空字串）——
    openpyxl 寫入空字串會產生 `<c t="inlineStr"></c>`（缺少必要的 `<is>` 子
    元素），是不合法的 OOXML，Excel 開啟時會跳出「發現部分內容有問題」的
    修復對話框（2026-08-26 實測發現，範例檔案幾乎每個空白欄位都中招）。
    None 才會讓 openpyxl 產生真正空白、合法的儲存格。HTML 那邊 `_esc(None)`
    本來就會轉成空字串，行為不變。"""
    if isinstance(v, bool):
        return "是" if v else "否"
    if v is None or v == "":
        return None
    return v


# 目前 10 大類明細裡唯一的 checkbox 型欄位（前端 network-plan-form.html 的
# DATA_TABS 對應同一個欄位標記 type:'checkbox'）——匯入時「是」/TRUE 要轉回
# Python bool，不能原樣存成字串，否則前端 checkbox 綁定會失效。
_BOOL_FIELDS = {("firewallRules", "enabled")}


def _parse_cell_value(section_key, col_key, raw):
    if raw is None or raw == "":
        return None
    if (section_key, col_key) in _BOOL_FIELDS:
        return raw in (True, "是", "TRUE", "True", "true", 1)
    if isinstance(raw, float) and raw.is_integer():
        return int(raw)
    if isinstance(raw, str):
        return raw.strip()
    return raw


def parse_plan_excel(content: bytes) -> dict:
    """讀取使用者匯入的 Excel（預期是本系統匯出的範本，可能已離線填寫過），
    回傳 {"sections": {sectionKey: [rows...]}, "warnings": [...]}。

    比對規則：分頁名稱＝`SECTIONS` 定義的標題（跟匯出時完全一致才會被辨識）；
    欄位用**表頭文字**比對（不是欄位順序），使用者調整過欄位順序、或刪掉不
    需要的欄位都還能正確對應，只有表頭文字被改掉的欄位會對不到、視為略過。
    完全空白的列（所有欄位皆空）不納入。只會拋出例外給呼叫端轉成 400 的
    情況是檔案本身不是合法 xlsx（openpyxl 無法開啟）。

    未辨識的分頁名稱一律回報在 warnings（不當成錯誤，因為使用者可能在
    Excel 裡自己加了輔助分頁），**這是 2026-08-26 特意補上的行為**：舊版
    分頁名稱（改分頁命名前，如帶全形斜線的「WAN／對外線路」）匯入到新版
    程式碼會完全比對不到、被略過，若靜默跳過使用者會完全看不出「為什麼
    匯入完成訊息裡少了這兩個分頁」，只能回頭翻程式碼或猜——曾實際發生過
    使用者拿舊範本測試匯入、成功訊息卻少兩個分頁又沒有任何提示的情況。"""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    title_to_section = {title[:31]: (key, columns) for key, title, columns in SECTIONS}

    sections = {}
    warnings = []
    for sheet_name in wb.sheetnames:
        if sheet_name not in title_to_section:
            warnings.append(f"分頁「{sheet_name}」無法辨識，已略過（若這是舊版範本的分頁，"
                             f"請改用最新匯出的範本重新填寫）")
            continue
        section_key, columns = title_to_section[sheet_name]
        ws = wb[sheet_name]

        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        col_idx_by_key = {}
        for idx, header in enumerate(header_cells):
            for k, label in columns:
                if header == label and k not in col_idx_by_key:
                    col_idx_by_key[k] = idx
                    break
        if not col_idx_by_key:
            warnings.append(f"「{sheet_name}」找不到任何符合的欄位標題，已略過此分頁")
            continue

        rows_out = []
        for row_cells in ws.iter_rows(min_row=2, values_only=True):
            record = {}
            for k, idx in col_idx_by_key.items():
                if idx >= len(row_cells):
                    continue
                v = _parse_cell_value(section_key, k, row_cells[idx])
                if v is not None:
                    record[k] = v
            if record:
                rows_out.append(record)
        sections[section_key] = rows_out

    if not sections:
        warnings.append("找不到任何符合本系統格式的分頁，請確認上傳的是本系統匯出的 Excel 範本")
    return {"sections": sections, "warnings": warnings}


# ── Excel ─────────────────────────────────────────────────────────────────

def build_plan_excel(plan: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    data = plan.get("data") or {}

    header_font = Font(bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="111827")

    ws = wb.create_sheet("封面")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 44
    ws["A1"] = _COMPANY
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = _COMPANY2
    ws["A2"].font = Font(size=9, color="6B7280")
    ws["A3"] = "網路架構規劃書"
    ws["A3"].font = Font(bold=True, size=12)
    cover_rows = [
        ("規劃書編號", plan.get("planNo", "")),
        ("案場名稱", plan.get("siteName", "")),
        ("綁定案件", plan.get("quoteNo") or "（獨立建立）"),
        ("聯絡人", plan.get("contactName", "")),
        ("聯絡電話", plan.get("contactPhone", "")),
        ("狀態", plan.get("status", "")),
        ("最後更新", (plan.get("updatedAt") or "")[:16].replace("T", " ")),
        ("產製時間", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    r = 5
    for label, val in cover_rows:
        ws.cell(row=r, column=1, value=label).font = Font(color="6B7280", size=10)
        ws.cell(row=r, column=2, value=val or None)
        r += 1

    for key, title, columns in SECTIONS:
        items = data.get(key) or []
        sheet = wb.create_sheet(title[:31])
        for i, (_, header) in enumerate(columns, start=1):
            c = sheet.cell(row=1, column=i, value=header)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(vertical="center")
        for r_i, item in enumerate(items, start=2):
            for c_i, (k, _) in enumerate(columns, start=1):
                sheet.cell(row=r_i, column=c_i, value=_cell_text(item.get(k)))
        for i in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 16
        sheet.freeze_panes = "A2"

    ws = wb.create_sheet("修訂紀錄")
    for i, header in enumerate(["版本", "日期", "修改人", "說明"], start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = header_font
        c.fill = header_fill
    for r_i, log in enumerate(data.get("revisionLog") or [], start=2):
        ws.cell(row=r_i, column=1, value=log.get("version"))
        ws.cell(row=r_i, column=2, value=(log.get("date") or "")[:16].replace("T", " "))
        ws.cell(row=r_i, column=3, value=log.get("editor"))
        ws.cell(row=r_i, column=4, value=log.get("note"))
    for col, w in zip("ABCD", [8, 18, 14, 50]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────

def _esc(s):
    return (str(s) if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _section_table_html(title, columns, items):
    if not items:
        return ""
    thead = "".join(f"<th>{_esc(h)}</th>" for _, h in columns)
    rows = ""
    for item in items:
        cells = "".join(f"<td>{_esc(_cell_text(item.get(k)))}</td>" for k, _ in columns)
        rows += f"<tr>{cells}</tr>"
    return (
        f'<div class="section-label">{_esc(title)}</div>'
        f"<table><thead><tr>{thead}</tr></thead><tbody>{rows}</tbody></table>"
    )


def build_plan_html(plan: dict) -> str:
    data = plan.get("data") or {}
    try:
        topo_svg = (build_topology_svg(data) or {}).get("html")
    except Exception:
        topo_svg = None
    topo_html = (
        f'<div class="section-label">網路拓樸圖</div>'
        f'<div style="overflow-x:auto;page-break-inside:avoid;break-inside:avoid">{topo_svg}</div>'
    ) if topo_svg else ""
    sections_html = "".join(
        _section_table_html(title, columns, data.get(key) or [])
        for key, title, columns in SECTIONS
    )
    rev_rows = "".join(
        f'<tr><td>v{_esc(log.get("version"))}</td>'
        f'<td>{_esc((log.get("date") or "")[:16].replace("T", " "))}</td>'
        f'<td>{_esc(log.get("editor"))}</td><td>{_esc(log.get("note"))}</td></tr>'
        for log in (data.get("revisionLog") or [])
    )
    rev_html = ""
    if rev_rows:
        rev_html = (
            '<div class="section-label">修訂紀錄</div>'
            "<table><thead><tr><th>版本</th><th>日期</th><th>修改人</th><th>說明</th></tr></thead>"
            f"<tbody>{rev_rows}</tbody></table>"
        )

    return (
        '<!DOCTYPE html>\n<html lang="zh-Hant">\n<head>\n<meta charset="UTF-8">\n'
        f'<title>{_esc(plan.get("planNo", ""))} 網路架構規劃書</title>\n'
        "<style>\n"
        "  *{box-sizing:border-box;margin:0;padding:0}\n"
        '  body{font-family:"Microsoft JhengHei","PMingLiU",serif;font-size:10.5px;color:#0A0A0A;line-height:1.5;background:#fff}\n'
        "  #root{padding:16px 20px}\n"
        '  @page{size:A4 landscape;margin:8mm;@bottom-center{content:counter(page);font-family:Arial,sans-serif;font-size:9px;color:#888}}\n'
        "  @media print{html,body{margin:0;padding:0;background:#fff}tr{page-break-inside:avoid}}\n"
        "  .accent-bar{height:3px;background:#0A0A0A;margin-bottom:12px}\n"
        "  .header{display:flex;justify-content:space-between;align-items:flex-start;padding-bottom:10px;border-bottom:1px solid #0A0A0A;margin-bottom:12px}\n"
        "  .co-name{font-size:14px;font-weight:700;letter-spacing:.06em}\n"
        "  .co-sub{font-size:9px;color:#888;margin-top:2px;font-family:Arial,sans-serif}\n"
        "  .doc-title{font-size:18px;font-weight:700;letter-spacing:.16em;text-align:right}\n"
        "  .meta{display:grid;grid-template-columns:repeat(4,1fr);gap:4px;margin-bottom:12px;font-size:10.5px;background:#FAFAF8;padding:8px 10px;border-radius:4px;border:1px solid #EDEAE4}\n"
        "  .meta span{color:#888;font-family:Arial,sans-serif;font-size:9.5px}\n"
        '  .section-label{font-size:9px;font-family:Arial,sans-serif;letter-spacing:.1em;text-transform:uppercase;color:#888;font-weight:600;margin:12px 0 6px;display:flex;align-items:center;gap:8px}\n'
        '  .section-label::after{content:"";flex:1;height:1px;background:#EDEAE4}\n'
        "  table{width:100%;border-collapse:collapse;margin-bottom:10px;table-layout:fixed}\n"
        "  thead th{background:#0A0A0A;color:#F5F4F0;padding:5px 6px;text-align:left;font-size:9px;font-weight:500;font-family:Arial,sans-serif;word-break:break-all}\n"
        "  tbody td{padding:5px 6px;border-bottom:1px solid #EDEAE4;font-size:9.5px;word-break:break-all}\n"
        "  tbody tr:nth-child(even) td{background:#FAFAF8}\n"
        "  .footer{text-align:center;font-size:9px;color:#888;margin-top:14px;padding-top:10px;border-top:1px solid #EDEAE4;font-family:Arial,sans-serif}\n"
        "</style>\n</head>\n<body>\n<div id=\"root\">\n"
        '<div class="accent-bar"></div>\n'
        f'<div class="header">\n  <div>\n    <div class="co-name">{_esc(_COMPANY)}</div>\n'
        f'    <div class="co-sub">{_esc(_COMPANY2)}</div>\n'
        '    <div class="co-sub" style="margin-top:3px">統一編號：60575481　｜　電話：04-3610-6566　｜　info@miactw.com</div>\n'
        '  </div>\n  <div>\n    <div class="doc-title">網路架構規劃書</div>\n  </div>\n</div>\n'
        '<div class="meta">\n'
        f'  <div><span>規劃書編號：</span>{_esc(plan.get("planNo", ""))}</div>\n'
        f'  <div><span>案場名稱：</span>{_esc(plan.get("siteName", ""))}</div>\n'
        f'  <div><span>綁定案件：</span>{_esc(plan.get("quoteNo") or "獨立建立")}</div>\n'
        f'  <div><span>狀態：</span>{_esc(plan.get("status", ""))}</div>\n'
        "</div>\n"
        f"{topo_html}"
        f"{sections_html}"
        f"{rev_html}"
        f'<div class="footer">{_esc(_COMPANY2)} 允碩整合集創 ｜ info@miactw.com ｜ Tel: 04-3610-6566 ｜ 統一編號: 60575481　｜　產製時間：{datetime.now().strftime("%Y-%m-%d %H:%M")}</div>\n'
        "</div>\n</body>\n</html>"
    )


def _render_pdf_via_edge(html_content: str, virtual_time_budget: int = None) -> bytes:
    """共用的 HTML→PDF 轉檔（Edge headless），供完整規劃書與純拓樸圖快速工具
    共用，避免兩處各自維護一份幾乎一樣的 subprocess 邏輯。

    virtual_time_budget：比照使用者原本個案腳本 b1f_topology.py 的轉檔方式
    （見 build_topology_only_pdf_bytes）——當頁面內有轉檔前必須先跑完的非同步
    JS（如動態量測內容尺寸再設定 @page 大小），單純等 window.onload 不夠，
    需要這個旗標讓 headless 把頁面內部的計時器/Promise 佇列往前推進足夠時間
    再截圖轉檔。一般靜態頁面（完整規劃書 PDF）不需要，維持 None 走原本的
    --run-all-compositor-stages-before-draw。"""
    edge = _get_edge_path()
    tmp_html = tmp_pdf = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", suffix=".html", encoding="utf-8", delete=False) as f:
            f.write(html_content)
            tmp_html = f.name
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            tmp_pdf = f.name
        file_url = "file:///" + tmp_html.replace("\\", "/")
        wait_flag = (f"--virtual-time-budget={virtual_time_budget}" if virtual_time_budget
                     else "--run-all-compositor-stages-before-draw")
        subprocess.run(
            [edge, "--headless", "--disable-gpu", "--no-sandbox",
             f"--print-to-pdf={tmp_pdf}",
             "--no-pdf-header-footer",
             wait_flag,
             file_url],
            timeout=40, check=False,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        if not os.path.exists(tmp_pdf) or os.path.getsize(tmp_pdf) == 0:
            raise ValueError("Edge 執行完畢但未產生 PDF 檔案")
        with open(tmp_pdf, "rb") as f:
            return f.read()
    finally:
        for p in (tmp_html, tmp_pdf):
            if p:
                try:
                    os.unlink(p)
                except Exception:
                    pass


def build_plan_pdf_bytes(plan: dict) -> bytes:
    return _render_pdf_via_edge(build_plan_html(plan))


# ── 快速拓樸圖（不建立規劃書，純畫圖用，見 routers/network_plans_quick.py） ──

def build_topology_only_html(data: dict, title: str = "", floor_tag: str = "", footer: str = "") -> str:
    """快速拓樸圖工具專用：標題＋拓樸圖＋文字版埠位對照表＋頁尾的極簡頁面，
    不含規劃書的 WAN／VLAN／IP 等其餘章節——呼叫端（routers/network_plans_quick.py）
    已經先確認過 build_topology_svg 有東西可畫才會呼叫這裡。

    2026-09-06 使用者要求「格式跟分頁方式要完全一樣」，改回逐項比照使用者
    原始個案腳本 b1f_topology.py 實際輸出（B1F_topology.html）：
    - 版面（CSS 變數、淺色放射漸層背景、.board 白卡＋陰影圓角、h1+徽章式
      標題列、footer 靠右對齊樣式）逐一比照原腳本 CSS，不再沿用規劃書 PDF
      那套企業合約書風格（accent-bar／co-name／doc-title）。
    - 型號徽章（.tag.mod）比照原腳本「型號　(N×GbE + M×SFP)」樣式，從
      build_topology_svg() 回傳的 models／uniform_ports 動態組出（原腳本是
      寫死文字，這裡改成依實際資料算，多型號或埠數不一致時只顯示型號、
      不強加可能失真的埠數字樣）。
    - 分頁方式改回原腳本手法：不用固定 A4，改用頁尾 <script> 等頁面
      load+字型 ready 後，量測實際渲染尺寸再動態產生剛好等於內容大小的
      @page（永遠一頁、不強制切成 A4 多頁；本次 2026-09-06 對話已跟使用者
      確認過取捨——這會使匯出的 PDF 頁面尺寸不是標準 A4，直接送實體印表機
      可能被印表機驅動縮放或裁切，不影響數位保存/瀏覽器開啟）——取代
      2026-09-04 當時因應印表分頁切斷問題而暫時採用的 A4 直版＋
      page-break-inside:avoid 版本。
    - 「埠位對照表 Port Assignment」＋固定兩欄 .tbl-wrap CSS Grid維持不變
      （2026-09-04 已比照原腳本做過，本次沿用）；埠號欄位改用等寬字體＋
      粗體（比照原腳本 td.port），純 CSS 選取器達成，不需更動
      build_topology_text_summary_html() 的表格 HTML 結構。"""
    try:
        topo = build_topology_svg(data) or {}
    except Exception:
        topo = {}
    topo_svg = topo.get("html") or "<div style='color:#888'>（尚無可畫的交換器資料）</div>"
    try:
        summary_tables = build_topology_text_summary_html(data)
    except Exception:
        summary_tables = ""
    summary_block = (
        f'<h2>埠位對照表 Port Assignment</h2>\n<div class="tbl-wrap">{summary_tables}</div>\n'
    ) if summary_tables else ""
    title = (title or "").strip() or "網路埠拓樸圖"

    models = topo.get("models") or []
    mod_tag = " / ".join(models)
    uniform_ports = topo.get("uniform_ports")
    if uniform_ports:
        copper, sfp = uniform_ports
        spec_bits = [b for b in (f"{copper}×GbE" if copper else "", f"{sfp}×SFP" if sfp else "") if b]
        if spec_bits:
            mod_tag = (mod_tag + "　" if mod_tag else "") + "(" + " + ".join(spec_bits) + ")"
    mod_tag_html = f'<span class="tag mod">{_esc(mod_tag)}</span>' if mod_tag else ""
    floor_tag_html = f'<span class="tag">{_esc(floor_tag)}</span>' if (floor_tag or "").strip() else ""
    footer_text = (footer or "").strip() or f"{_COMPANY2} 允碩整合集創 ｜ 產製時間：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    return (
        '<!DOCTYPE html>\n<html lang="zh-Hant">\n<head>\n<meta charset="UTF-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1.0">\n'
        f'<title>{_esc(title)}</title>\n'
        "<style>\n"
        "  :root{--ink:#0f172a;--sub:#475569;--line:#cbd5e1;--paper:#fff;}\n"
        "  *{box-sizing:border-box}\n"
        "  body{margin:0;padding:30px 26px 54px;color:var(--ink);"
        "font-family:'Microsoft JhengHei','Noto Sans TC',system-ui,sans-serif;\n"
        "   background:radial-gradient(1200px 600px at 12% -10%,#e7eef7 0%,transparent 60%),\n"
        "   radial-gradient(1000px 500px at 100% 0%,#eef0f8 0%,transparent 55%),#eef2f6;\n"
        "   -webkit-print-color-adjust:exact;print-color-adjust:exact}\n"
        "  .wrap{margin:0 auto}\n"
        "  header{display:flex;align-items:baseline;gap:14px;flex-wrap:wrap}\n"
        "  h1{font-weight:900;font-size:28px;letter-spacing:.5px;margin:0}\n"
        "  .tag{font-family:Consolas,'Courier New',monospace;font-weight:700;font-size:13px;"
        "background:var(--ink);color:#fff;padding:4px 11px;border-radius:6px}\n"
        "  .tag.mod{background:#334155}\n"
        "  .board{background:var(--paper);border:1px solid var(--line);border-radius:16px;"
        "padding:22px 20px 26px;box-shadow:0 10px 30px -18px rgba(15,23,42,.35);display:inline-block}\n"
        "  .board svg{display:block;max-width:100%;height:auto}\n"
        "  h2{font-weight:900;font-size:18px;margin:32px 0 10px}\n"
        "  .tbl-wrap{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:16px}\n"
        "  table{width:100%;border-collapse:collapse;font-size:13px;background:var(--paper);"
        "border:1px solid var(--line);border-radius:12px;overflow:hidden;table-layout:fixed}\n"
        "  caption{caption-side:top;text-align:left;font-weight:700;font-size:14px;padding:0 0 7px 2px}\n"
        "  th,td{padding:6px 10px;text-align:left;border-bottom:1px solid #e2e8f0;word-break:break-all}\n"
        "  th{background:#f1f5f9;font-weight:700;font-size:12px}\n"
        "  td:first-child{font-family:Consolas,'Courier New',monospace;font-weight:700}\n"
        "  tr:last-child td{border-bottom:none}\n"
        "  footer{display:block;margin-top:28px;color:var(--sub);font-size:12px;text-align:right}\n"
        "</style>\n</head>\n<body>\n<div class=\"wrap\">\n"
        f'<header>\n  <h1>{_esc(title)}</h1>\n  {floor_tag_html}\n  {mod_tag_html}\n</header>\n'
        f'<div class="board">{topo_svg}</div>\n'
        f"{summary_block}"
        f'<footer>{_esc(footer_text)}</footer>\n'
        "</div>\n"
        "<script>\n"
        "function fitPageToContent() {\n"
        '  var board = document.querySelector(".board");\n'
        '  var wrap = document.querySelector(".wrap");\n'
        "  if (board && wrap) {\n"
        '    wrap.style.maxWidth = Math.ceil(board.getBoundingClientRect().width) + "px";\n'
        "  }\n"
        "  var w = Math.ceil(document.documentElement.scrollWidth);\n"
        "  var h = Math.ceil(document.documentElement.scrollHeight) + 16;\n"
        '  var style = document.createElement("style");\n'
        '  style.textContent = "@page { size: " + w + "px " + h + "px; margin: 0; } body { margin: 0; }";\n'
        "  document.head.appendChild(style);\n"
        "}\n"
        "var loaded = new Promise(function (res) {\n"
        '  if (document.readyState === "complete") res();\n'
        '  else window.addEventListener("load", res);\n'
        "});\n"
        "var fontsReady = (document.fonts && document.fonts.ready) ? document.fonts.ready : Promise.resolve();\n"
        "Promise.all([loaded, fontsReady]).then(fitPageToContent);\n"
        "</script>\n"
        "</body>\n</html>"
    )


def build_topology_only_pdf_bytes(data: dict, title: str = "", floor_tag: str = "", footer: str = "") -> bytes:
    html = build_topology_only_html(data, title, floor_tag, footer)
    return _render_pdf_via_edge(html, virtual_time_budget=8000)

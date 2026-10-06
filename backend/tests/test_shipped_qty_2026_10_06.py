# -*- coding: utf-8 -*-
"""第 43 班：報價品項「已出貨數量」（使用者裁示選項 B：材料申請出貨＋從報價單帶入的出貨列都要能歸屬）。

守門：
 ① 提供者 `shipping.quote_item_shipped`（M03）：只算 有 quoteItemId、非標題、沒有 materialLink、沒有料號／序號、數量為正 的列；已核准＝shipped、待審核／簽核中＝reserved、草稿／已退回不計
 ② 帶 materialLink 的列即使也帶 quoteItemId 只算一次（走材料申請那一邊）——突變：拿掉 materialLink 排除 ⇒ 重複計算 ⇒ 紅
 ③ `case.item_shipped.shipped_by_item`：材料申請連到報價品項＋報價帶入列加總；沒有可歸屬紀錄 ⇒ attributed=False（畫面「—」，不是 0）
 ④ 端點 `GET /api/quotations/{no}/item-shipped`：只有數量（回應沒有任何金額／成本欄位）；沒登入 401；看不到的案件 404
 ⑤ 預設不變：沒有任何出貨時全部 shipped=0、reserved=0、attributed=False；既有的材料申請出貨提供者回傳不變
 ⑥ 營運報表匯出：「毛利分析」多兩欄（附在最右，既有欄號不動）；提供者不在 ⇒ 空白、不影響匯出
斷言打在提供者回傳與 API 回應，不打畫面文字。"""
import json

import pytest

from tests._requires import requires_module
from modules.supply.tests.test_shipping_material_link_2026_10_03 import DOC, ITEM, Q, _note, _set, world, L  # noqa: F401

pytestmark = requires_module("case", "已出貨數量需要 M01 的材料申請")


@pytest.fixture
def sq(world):
    """在連結測試的案件上補報價品項：q1 電纜 10 捲（材料申請 ITEM 連到它）、q2 攝影機 8 台、q3 配件 5 個、q4 另一品項 3 組。"""
    client, h = world
    import db
    c = db.get_db()
    try:
        row = c.execute("SELECT data_json FROM quotations WHERE quote_no=?", (Q,)).fetchone()
        data = json.loads(row["data_json"])
        data["items"] = [{"id": "q1", "description": "電纜", "qty": 10, "unit": "捲", "cost": 100, "unitPrice": 150},
                         {"type": "header", "description": "標題"},
                         {"id": "q2", "description": "攝影機", "qty": 8, "unit": "台", "cost": 3000, "unitPrice": 4500},
                         {"id": "q3", "description": "配件", "qty": 5, "unit": "個"},
                         {"id": "q4", "description": "另一品項", "qty": 3, "unit": "組"}]
        c.execute("UPDATE quotations SET data_json=? WHERE quote_no=?", (json.dumps(data, ensure_ascii=False), Q))
        c.commit()
    finally:
        c.close()
    return client, h


def Qi(qty, qid="q2", **extra):
    d = {"description": "攝影機", "qty": qty, "unit": "台", "quoteItemId": qid}
    d.update(extra)
    return d


def _prov():
    from core import registry
    import db
    return db.get_db(), registry.providers("shipping.quote_item_shipped")["supply"]


def _by_item():
    import db
    from modules.case import item_shipped as IS
    c = db.get_db()
    try:
        data = json.loads(c.execute("SELECT data_json FROM quotations WHERE quote_no=?", (Q,)).fetchone()["data_json"])
        return IS.shipped_by_item(c, Q, data)
    finally:
        c.close()


# ── ① 提供者 ─────────────────────────────────────────────────────────

def test_quote_item_provider_counts_reserved_and_shipped_but_not_draft_or_returned(sq):
    client, h = sq
    for st, q in (("草稿", 1), ("待審核", 2), ("簽核中", 3), ("已核准", 4), ("已退回", 8)):
        _note(client, h, [Qi(q)], status=st)
    c, fn = _prov()
    try:
        out = fn(c, Q)
    finally:
        c.close()
    assert out["q2"]["reserved"] == 5.0 and out["q2"]["shipped"] == 4.0 and len(out["q2"]["notes"]) == 3


def test_quote_item_provider_skips_headers_stock_lines_material_links_and_bad_qty(sq):
    client, h = sq
    lines = [Qi(2),                                                                      # ✔
             {"type": "header", "description": "標題", "quoteItemId": "q2", "qty": 9},       # 標題列
             Qi(3, part_no="PN-1"),                                                      # 庫存料號列
             Qi(3, serials=["S1", "S2", "S3"]),                                          # 庫存序號列
             Qi(0), Qi(-1), Qi("x"),                                                     # 數量不是正數
             {"description": "舊列", "qty": 7},                                           # 沒有 quoteItemId（舊單／手動列）
             Qi(6, qid="  ")]                                                            # 空白 id
    _note(client, h, lines, status="已核准")
    c, fn = _prov()
    try:
        out = fn(c, Q)
    finally:
        c.close()
    assert out == {"q2": {"reserved": 0.0, "shipped": 2.0, "notes": out["q2"]["notes"]}}, out


def test_a_line_with_material_link_and_quote_item_id_is_counted_once_via_the_material_side(sq):
    """② 突變對照：拿掉 `materialLink is not None` 排除 ⇒ 這條會同時出現在 quote_item_shipped ⇒ 這題紅。"""
    client, h = sq
    _note(client, h, [{**L(4), "quoteItemId": "q1", "qty": 4}], status="已核准")
    c, fn = _prov()
    try:
        assert fn(c, Q) == {}, "帶 materialLink 的列不可在報價品項提供者再算一次"
    finally:
        c.close()
    assert _by_item()["q1"]["shipped"] == 4.0, "材料申請那一邊要算到（ITEM 連到 q1）"


# ── ③ 加總與歸屬 ─────────────────────────────────────────────────────

def test_shipped_by_item_adds_material_and_quote_copied_lines(sq):
    client, h = sq
    _note(client, h, [L(4)], status="已核准")                                            # 材料申請 ⇒ q1 已出貨 4
    _note(client, h, [L(2)], status="待審核")                                            # q1 占用 2
    _note(client, h, [Qi(3, qid="q1")], status="已核准")                                 # 報價帶入 q1 ⇒ 再 +3
    _note(client, h, [Qi(5)], status="已核准")                                           # q2 已出貨 5
    _note(client, h, [Qi(2)], status="簽核中")                                           # q2 占用 2
    r = _by_item()
    assert (r["q1"]["ordered"], r["q1"]["unit"], r["q1"]["shipped"], r["q1"]["reserved"], r["q1"]["attributed"]) == (10.0, "捲", 7.0, 2.0, True)
    assert (r["q2"]["ordered"], r["q2"]["shipped"], r["q2"]["reserved"]) == (8.0, 5.0, 2.0)
    assert r["q3"]["attributed"] is False and r["q3"]["shipped"] == 0.0 and r["q3"]["ordered"] == 5.0, "沒有可歸屬紀錄 ⇒ 顯示「—」不是 0"
    assert list(r) == ["q1", "q2", "q3", "q4"], "標題列不算品項、保持報價順序"


def test_old_notes_without_quote_item_id_stay_unattributed(sq):
    client, h = sq
    _note(client, h, [{"description": "攝影機", "qty": 8, "unit": "台"}], status="已核准")           # 舊單：沒有 quoteItemId、沒有連結
    r = _by_item()
    assert all(v["attributed"] is False and v["shipped"] == 0.0 for v in r.values())


def test_unknown_quote_item_id_is_ignored(sq):
    client, h = sq
    _note(client, h, [Qi(3, qid="not-an-item")], status="已核准")
    assert all(v["attributed"] is False for v in _by_item().values())


def test_provider_failure_only_drops_that_source(sq, monkeypatch):
    client, h = sq
    _note(client, h, [L(4)], status="已核准")
    _note(client, h, [Qi(5)], status="已核准")
    from core import registry
    real = registry.providers

    def boom_providers(name):
        d = real(name)
        if name == "shipping.quote_item_shipped":
            return {"supply": lambda *a, **k: (_ for _ in ()).throw(RuntimeError("x"))}
        return d
    monkeypatch.setattr(registry, "providers", boom_providers)
    r = _by_item()
    assert r["q1"]["shipped"] == 4.0 and r["q2"]["attributed"] is False, "一個來源失敗不影響另一個來源，也不丟例外"


# ── ④ 端點 ───────────────────────────────────────────────────────────

def test_endpoint_returns_quantities_only(sq):
    client, h = sq
    _note(client, h, [Qi(5)], status="已核准")
    r = client.get("/api/quotations/%s/item-shipped" % Q, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["items"]["q2"]["shipped"] == 5.0 and d["materialToItem"] == {ITEM: "q1"} and d["totals"]["ordered"] == 26.0
    blob = json.dumps(d, ensure_ascii=False).lower()
    for bad in ("cost", "price", "amount", "total_price", "unitprice", "margin", "金額", "成本"):
        assert bad not in blob, "回應不可有金額／成本欄位：" + bad


def _login_user(client, make_user, name, **kw):
    u, p = make_user(username=name, **kw)[:2]
    r = client.post("/api/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}


def test_endpoint_requires_login_row_access_and_a_shipping_module(sq, make_user):
    """突變對照：拿掉 `require_case` ⇒ 沒有案件權限的人拿到 200 ⇒ 這題紅；拿掉模組檢查 ⇒ 沒有任何出貨相關模組的人拿到 200 ⇒ 紅。"""
    client, h = sq
    url = "/api/quotations/%s/item-shipped" % Q
    assert client.get(url).status_code == 401
    assert client.get("/api/quotations/NO-SUCH-CASE/item-shipped", headers=h).status_code == 404
    outsider = _login_user(client, make_user, "sq_out", role="viewer", modules=["dashboard", "quotation"])       # 有 quotation 模組、但不在這個案件的可見範圍
    assert client.get(url, headers=outsider).status_code == 404, "沒有案件列權限 ⇒ 與查無相同的 404"
    nomod = _login_user(client, make_user, "sq_nomod", role="viewer", modules=["dashboard"])
    assert client.get(url, headers=nomod).status_code == 403, "沒有 case_manage／quotation／financial_view ⇒ 403"


def test_exclude_note_drops_the_edited_notes_own_reservation(sq):
    """編輯中的單（待審核）不把自己的占用算進「其他出貨單」。突變對照：忽略 exclude 參數 ⇒ 紅。"""
    client, h = sq
    _note(client, h, [Qi(5)], status="已核准")
    mine = _note(client, h, [Qi(2)], status="待審核")
    mat = _note(client, h, [L(3)], status="待審核")
    full = client.get("/api/quotations/%s/item-shipped" % Q, headers=h).json()["items"]
    assert full["q2"]["reserved"] == 2.0 and full["q1"]["reserved"] == 3.0
    ex = client.get("/api/quotations/%s/item-shipped?exclude_note=%s" % (Q, mine), headers=h).json()["items"]
    assert ex["q2"]["reserved"] == 0.0 and ex["q2"]["shipped"] == 5.0 and ex["q1"]["reserved"] == 3.0, "只排除那一張"
    ex2 = client.get("/api/quotations/%s/item-shipped?exclude_note=%s" % (Q, mat), headers=h).json()["items"]
    assert ex2["q1"]["reserved"] == 0.0, "材料申請出貨那一邊也要排除"


def test_notes_list_has_each_note_once_even_with_several_lines(sq):
    """同一張單有兩列同品項 ⇒ 單號只出現一次。突變對照：拿掉去重 ⇒ 紅。"""
    client, h = sq
    no = _note(client, h, [Qi(2), Qi(3)], status="已核准")
    c, fn = _prov()
    try:
        out = fn(c, Q)
    finally:
        c.close()
    assert out["q2"]["shipped"] == 5.0 and out["q2"]["notes"] == [no], out
    _set(no, items=[{**L(1), "qty": 1}, {**L(2), "qty": 2}])
    assert _by_item()["q1"]["notes"] == [no]


# ── ⑤ 預設不變 ───────────────────────────────────────────────────────

def test_without_any_shipment_everything_is_zero_and_unattributed(sq):
    r = _by_item()
    assert all(v["shipped"] == 0.0 and v["reserved"] == 0.0 and v["attributed"] is False for v in r.values())
    from core import registry
    import db
    c = db.get_db()
    try:
        assert registry.providers("shipping.material_shipped")["supply"](c, Q) == {}, "既有提供者回傳不變"
    finally:
        c.close()


def test_stamping_the_field_does_not_change_the_existing_material_provider(sq):
    client, h = sq
    _note(client, h, [L(4), Qi(5)], status="已核准")
    from core import registry
    import db
    c = db.get_db()
    try:
        out = registry.providers("shipping.material_shipped")["supply"](c, Q)
    finally:
        c.close()
    assert out == {ITEM: {"reserved": 0.0, "shipped": 4.0, "notes": out[ITEM]["notes"]}}


# ── ⑥ 營運報表匯出 ───────────────────────────────────────────────────

def test_case_shipped_summary_provider_and_report_helper(sq):
    client, h = sq
    _note(client, h, [Qi(5), Qi(2, qid="q3")], status="已核准")
    from core import registry
    from modules.analytics.api import reports as R
    fn = registry.single_provider("case.shipped_summary")
    import db
    c = db.get_db()
    try:
        assert fn(c, Q) == {"ordered": 26.0, "shipped": 7.0, "reserved": 0.0}
        assert fn(c, "NO-SUCH") == {"ordered": 0.0, "shipped": 0.0, "reserved": 0.0}
    finally:
        c.close()
    assert R._case_ship_summaries([Q]) == {Q: {"ordered": 26.0, "shipped": 7.0}}


def test_report_helper_is_blank_when_the_provider_is_missing_or_fails(sq, monkeypatch):
    from modules.analytics.api import reports as R
    monkeypatch.setattr(R._registry, "single_provider", lambda name: None)
    assert R._case_ship_summaries([Q]) == {}
    monkeypatch.setattr(R._registry, "single_provider", lambda name: (lambda conn, qn: (_ for _ in ()).throw(RuntimeError("x"))))
    assert R._case_ship_summaries([Q]) == {}, "單一案件失敗只留白、不丟例外"


def test_excel_export_columns_order_rows_and_totals(sq, monkeypatch):
    """「毛利分析」最右兩欄：已出貨數量合計、報價品項數量合計（順序同畫面「已出貨/數量」）；逐案列與合計列的值都要對（突變：欄位互換 ⇒ 紅）。"""
    import io
    from openpyxl import load_workbook
    from modules.analytics.api import reports as R
    client, h = sq
    _note(client, h, [Qi(5), Qi(2, qid="q3")], status="已核准")                           # Q：已出貨 7 / 訂購 26
    real = R._collect

    def fake_collect(d0, d1, dept=None):
        d = real(d0, d1, dept)
        base = {"customer": "客戶", "project": "專案", "salesPerson": "業務", "dealTag": "已成案", "pretax": 1000, "netMarginPct": 10.0, "actualMarginPct": 12.0,
                "grossProfit": 120, "settleSummary": {}, "settleStatus": "finalized", "settleDate": "", "settleBy": ""}
        d["marginCases"] = [{**base, "quoteNo": Q}, {**base, "quoteNo": "NO-SHIP-CASE"}]
        return d
    monkeypatch.setattr(R, "_collect", fake_collect)
    r = client.get("/api/reports/financial/excel?period=2026&basis=accrual", headers=h)
    assert r.status_code == 200, r.text[:200]
    ws = load_workbook(io.BytesIO(r.content))["毛利分析"]
    hdr = [c.value for c in ws[3]]
    assert hdr[21:26] == ["報價預留間接成本", "其中：預留未被實際成本抵用", "其中：其他", "已出貨數量合計", "報價品項數量合計"], hdr
    assert hdr[0] == "案件號" and len([x for x in hdr if x]) == 26, "既有欄位不動、只在最右加兩欄"
    row = {ws.cell(row=r_i, column=1).value: [ws.cell(row=r_i, column=c).value for c in (25, 26)] for r_i in (4, 5, 6)}
    assert row[Q] == [7, 26], row                                                      # 已出貨 7、訂購 26（欄位順序：已出貨在前）
    assert row["NO-SHIP-CASE"] == [0, 0], row                                          # 案件不存在／沒有品項 ⇒ 0（不是空白也不是別案的值）
    assert row["合計"] == [7, 26], row



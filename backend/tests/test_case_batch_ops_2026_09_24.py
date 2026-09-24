"""案件清單批次操作（2026-09-24 使用者表單）：多選後批次改執行負責／成員、批次匯出。

- POST /api/case-batch/assign {quote_nos, executor?, add_members?, remove_members?}
  管理員以上（與單筆成員分配同一個門檻）；只改進行中（已成案）的案件，已結案的列在 skipped 並說明原因；
  指定的帳號不存在或已停用 ⇒ 400，整批不寫
- POST /api/case-batch/export {quote_nos} ⇒ xlsx；只含呼叫者看得到的案件；看不到金額的帳號金額欄留空
路徑刻意不放在 /api/quotations/ 底下：/api/quotations/{quote_no}/export 會把 batch 當成單號。
"""
import io
import json

import pytest

ASSIGN = "/api/case-batch/assign"
EXPORT = "/api/case-batch/export"


def _login(client, username, password):
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _uid(username):
    import db
    conn = db.get_db()
    try:
        return conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()[0]
    finally:
        conn.close()


def _case(no, *, deal_tag="已成案", assigned=(), sales="", total=1000):
    import db
    now = "2026-01-01T00:00:00"
    cr = {"roles": {"filler": "", "sales": "", "executor": ""}, "payment": {"items": []}}
    conn = db.get_db()
    try:
        conn.execute(
            "INSERT INTO quotations (quote_no, status, customer_name, project_name, total, pretax, data_json,"
            " created_at, updated_at, deal_tag, quote_date, sales_person, assigned_user_ids)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (no, "已送出", "客戶" + no[-1], "專案", total, total,
             json.dumps({"dealTag": deal_tag, "caseRecord": cr}, ensure_ascii=False),
             now, now, deal_tag, "2026-08-01", sales, json.dumps(list(assigned))))
        conn.commit()
    finally:
        conn.close()


def _row(no):
    import db
    conn = db.get_db()
    try:
        r = conn.execute("SELECT data_json, assigned_user_ids FROM quotations WHERE quote_no=?", (no,)).fetchone()
        return json.loads(r["data_json"]), json.loads(r["assigned_user_ids"] or "[]")
    finally:
        conn.close()


def test_batch_sets_executor_and_members_on_active_cases_only(client, make_user):
    h = _login(client, *make_user(username="bt_admin", role="admin"))
    make_user(username="bt_exec", role="engineer")
    make_user(username="bt_m1", role="engineer")
    make_user(username="bt_m2", role="engineer")
    m1, m2 = _uid("bt_m1"), _uid("bt_m2")
    _case("MQ-BT-1", assigned=[m2])
    _case("MQ-BT-2")
    _case("MQ-BT-C", deal_tag="已結案")
    r = client.post(ASSIGN, headers=h, json={"quote_nos": ["MQ-BT-1", "MQ-BT-2", "MQ-BT-C"],
                                             "executor": "bt_exec", "add_members": [m1], "remove_members": [m2]})
    assert r.status_code == 200, r.text
    body = r.json()
    assert sorted(body["updated"]) == ["MQ-BT-1", "MQ-BT-2"]
    assert [s["quoteNo"] for s in body["skipped"]] == ["MQ-BT-C"] and body["skipped"][0]["reason"]
    for no in ("MQ-BT-1", "MQ-BT-2"):
        d, members = _row(no)
        assert d["caseRecord"]["roles"]["executor"]["username"] == "bt_exec"
        assert members == [m1]
    d, members = _row("MQ-BT-C")
    assert d["caseRecord"]["roles"]["executor"] == "" and members == []


def test_batch_assign_needs_admin(client, make_user):
    h = _login(client, *make_user(username="bt_sales", role="sales"))
    _case("MQ-BT-S", sales="bt_sales")
    r = client.post(ASSIGN, headers=h, json={"quote_nos": ["MQ-BT-S"], "add_members": [1]})
    assert r.status_code == 403, r.text


def test_batch_assign_rejects_unknown_executor_and_writes_nothing(client, make_user):
    h = _login(client, *make_user(username="bt_admin2", role="admin"))
    _case("MQ-BT-X")
    r = client.post(ASSIGN, headers=h, json={"quote_nos": ["MQ-BT-X"], "executor": "no_such_user"})
    assert r.status_code == 400, r.text
    d, _ = _row("MQ-BT-X")
    assert d["caseRecord"]["roles"]["executor"] == ""


def test_batch_export_is_xlsx_with_visible_cases_and_masked_amounts(client, make_user):
    from openpyxl import load_workbook
    h = _login(client, *make_user(username="bt_admin3", role="admin"))
    _case("MQ-BT-E1", total=1234)
    _case("MQ-BT-E2", total=5678)
    r = client.post(EXPORT, headers=h, json={"quote_nos": ["MQ-BT-E1", "MQ-BT-E2"]})
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    rows = {row[0]: dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"MQ-BT-E1", "MQ-BT-E2"}
    assert rows["MQ-BT-E1"]["金額"] == 1234

    # 工程師：只看得到被分配的那一件，而且金額留空
    make_user(username="bt_eng", role="engineer", modules=["case_manage"])
    import db
    conn = db.get_db()
    try:
        conn.execute("UPDATE quotations SET assigned_user_ids=? WHERE quote_no='MQ-BT-E1'",
                     (json.dumps([_uid("bt_eng")]),))
        conn.commit()
    finally:
        conn.close()
    he = _login(client, "bt_eng", "Test-Pass-123")
    r = client.post(EXPORT, headers=he, json={"quote_nos": ["MQ-BT-E1", "MQ-BT-E2"]})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert [x["單號"] for x in rows] == ["MQ-BT-E1"]
    assert rows[0]["金額"] is None


def test_batch_size_limit(client, make_user):
    h = _login(client, *make_user(username="bt_admin4", role="admin"))
    r = client.post(EXPORT, headers=h, json={"quote_nos": [f"MQ-{i}" for i in range(201)]})
    assert r.status_code == 400, r.text

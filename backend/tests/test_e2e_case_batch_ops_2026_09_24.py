"""案件清單批次操作：頁面端（2026-09-24 使用者表單）。

「多選」⇒ 卡片出現勾選框；勾兩件、選執行負責、套用 ⇒ 資料庫兩件都改；匯出 ⇒ 下載 xlsx 且只含勾選的。
非管理員看不到批次改負責人／成員（只剩匯出）。觀測點：資料庫落地值、下載檔內容。
"""
from tests._requires import requires_module, skip_module_unless  # noqa: E402  M01 ④(c)（稽核 D M4-M3）
import io
import json
import threading
import time

import pytest

pytest.importorskip("playwright.sync_api")
from tests._e2e_login import inject_login  # noqa: E402
pytestmark = requires_module("case", '本檔的題打 M01（案件）的端點或讀寫 M01 的資料（報價單／案件）；M01 不在時沒有對象（稽核 D M4-M3）')

DATA_JS = "Alpine.$data(document.querySelector('[x-data]'))"


def _case(no, sales=""):
    import db
    now = "2026-01-01T00:00:00"
    cr = {"roles": {"filler": "", "sales": "", "executor": ""}, "payment": {"items": []}}
    conn = db.get_db()
    try:
        conn.execute(
            "INSERT INTO quotations (quote_no, status, customer_name, project_name, total, pretax, data_json,"
            " created_at, updated_at, deal_tag, quote_date, sales_person) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (no, "已送出", "客戶", "專案", 1000, 952,
             json.dumps({"dealTag": "已成案", "caseRecord": cr}, ensure_ascii=False),
             now, now, "已成案", "2026-08-01", sales))
        conn.commit()
    finally:
        conn.close()


def _executor(no):
    import db
    conn = db.get_db()
    try:
        d = json.loads(conn.execute("SELECT data_json FROM quotations WHERE quote_no=?", (no,)).fetchone()[0])
        return (d["caseRecord"]["roles"]["executor"] or {}).get("username") if isinstance(
            d["caseRecord"]["roles"]["executor"], dict) else d["caseRecord"]["roles"]["executor"]
    finally:
        conn.close()




def _open(browser, base, user):
    page = browser.new_context(accept_downloads=True).new_page()
    page.on("dialog", lambda d: d.accept())
    inject_login(page, base, user[0], user[1])
    page.goto(f"{base}/pages/case-management.html")
    page.wait_for_function(f"() => {DATA_JS} && {DATA_JS}.session && {DATA_JS}.session.token"
                           f" && !{DATA_JS}.loading && {DATA_JS}.selectableUsers.length", timeout=20000)
    return page


def _check(page, no):
    page.locator(f".cm-card[data-quote-no='{no}'] [data-testid=batch-check]").click()


@pytest.mark.e2e
def test_batch_executor_and_export(live_server, make_user, e2e_browser):
    from openpyxl import load_workbook
    u = make_user(username="be_admin", role="admin")
    make_user(username="be_exec", role="engineer")
    for no in ("MQ-BE-1", "MQ-BE-2", "MQ-BE-3"):
        _case(no)
    browser = e2e_browser
    page = _open(browser, live_server, u)
    page.click("[data-testid=batch-toggle]")
    _check(page, "MQ-BE-1")
    _check(page, "MQ-BE-3")
    assert page.locator("[data-testid=batch-count]").inner_text() == "已選 2 件"
    assert page.evaluate(f"() => {DATA_JS}.selected") is None, "勾選不可以順便打開案件"
    page.select_option("[data-testid=batch-exec]", "be_exec")
    page.click("[data-testid=batch-exec-apply]")
    page.locator("[data-testid=batch-msg]").wait_for(state="visible", timeout=10000)
    assert "已變更 2 件" in page.locator("[data-testid=batch-msg]").inner_text()
    assert (_executor("MQ-BE-1"), _executor("MQ-BE-2"), _executor("MQ-BE-3")) == ("be_exec", "", "be_exec")
    with page.expect_download() as dl:
        page.click("[data-testid=batch-export]")
    path = dl.value.path()
    with open(path, "rb") as f:           # 下載暫存檔沒有副檔名，openpyxl 依副檔名拒讀
        ws = load_workbook(io.BytesIO(f.read())).active
    assert sorted(r[0] for r in ws.iter_rows(min_row=2, values_only=True)) == ["MQ-BE-1", "MQ-BE-3"]


@pytest.mark.e2e
def test_non_admin_sees_export_only(live_server, make_user, e2e_browser):
    u = make_user(username="be_sales", role="sales")
    _case("MQ-BE-S", sales="be_sales")
    browser = e2e_browser
    page = _open(browser, live_server, u)
    page.click("[data-testid=batch-toggle]")
    _check(page, "MQ-BE-S")
    page.locator("[data-testid=batch-export]").wait_for(state="visible", timeout=5000)
    assert page.locator("[data-testid=batch-exec]").count() == 0
    assert page.locator("[data-testid=batch-member]").count() == 0

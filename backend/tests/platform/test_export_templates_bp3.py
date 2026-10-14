# -*- coding: utf-8 -*-
"""BP-3 匯出範本（R0 模型＋驗證、R1 先導匯出、R2 預覽與治理）。設定 builder_exports.enabled 預設關。

- R0：範本只收宣告式資料（欄位 id 必須在欄位目錄、標題是純文字、沒有占位符／運算式／控制字元）；CSV 沒有頁首頁尾／浮水印；浮水印不得含官方狀態字樣
- 遮蔽優先：money／有 visible 的欄位在輸出端再驗一次權限，沒權限 ⇒ 整欄省略；money 沒有 visible ⇒ 一律不顯示（fail-closed）
- R1（每日任務歷程 CSV）：旗標關 ⇒ 輸出與預設完全相同，已發布的範本也不套用；開且有範本 ⇒ 欄位／標題／順序依範本、每格仍經 csv_safe
- R2：預覽端點（關 404、非最高管理者 403、不合法 422、不寫資料）；治理開啟時 export_template 發布要帶預覽雜湊、放寬（新增欄位）變待核准
"""
import csv
import io

import pytest

from helpers import export_templates as ET


# ── 測試用的匯出目錄（不依賴任何 L2 模組）────────────────────────────────
def _is_super(u):
    return (u or {}).get("role") == "superadmin"


ET.register_export(
    "t_bp3:csv", label="測試 CSV", formats=("csv",),
    columns=[{"id": "name", "label": "名稱"}, {"id": "note", "label": "備註"},
             {"id": "amount", "label": "金額", "money": True, "visible": _is_super},
             {"id": "secret", "label": "隱藏", "money": True}],                       # money 沒有 visible ⇒ fail-closed
    sample_rows=[{"name": "甲", "note": "=cmd", "amount": 100, "secret": "x"}, {"name": "乙", "note": "-VLAN", "amount": 200, "secret": "y"}])
ET.register_export(
    "t_bp3:pdf", label="測試 PDF", formats=("pdf", "xlsx"),
    columns=[{"id": "a", "label": "甲欄"}, {"id": "b", "label": "乙欄"}], sample_rows=[{"a": 1, "b": 2}])


def _tpl(cols, **kw):
    return dict({"version": 1, "format": "csv", "columns": [{"id": i, "label": l} for i, l in cols]}, **kw)


def _bad(body, key="t_bp3:csv"):
    return ET.definition_problems(key, body)


def _paths(probs):
    return {p["path"] for p in probs}


# ── R0：模型與驗證 ────────────────────────────────────────────────────────
def test_default_template_is_valid_and_matches_the_catalog_order():
    d = ET.default_template("t_bp3:csv")
    assert d["columns"] == [{"id": "name", "label": "名稱"}, {"id": "note", "label": "備註"}, {"id": "amount", "label": "金額"}, {"id": "secret", "label": "隱藏"}]
    assert _bad(d) == []


@pytest.mark.parametrize("body,path", [
    ({"version": 2, "format": "csv", "columns": [{"id": "name", "label": "n"}]}, "version"),
    ({"version": 1, "format": "xlsx", "columns": [{"id": "name", "label": "n"}]}, "format"),
    ({"version": 1, "format": "csv", "columns": []}, "columns"),
    ({"version": 1, "format": "csv", "columns": [{"id": "nope", "label": "n"}]}, "columns[0].id"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "a"}, {"id": "name", "label": "b"}]}, "columns[1].id"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "同"}, {"id": "note", "label": "同"}]}, "columns[1].label"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "{{x}}"}]}, "columns[0].label"),            # 占位符
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "<b>x"}]}, "columns[0].label"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "=SUM(A1)"}]}, "columns[0].label"),          # 開頭觸發字元
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "x" * 21}]}, "columns[0].label"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "a‮b"}]}, "columns[0].label"),          # 雙向覆寫
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "a\nb"}]}, "columns[0].label"),             # 控制字元
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "n", "formula": "a+b"}]}, "columns[0].formula"),  # 沒有運算式
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "n", "format": "#,##0"}]}, "columns[0].format"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "n"}], "script": "x"}, "script"),
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "n"}], "header": {"text": "H"}}, "header"),       # CSV 沒有頁首
    ({"version": 1, "format": "csv", "columns": [{"id": "name", "label": "n"}], "watermark": {"text": "W"}}, "watermark"),
])
def test_validator_rejects(body, path):
    assert path in _paths(_bad(body)), _bad(body)


def test_validator_rejects_unknown_export_and_non_object():
    assert _bad(_tpl([("name", "n")]), key="nope:x")
    assert _bad([1, 2])


def test_header_footer_watermark_only_where_the_format_has_them_and_reserved_words_are_refused():
    base = {"version": 1, "format": "pdf", "columns": [{"id": "a", "label": "甲"}]}
    assert _bad(dict(base, header={"text": "內部資料"}, footer={"text": "第一頁"}, watermark={"text": "內部使用"}), key="t_bp3:pdf") == []
    xl = dict(base, format="xlsx", header={"text": "H"})
    assert _bad(xl, key="t_bp3:pdf") == []
    assert "watermark" in _paths(_bad(dict(xl, watermark={"text": "W"}), key="t_bp3:pdf")), "xlsx 沒有浮水印"
    for w in ("正本", "已核准", "已 付 款", "已簽回", "作廢", "已核可"):
        assert "watermark.text" in _paths(_bad(dict(base, watermark={"text": w}), key="t_bp3:pdf")), w
    assert "watermark.text" in _paths(_bad(dict(base, watermark={"text": "已核准"}), key="t_bp3:pdf")), "全形／近似也擋（NFKC）"
    assert "watermark.text" in _paths(_bad(dict(base, watermark={"text": "ｗ" * 41}), key="t_bp3:pdf"))
    assert "header.text" in _paths(_bad(dict(base, header={"text": "{x}"}), key="t_bp3:pdf"))
    assert "header.x" in _paths(_bad(dict(base, header={"text": "H", "x": 1}), key="t_bp3:pdf"))


# ── 遮蔽優先 ──────────────────────────────────────────────────────────────
def test_money_columns_need_the_permission_check_and_fail_closed_without_one():
    t = _tpl([("name", "名稱"), ("amount", "金額"), ("secret", "隱藏"), ("note", "備註")])
    cols, omitted = ET.effective_columns("t_bp3:csv", t, {"role": "superadmin"})
    assert [c["id"] for c in cols] == ["name", "amount", "note"]
    assert omitted == [{"id": "secret", "reason": "這位使用者沒有這個欄位的檢視權限"}], "money 沒有 visible ⇒ 對誰都不顯示"
    cols, omitted = ET.effective_columns("t_bp3:csv", t, {"role": "viewer"})
    assert [c["id"] for c in cols] == ["name", "note"] and {o["id"] for o in omitted} == {"amount", "secret"}
    # visible 丟例外 ⇒ 當沒權限
    ET._EXPORTS["t_bp3:csv"]["columns"][2]["visible"] = lambda u: 1 / 0
    try:
        cols, _ = ET.effective_columns("t_bp3:csv", t, {"role": "superadmin"})
        assert "amount" not in [c["id"] for c in cols]
    finally:
        ET._EXPORTS["t_bp3:csv"]["columns"][2]["visible"] = _is_super


def test_render_csv_rows_reorders_renames_and_neutralizes_every_cell():
    t = _tpl([("note", "備註欄"), ("name", "名稱")])
    rows = [{"name": "=HYPERLINK(1)", "note": "-VLAN", "amount": 5}, {"name": "乙", "note": None}]
    table, omitted = ET.render_csv_rows("t_bp3:csv", t, rows, {"role": "viewer"})
    assert table == [["備註欄", "名稱"], ["'-VLAN", "'=HYPERLINK(1)"], [None, "乙"]] and omitted == []
    assert not any("amount" in str(r) for r in table), "範本沒選的欄位不會輸出（範本不能新增資料，只能選目錄裡的）"


def test_default_template_renders_the_same_header_and_cells_as_the_hard_coded_export():
    """R1 的『預設＝今天』證明：預設範本渲染 ＝ 端點原本寫死的標題與欄位順序。"""
    ET_ID = "daily_tasks:task-history"
    import modules.daily_tasks.api as dt                      # 註冊 ET 的匯出目錄
    d = ET.default_template(ET_ID)
    row = {c: "v%d" % i for i, c in enumerate(dt._HISTORY_COLS)}
    table, omitted = ET.render_csv_rows(ET_ID, d, [row], {"role": "superadmin"})
    assert table[0] == ["工作事項", "日期", "星期", "指派人員", "完成狀態", "完成時間", "回報內容"] and omitted == []
    assert table[1] == ["v%d" % i for i in range(7)]


# ── API 共用 ──────────────────────────────────────────────────────────────
def _login(client, make_user, name, role="superadmin", modules=None):
    u, pw = make_user(username=name, role=role, **({"modules": modules} if modules is not None else {}))[:2]
    return u, {"Authorization": "Bearer " + client.post("/api/auth/login", json={"username": u, "password": pw}).json()["token"]}


def _flag(monkeypatch, on, governance=False):
    monkeypatch.setattr(ET, "enabled", lambda conn=None: on)
    import routers.def_governance as G
    import routers.def_pipeline as P
    import routers.definitions as DF
    monkeypatch.setattr(G, "enabled", lambda conn=None: governance)
    monkeypatch.setattr(P, "enabled", lambda conn=None: governance)
    monkeypatch.setattr(DF, "_governance_enabled", lambda conn=None: governance)


EID = "daily_tasks:task-history"


def _seed_task(client, hd, username, title="=巡檢", report="-VLAN10 異常"):
    from datetime import date
    r = client.post("/api/daily-tasks", json={"task_date": date.today().isoformat(), "title": title, "assigned_to": [username]}, headers=hd)
    assert r.status_code == 201, r.text
    tid = r.json()["id"]
    c = client.patch("/api/daily-tasks/%d/complete" % tid, json={"completed": True, "report": report}, headers=hd)
    assert c.status_code == 200, c.text
    return tid


def _export(client, hd, tid):
    r = client.get("/api/daily-tasks/%d/history/export" % tid, headers=hd)
    assert r.status_code == 200, r.text
    return r.content


def _publish(client, hd, body, key=EID, note="t", preview_hash=None):
    url = "/api/definitions/export_template/%s" % key
    s = client.put(url + "/draft?scope=company", json={"body": body}, headers=hd)
    assert s.status_code == 200, s.text
    payload = {"note": note}
    if preview_hash:
        payload["preview_hash"] = preview_hash
    return client.post(url + "/publish?scope=company", json=payload, headers=hd)


def _rows(raw):
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))


# ── R1：先導匯出 ──────────────────────────────────────────────────────────
def test_flag_off_export_is_identical_to_default_and_ignores_a_published_template(client, make_user, monkeypatch):
    u, hd = _login(client, make_user, "bp3_sa", "superadmin")
    tid = _seed_task(client, hd, u)
    _flag(monkeypatch, False)
    off = _export(client, hd, tid)
    rows = _rows(off)
    assert rows[0] == ["工作事項", "日期", "星期", "指派人員", "完成狀態", "完成時間", "回報內容"], "旗標關＝今天寫死的標題"
    assert rows[1][0] == "'=巡檢" and rows[1][6] == "'-VLAN10 異常", "公式中和（csv_safe）照舊"
    assert off.startswith(b"\xef\xbb\xbf"), "UTF-8 BOM 照舊"
    _flag(monkeypatch, True)
    assert _export(client, hd, tid) == off, "旗標開、沒有已發布範本 ⇒ 仍與預設逐位元相同"
    r = _publish(client, hd, _tpl([("report", "回報"), ("title", "事項")]))
    assert r.status_code == 200 and not r.json().get("pending"), r.text
    _flag(monkeypatch, False)
    assert _export(client, hd, tid) == off, "旗標關回去 ⇒ 已發布的範本不再套用，輸出與預設逐位元相同"


def test_flag_on_published_template_drives_columns_labels_order_and_keeps_neutralization(client, make_user, monkeypatch):
    u, hd = _login(client, make_user, "bp3_sa2", "superadmin")
    tid = _seed_task(client, hd, u)
    _flag(monkeypatch, True)
    r = _publish(client, hd, _tpl([("report", "回報內容"), ("title", "事項"), ("date", "日期")]))
    assert r.status_code == 200, r.text
    rows = _rows(_export(client, hd, tid))
    assert rows[0] == ["回報內容", "事項", "日期"], "欄位／標題／順序依範本；沒選的欄位不輸出"
    assert rows[1][0] == "'-VLAN10 異常" and rows[1][1] == "'=巡檢", "範本路徑的每個儲存格仍經 csv_safe"
    assert len(rows) == 2


# ── R2：預覽與定義端點 ────────────────────────────────────────────────────
def test_preview_is_404_when_off_403_for_non_superadmin_422_when_invalid_and_writes_nothing(client, make_user, monkeypatch):
    import db
    _, h_sa = _login(client, make_user, "bp3_sa3", "superadmin")
    _, h_ad = _login(client, make_user, "bp3_ad", "admin", modules=["daily_task"])
    url = "/api/export-templates/%s/preview" % EID
    good = {"template": ET.default_template(EID)}
    _flag(monkeypatch, False)
    assert client.post(url, json=good, headers=h_sa).status_code == 404
    _flag(monkeypatch, True)
    assert client.post(url, json=good, headers=h_ad).status_code == 403
    assert client.post(url, json=good).status_code in (401, 403)
    assert client.post("/api/export-templates/nope/preview", json=good, headers=h_sa).status_code == 404
    bad = client.post(url, json={"template": _tpl([("nope", "x")])}, headers=h_sa)
    assert bad.status_code == 422 and bad.json()["detail"]["problems"]
    assert client.post(url, json={"template": good["template"], "n": "x"}, headers=h_sa).status_code == 422

    def count():
        c = db.get_db()
        try:
            return c.execute("SELECT (SELECT COUNT(*) FROM ui_definitions)+(SELECT COUNT(*) FROM audit_log WHERE action LIKE 'export.%')").fetchone()[0]
        finally:
            c.close()
    n0 = count()
    ok = client.post(url, json={"template": _tpl([("title", "事項"), ("report", "回報")])}, headers=h_sa)
    assert ok.status_code == 200, ok.text
    d = ok.json()
    assert d["columns"] == ["事項", "回報"] and d["sample"] == "sample" and d["rows"][0][0] == "每週設備巡檢"
    assert count() == n0, "預覽不寫定義、不留 export.* 稽核、不計匯出"


def test_preview_as_another_user_uses_that_users_visible_rows(client, make_user, monkeypatch):
    u, h_sa = _login(client, make_user, "bp3_sa4", "superadmin")
    other, _ = _login(client, make_user, "bp3_other", "admin", modules=["daily_task"])
    _seed_task(client, h_sa, u, title="只有我的")
    _flag(monkeypatch, True)
    url = "/api/export-templates/%s/preview" % EID
    t = {"template": _tpl([("title", "事項")]), "sample": "visible"}
    mine = client.post(url, json=t, headers=h_sa).json()
    assert mine["sample"] == "visible" and any("只有我的" in r[0] for r in mine["rows"])
    theirs = client.post(url, json=dict(t, user=other), headers=h_sa).json()
    assert theirs["user"] == other and theirs["sample"] == "sample" and theirs["warnings"], "別人沒被指派 ⇒ 看不到這筆，退回樣本並說明"
    assert not any("只有我的" in r[0] for r in theirs["rows"])
    assert client.post(url, json=dict(t, user="no_such_user"), headers=h_sa).status_code == 404


def test_status_endpoint_hides_the_editor_unless_on_and_superadmin(client, make_user, monkeypatch):
    _, h_sa = _login(client, make_user, "bp3_sa5", "superadmin")
    _, h_ad = _login(client, make_user, "bp3_ad2", "admin", modules=["daily_task"])
    url = "/api/export-templates/%s/status" % EID
    _flag(monkeypatch, False)
    assert client.get(url, headers=h_sa).json() == {"enabled": False, "canEdit": False}
    _flag(monkeypatch, True)
    s = client.get(url, headers=h_sa).json()
    assert s["enabled"] and s["canEdit"] and [c["id"] for c in s["columns"]] == ["title", "date", "weekday", "assignee", "status", "at", "report"]
    a = client.get(url, headers=h_ad).json()
    assert a == {"enabled": True, "canEdit": False}, "非最高管理者拿不到欄位目錄"
    assert client.get("/api/export-templates/nope/status", headers=h_sa).status_code == 404
    assert client.get(url).status_code in (401, 403)


def test_definition_is_admin_only_and_rejected_while_off(client, make_user, monkeypatch):
    _, h_sa = _login(client, make_user, "bp3_sa6", "superadmin")
    _, h_ad = _login(client, make_user, "bp3_ad3", "admin", modules=["daily_task"])
    d = "/api/definitions/export_template/%s" % EID
    _flag(monkeypatch, True)
    for who in (h_ad,):
        assert client.get(d + "?scope=company", headers=who).status_code in (401, 403)
        assert client.put(d + "/draft?scope=company", json={"body": ET.default_template(EID)}, headers=who).status_code in (401, 403)
    _flag(monkeypatch, False)
    v = client.post(d + "/validate", json={"body": ET.default_template(EID)}, headers=h_sa)
    assert v.status_code == 200 and any(ET.OFF_REASON in p["message"] for p in v.json()["problems"]), "旗標關：發布驗證擋"
    kinds = client.get("/api/definition-kinds", headers=h_sa).json()
    kinds = kinds.get("kinds", kinds) if isinstance(kinds, dict) else kinds
    assert "export_template" in [k["kind"] if isinstance(k, dict) else k for k in kinds], "kind 有登記（版本／還原／帳本共用定義文件庫）"


# ── R2：治理（版本／預覽雜湊／待核准／還原）──────────────────────────────
def test_governance_requires_preview_hash_and_turns_added_columns_into_pending(client, make_user, monkeypatch):
    u, hd = _login(client, make_user, "bp3_sa7", "superadmin")
    tid = _seed_task(client, hd, u)
    _flag(monkeypatch, True, governance=False)
    assert _publish(client, hd, _tpl([("title", "事項"), ("date", "日期")])).status_code == 200      # 基準：兩欄（治理關 ⇒ 直接發布）
    narrow = _rows(_export(client, hd, tid))
    assert narrow[0] == ["事項", "日期"]
    _flag(monkeypatch, True, governance=True)
    wide = _tpl([("title", "事項"), ("date", "日期"), ("report", "回報內容")])
    url = "/api/definitions/export_template/%s" % EID
    assert client.put(url + "/draft?scope=company", json={"body": wide}, headers=hd).status_code == 200
    no_hash = client.post(url + "/publish?scope=company", json={"note": "x"}, headers=hd)
    assert no_hash.status_code == 409 and no_hash.json()["code"] == "preview_required", "名單內的 kind：治理開啟後發布必須先看過預覽"
    im = client.get("/api/def-governance/export_template/%s/impact?scope=company" % EID, headers=hd)
    assert im.status_code == 200 and im.json()["previewHash"], im.text
    r = client.post(url + "/publish?scope=company", json={"note": "x", "preview_hash": im.json()["previewHash"]}, headers=hd)
    d = r.json()
    assert r.status_code == 200 and d["pending"] is True and d["changeId"] and d["effectiveAt"] and any("新增匯出欄位 report" in x for x in d["reasons"]), d
    assert _rows(_export(client, hd, tid))[0] == ["事項", "日期"], "待核准的放寬還沒生效：匯出仍是舊範本"


def test_governance_tightening_publishes_immediately_and_restore_goes_through_the_same_pipeline(client, make_user, monkeypatch):
    u, hd = _login(client, make_user, "bp3_sa8", "superadmin")
    tid = _seed_task(client, hd, u)
    _flag(monkeypatch, True, governance=False)
    assert _publish(client, hd, _tpl([("title", "事項"), ("date", "日期"), ("report", "回報")])).status_code == 200       # v1（三欄）
    _flag(monkeypatch, True, governance=True)
    url = "/api/definitions/export_template/%s" % EID
    client.put(url + "/draft?scope=company", json={"body": _tpl([("title", "事項"), ("date", "日期")])}, headers=hd)
    h = client.get("/api/def-governance/export_template/%s/impact?scope=company" % EID, headers=hd).json()["previewHash"]
    r = client.post(url + "/publish?scope=company", json={"note": "收緊", "preview_hash": h}, headers=hd)
    assert r.status_code == 200 and not r.json().get("pending"), r.text           # 移除欄位＝收緊 ⇒ 立即生效（照舊）
    assert _rows(_export(client, hd, tid))[0] == ["事項", "日期"]
    back = client.post(url + "/restore/1?scope=company", json={}, headers=hd)       # 還原到三欄版＝放寬 ⇒ 待核准，不立即生效
    assert back.status_code == 200 and back.json().get("pending") is True, back.text
    assert _rows(_export(client, hd, tid))[0] == ["事項", "日期"]


def test_pipeline_comparator_classes():
    from routers.def_pipeline import classify
    a = _tpl([("title", "事項"), ("date", "日期")])
    assert classify("export_template", a, _tpl([("date", "日期"), ("title", "標題改了")]))[0] == "neutral", "只改順序／標題 ⇒ 中性"
    assert classify("export_template", a, _tpl([("title", "事項")]))[0] == "tighten"
    c, why = classify("export_template", a, _tpl([("title", "事項"), ("date", "日期"), ("report", "回報")]))
    assert c == "loosen" and any("report" in x for x in why)
    assert classify("export_template", a, dict(a, format="xlsx"))[0] == "loosen"
    assert classify("export_template", a, dict(a, extra=1))[0] == "loosen", "無法判斷 ⇒ 放寬"
    with_wm = dict(_tpl([("title", "t")]), watermark={"text": "W"})
    assert classify("export_template", with_wm, _tpl([("title", "t")]))[0] == "tighten", "拿掉浮水印＝收緊"


def test_setting_default_is_off_and_registered():
    from helpers import settings_groups  # noqa: F401
    from helpers import settings_registry as SR
    f = {x["key"]: x for x in SR.groups()["builder_exports"]["fields"]}["enabled"]
    assert f["default"] == 0 and (f["min"], f["max"]) == (0, 1)
    assert ET.enabled() is False, "預設關"


def test_menu_response_gains_builder_exports_only_when_on_and_superadmin(client, make_user, monkeypatch):
    """OFF 時頁面不載入編輯器、不多任何請求：旗標由已經會發的 /api/platform/menu 帶（鍵只在開啟＋最高管理者時才出現）。"""
    _, h_sa = _login(client, make_user, "bp3_sa9", "superadmin")
    _, h_ad = _login(client, make_user, "bp3_ad4", "admin", modules=["daily_task"])
    _flag(monkeypatch, False)
    assert "builderExports" not in client.get("/api/platform/menu", headers=h_sa).json()
    _flag(monkeypatch, True)
    assert client.get("/api/platform/menu", headers=h_sa).json().get("builderExports") is True
    assert "builderExports" not in client.get("/api/platform/menu", headers=h_ad).json()


def test_page_declares_the_export_id_and_only_sidebar_loads_the_editor():
    from pathlib import Path
    root = Path(__file__).resolve().parents[3] / "frontend"
    html = (root / "pages" / "daily-tasks.html").read_text(encoding="utf-8")
    assert '<meta name="motrix-export-id" content="daily_tasks:task-history">' in html
    assert "export-editor.js" not in html, "頁面不直接載入編輯器（關著時不能多請求）"
    sb = (root / "static" / "sidebar.js").read_text(encoding="utf-8")
    assert "d.builderExports" in sb and "/static/export-editor.js" in sb


# ═══ R3：xlsx 範本（頁首／頁尾＝列印頁首頁尾；標題鎖定的匯出；先導＝外包名冊）═══════════════════════
ET.register_export(
    "t_bp3:xlsx", label="測試 XLSX", formats=("xlsx",), lock_labels=True,
    columns=[{"id": "name", "label": "姓名"}, {"id": "phone", "label": "電話"}], sample_rows=[{"name": "甲", "phone": "+886 2"}])


def test_lock_labels_blocks_renaming_but_allows_omit_and_reorder_and_print_text():
    ok = {"version": 1, "format": "xlsx", "columns": [{"id": "phone", "label": "電話"}, {"id": "name", "label": "姓名"}],
          "header": {"text": "外包名冊"}, "footer": {"text": "內部資料"}}
    assert _bad(ok, key="t_bp3:xlsx") == []
    assert _bad(dict(ok, columns=[{"id": "name", "label": "姓名"}]), key="t_bp3:xlsx") == []
    renamed = dict(ok, columns=[{"id": "name", "label": "全名"}])
    assert "columns[0].label" in _paths(_bad(renamed, key="t_bp3:xlsx")), "匯入靠表頭文字對應 ⇒ 標題不能改"
    assert "watermark" in _paths(_bad(dict(ok, watermark={"text": "W"}), key="t_bp3:xlsx")), "xlsx 沒有浮水印"
    assert "header.text" in _paths(_bad(dict(ok, header={"text": "a&b"}), key="t_bp3:xlsx")), "& 是列印頁首的控制字元，一律拒絕"


def test_xlsx_cells_use_the_xlsx_neutralizer_only_equals_is_prefixed():
    t = {"version": 1, "format": "xlsx", "columns": [{"id": "name", "label": "姓名"}, {"id": "phone", "label": "電話"}]}
    table, _ = ET.render_rows("t_bp3:xlsx", t, [{"name": "=HYPERLINK(1)", "phone": "+886 2"}, {"name": "-VLAN10", "phone": None}], {"role": "viewer"})
    assert table == [["姓名", "電話"], ["'=HYPERLINK(1)", "+886 2"], ["-VLAN10", None]], "xlsx 只補 '='；電話 +886 不多可見單引號"


def test_print_header_footer_land_in_the_workbook_and_the_pdf_sibling_html_and_default_html_is_unchanged(tmp_path):
    import io
    import openpyxl
    from helpers.xlsx_out import xlsx_to_html
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名冊"
    ws.append(["姓名"])
    ws.append(["甲"])
    buf = io.BytesIO()
    wb.save(buf)
    plain = xlsx_to_html(buf.getvalue(), "名冊")
    assert "class='meta'>" in plain and plain.count("class='meta'") == 1, "沒設定頁首頁尾 ⇒ 只有原本那一個 meta（匯出時間）"
    ET.apply_print_text(ws, {"header": {"text": "外包名冊"}, "footer": {"text": "內部資料"}})
    buf2 = io.BytesIO()
    wb.save(buf2)
    wb3 = openpyxl.load_workbook(io.BytesIO(buf2.getvalue()))
    assert wb3.active.oddHeader.center.text == "外包名冊" and wb3.active.oddFooter.center.text == "內部資料"
    html = xlsx_to_html(buf2.getvalue(), "名冊")
    assert "外包名冊" in html and "內部資料" in html and html.count("class='meta'") == 3
    assert wb3.active.max_row == 2, "頁首頁尾不佔儲存格，資料區位置不變"


def _xl_rows(raw):
    import io
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    return ws, [[c.value for c in row] for row in ws.iter_rows()]


CID = "subcontract:contractors"


def _seed_contractor(client, hd):
    r = client.post("/api/contractors", json={"name": "=王小明", "id_number": "A123456789", "phone": "+886 2 1234", "bank_code": "812",
                                               "bank_account_name": "王", "bank_account_number": "1234567890", "notes": "-備註"}, headers=hd)
    assert r.status_code in (200, 201), r.text


def test_roster_pilot_off_is_identical_template_picks_reorders_and_keeps_masking(client, make_user, monkeypatch):
    _, hd = _login(client, make_user, "bp3_sa10", "superadmin")
    _seed_contractor(client, hd)
    _flag(monkeypatch, False)
    base = client.get("/api/contractors/export", headers=hd)
    assert base.status_code == 200
    ws0, rows0 = _xl_rows(base.content)
    assert rows0[0][:3] == ["姓名", "證件號碼", "國籍"] and len(rows0[0]) == 15
    row = rows0[1]
    assert row[0] == "'=王小明" and row[1].startswith("******") and row[1].endswith("6789"), "證件號碼遮成末 4 碼；姓名公式中和照舊"
    assert row[4] == "+886 2 1234" and row[13] == "-備註", "xlsx 新寫入端只補 '='（電話 +886 不加）"
    _flag(monkeypatch, True)
    assert _xl_rows(client.get("/api/contractors/export", headers=hd).content)[1] == rows0, "旗標開、沒有範本 ⇒ 逐格相同"
    tpl = {"version": 1, "format": "xlsx", "columns": [{"id": "phone", "label": "電話"}, {"id": "name", "label": "姓名"}, {"id": "id_number", "label": "證件號碼"}],
           "header": {"text": "外包名冊（內部）"}, "footer": {"text": "請勿外流"}}
    pub = _publish(client, hd, tpl, key=CID)
    assert pub.status_code == 200 and not pub.json().get("pending"), pub.text
    ws, rows = _xl_rows(client.get("/api/contractors/export", headers=hd).content)
    assert rows[0] == ["電話", "姓名", "證件號碼"], "只匯出選的欄位、依範本順序"
    assert rows[1] == [row[4], row[0], row[1]], "值與預設匯出同一組遮蔽（範本不能改變某人看到的值）"
    assert ws.oddHeader.center.text == "外包名冊（內部）" and ws.oddFooter.center.text == "請勿外流"
    _flag(monkeypatch, False)
    assert _xl_rows(client.get("/api/contractors/export", headers=hd).content)[1] == rows0, "旗標關回去 ⇒ 範本不再套用，逐格回到預設"


def test_roster_template_cannot_rename_headers_and_non_superadmin_preview_falls_back_to_masked_sample(client, make_user, monkeypatch):
    _, hd = _login(client, make_user, "bp3_sa11", "superadmin")
    other, _ = _login(client, make_user, "bp3_adm5", "admin", modules=["contractor_list"])
    _seed_contractor(client, hd)
    _flag(monkeypatch, True)
    url = "/api/export-templates/%s/preview" % CID
    renamed = {"version": 1, "format": "xlsx", "columns": [{"id": "name", "label": "全名"}]}
    assert client.post(url, json={"template": renamed}, headers=hd).status_code == 422
    good = {"version": 1, "format": "xlsx", "columns": [{"id": "name", "label": "姓名"}, {"id": "bank_account_number", "label": "帳號"}, {"id": "id_number", "label": "證件號碼"}]}
    mine = client.post(url, json={"template": good, "sample": "visible"}, headers=hd).json()
    assert mine["sample"] == "visible" and mine["columns"] == ["姓名", "帳號", "證件號碼"]
    assert any(r[1] == "1234567890" for r in mine["rows"]), "最高管理者看得到全碼帳號"
    assert all(r[2].startswith("******") for r in mine["rows"]), "證件號碼一律遮罩"
    theirs = client.post(url, json={"template": good, "sample": "visible", "user": other}, headers=hd).json()
    assert theirs["sample"] == "sample" and theirs["warnings"] and not any("1234567890" in str(r) for r in theirs["rows"]), "匯出只開放最高管理者 ⇒ 別人沒有真實列，退回樣本"


# ═══ R4：PDF 頁尾／浮水印文字（版面固定的匯出＝decorOnly；先導＝開票申請憑據 PDF）═══════════════════
PDFID = "arap:invoice-voucher-pdf"


def test_pdf_decor_only_validator():
    import modules.arap.api.invoice_vouchers  # noqa: F401  註冊匯出目錄
    ok = {"version": 1, "format": "pdf", "footer": {"text": "內部資料"}, "watermark": {"text": "內部使用"}}
    assert _bad(ok, key=PDFID) == []
    assert "" in _paths(_bad({"version": 1, "format": "pdf"}, key=PDFID)), "至少一項"
    assert "columns" in _paths(_bad(dict(ok, columns=[{"id": "a", "label": "b"}]), key=PDFID)), "版面固定，沒有欄位可選"
    assert "header" in _paths(_bad(dict(ok, header={"text": "H"}), key=PDFID)), "這個匯出只開頁尾與浮水印（頁首會壓到憑據抬頭）"
    for w in ("正本", "已核准", "已 付 款", "作廢", "已簽回"):
        assert "watermark.text" in _paths(_bad(dict(ok, watermark={"text": w}), key=PDFID)), w
    assert "watermark.text" in _paths(_bad(dict(ok, watermark={"text": "x" * 41}), key=PDFID))
    assert _bad(dict(ok, footer={"text": "<script>alert(1)</script>"}), key=PDFID), "尖括號一律拒絕"


def test_apply_pdf_text_escapes_places_the_watermark_below_and_is_identity_without_text():
    html = "<html><body class='a'><div id='root'>內容</div></body></html>"
    assert ET.apply_pdf_text(html, {"version": 1, "format": "pdf"}) is html
    out = ET.apply_pdf_text(html, {"footer": {"text": "a<b>&"}, "watermark": {"text": "內部使用"}})
    assert out.index("<body") < out.index("mx-wm") < out.index("id='root'"), "插在 body 開頭、內容之前（浮水印在內容下層：z-index 0；頁尾 1）"
    assert "a&lt;b&gt;&amp;" in out and "a<b>" not in out, "文字一律跳脫"
    assert "rgba(0,0,0,.08)" in out and "z-index:0" in out
    assert out.replace(ET._PDF_CSS, "").count("<div") == 3 and "內容" in out, "原內容原樣保留"
    assert ET.apply_pdf_text("<p>no body tag</p>", {"footer": {"text": "F"}}).endswith("<p>no body tag</p>")


def test_decorate_pdf_returns_the_same_object_when_off_or_without_a_published_template(client, make_user, monkeypatch):
    import modules.arap.api.invoice_vouchers  # noqa: F401
    h = "<html><body>x</body></html>"
    _flag(monkeypatch, False)
    assert ET.decorate_pdf(PDFID, h) is h
    _flag(monkeypatch, True)
    assert ET.decorate_pdf(PDFID, h) is h, "旗標開但沒有已發布範本 ⇒ 仍是同一個字串"
    assert ET.decorate_pdf("no:such", h) is h
    _, hd = _login(client, make_user, "bp3_sa12", "superadmin")
    assert _publish(client, hd, {"version": 1, "format": "pdf", "footer": {"text": "內部資料"}}, key=PDFID).status_code == 200
    out = ET.decorate_pdf(PDFID, h)
    assert out != h and "內部資料" in out
    _flag(monkeypatch, False)
    assert ET.decorate_pdf(PDFID, h) is h, "旗標關回去 ⇒ 不再套用"


def test_invoice_voucher_builder_decorates_only_after_the_programs_own_banner():
    """掛鉤的位置（靜態）：『尚未核准』橫幅由程式先注入，範本文字（頁尾／浮水印）最後才加，不會蓋掉它；
    行為（OFF 同一個字串、ON 才加）由 test_decorate_pdf_returns_the_same_object… 證明。"""
    import pdf_gen
    src = open(pdf_gen.__file__, encoding="utf-8").read()
    hook = 'decorate_pdf("arap:invoice-voucher-pdf", out)'
    assert src.count(hook) == 1
    assert src.index("_inject_unapproved(html") < src.index(hook)
    assert "out = html if v.get(" in src


def test_pdf_comparator_adding_visible_text_is_a_loosening():
    from routers.def_pipeline import classify
    base = {"version": 1, "format": "pdf", "footer": {"text": "內部資料"}}
    assert classify("export_template", base, dict(base, footer={"text": "改過的文字"}))[0] == "neutral"
    c, why = classify("export_template", base, dict(base, watermark={"text": "內部使用"}))
    assert c == "loosen" and any("浮水印" in x for x in why)
    assert classify("export_template", base, {"version": 1, "format": "pdf"})[0] == "tighten"
    assert classify("export_template", {"version": 1, "format": "pdf", "footer": {"text": "a"}}, {"version": 1, "format": "pdf", "footer": {"text": "a"}, "header": {"text": "H"}})[0] == "loosen"


# ═══ SheetJS 三頁（客戶／料號／供應商）：先過伺服器稽核與冷卻 ═════════════════════════════════════
def test_client_export_log_audits_rate_limits_and_validates(client, make_user):
    import db
    _, hd = _login(client, make_user, "bp3_sa14", "superadmin")
    url = "/api/exports/client-log"
    assert client.post(url, json={"name": "customers", "rows": 3}).status_code in (401, 403)
    assert client.post(url, json={"name": "nope", "rows": 3}, headers=hd).status_code == 422
    for bad in (-1, 200001, "3", True, None, 1.5):
        assert client.post(url, json={"name": "customers", "rows": bad}, headers=hd).status_code == 422, bad
    r = client.post(url, json={"name": "customers", "rows": 12}, headers=hd)
    assert r.status_code == 200 and r.json() == {"ok": True}
    c = db.get_db()
    try:
        row = c.execute("SELECT detail FROM audit_log WHERE action='export.xlsx' AND target_id='customers' ORDER BY id DESC LIMIT 1").fetchone()
    finally:
        c.close()
    assert row is not None and '"rows": 12' in row[0] and '"client": true' in row[0] and "crm" in row[0]
    assert client.post(url, json={"name": "customers", "rows": 12}, headers=hd).status_code == 429, "冷卻中 ⇒ 429（前端不產檔）"


def test_the_three_sheetjs_pages_call_the_guard_before_writing_the_file():
    from pathlib import Path
    pages = Path(__file__).resolve().parents[3] / "frontend" / "pages"
    for page, name in (("customers", "customers"), ("parts", "parts"), ("suppliers", "suppliers")):
        src = (pages / (page + ".html")).read_text(encoding="utf-8")
        assert '<script src="../static/export-guard.js"></script>' in src, page
        g = "await MotrixExport.guard('%s', rows.length)" % name
        assert src.count(g) == 1, page
        assert src.index("async exportExcel()") < src.index(g) < src.index("XLSX.writeFile"), "%s：守門要在產檔之前" % page

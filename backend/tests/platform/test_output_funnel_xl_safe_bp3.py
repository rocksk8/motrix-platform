# -*- coding: utf-8 -*-
"""試算表公式注入守門（BP-3 小切片；FRAMEWORK §0B.3 BP-3「集中輸出＋公式中和」）。

規則：任何把『非字面值』寫進 xlsx／csv 儲存格的函式（`ws.append(...)`、`ws.cell(..., value=…)`、`x.value = …`、`ws["A1"] = …`、
`writer.writerow(...)`），整個最外層函式必須經過 `helpers.xlsx_out.xl_safe`（或 `set_row`）。openpyxl 會把 `=` 開頭的字串存成真公式，
CSV 在 Excel 開啟時 `= + - @ \\t \\r` 開頭都會被當公式 ⇒ 自由文字一律要中和。**只處理 str：真正的數字（含負數）不加前綴。**

粒度刻意粗（與 unscoped_reads 棘輪同）：管的是『有沒有人想過中和』，不是證明每一格都包了。
正對照：(1) 合成原始碼必須被抓到；(2) 把真實檔案裡的 `xl_safe(` 拿掉，掃描器必須抓到（證明它看得懂真實程式碼，不只看合成的）。
前端兩支 CSV 匯出（ledger-reports.js／ledger-statements.js）用同一組向量以 node 互驗。
"""
import ast
import re
import shutil
import subprocess
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[2]
FRONTEND_JS = BACKEND.parent / "frontend" / "js"

SAFE_NAMES = frozenset({"xl_safe", "set_row", "neutralize_cell"})
_WS_NAME = re.compile(r"^(ws|sheet|wsheet)\w*$|^(ws|sh)_\w+$")
_NUMERIC_CALLS = frozenset({"len", "sum", "round", "int", "float", "abs", "min", "max", "bool", "count"})
_FILE_HINT = re.compile(r"openpyxl|csv\.writer|_csv\.writer")
SKIP_FILES = ("helpers/xlsx_out.py",)

#: 已知『只寫程式自己產生的文字／數字，不含使用者自由文字』的函式。id = "相對路徑::最外層函式限定名"；原因必填。只准縮小。
KNOWN_PROGRAM_TEXT_ONLY: dict = {}


def _root_name(node):
    while isinstance(node, (ast.Attribute, ast.Subscript, ast.Call)):
        node = node.value if not isinstance(node, ast.Call) else node.func
    return node.id if isinstance(node, ast.Name) else None


def _is_literal_or_numeric(e) -> bool:
    if isinstance(e, ast.Constant):
        return True
    if isinstance(e, ast.JoinedStr):                                   # f"出納執行紀錄 — {x}"：開頭是字面文字 ⇒ 不可能以觸發字元開頭
        return bool(e.values) and isinstance(e.values[0], ast.Constant) and bool(e.values[0].value) and e.values[0].value[0] not in "=+-@\t\r"
    if isinstance(e, ast.BinOp) and isinstance(e.op, (ast.Sub, ast.Mult, ast.Div, ast.FloorDiv, ast.Mod, ast.Pow)):
        return not (isinstance(e.op, ast.Mod) and isinstance(e.left, ast.Constant) and isinstance(e.left.value, str))
    if isinstance(e, ast.UnaryOp) and isinstance(e.op, (ast.USub, ast.UAdd, ast.Not)):
        return _is_literal_or_numeric(e.operand)
    if isinstance(e, ast.Call):
        f = e.func
        n = f.id if isinstance(f, ast.Name) else f.attr if isinstance(f, ast.Attribute) else None
        return n in _NUMERIC_CALLS or n in SAFE_NAMES
    if isinstance(e, (ast.Compare,)):
        return True
    return False


def _value_exprs(arg):
    if isinstance(arg, (ast.List, ast.Tuple)):
        out = []
        for el in arg.elts:
            out.extend(_value_exprs(el.value if isinstance(el, ast.Starred) else el))
        return out
    return [arg]


def _sink_values(node):
    """⇒ 該節點若是寫儲存格／csv 的敘述，回傳被寫入的值運算式清單。"""
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
        a = node.func.attr
        if a in ("writerow", "writerows") and node.args:
            return _value_exprs(node.args[0])
        if a == "append" and len(node.args) == 1 and _WS_NAME.match(_root_name(node.func.value) or ""):
            return _value_exprs(node.args[0])
        if a == "cell":
            for kw in node.keywords:
                if kw.arg == "value":
                    return [kw.value]
            if len(node.args) >= 3:
                return [node.args[2]]
    if isinstance(node, ast.Assign) and len(node.targets) == 1:
        t = node.targets[0]
        if isinstance(t, ast.Attribute) and t.attr == "value" and (_root_name(t.value) or "").startswith(("ws", "sheet", "c", "msg", "hc", "cell")):
            return [node.value]
        if isinstance(t, ast.Subscript) and _WS_NAME.match(_root_name(t.value) or ""):
            return [node.value]
    return None


def scan_source(rel: str, source: str) -> set:
    """⇒ {id}：最外層函式含非字面值的寫入、卻沒有任何 xl_safe／set_row 參照。"""
    if not _FILE_HINT.search(source):
        return set()
    try:
        tree = ast.parse(source)
    except SyntaxError:
        return set()
    units = []
    mod_nodes = []
    for n in tree.body:
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)):
            units.append((n.name, n))
        elif isinstance(n, ast.ClassDef):
            for m in n.body:
                if isinstance(m, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    units.append(("%s.%s" % (n.name, m.name), m))
        else:
            mod_nodes.append(n)
    wrapper = ast.Module(body=mod_nodes, type_ignores=[])
    units.append(("<module>", wrapper))
    bad = set()
    for qual, unit in units:
        dirty = False
        safe_ref = False
        for node in ast.walk(unit):
            if isinstance(node, ast.Name) and node.id in SAFE_NAMES:
                safe_ref = True
            elif isinstance(node, ast.Attribute) and node.attr in SAFE_NAMES:
                safe_ref = True
            vals = _sink_values(node)
            if vals and any(not _is_literal_or_numeric(v) for v in vals):
                dirty = True
        if dirty and not safe_ref:
            bad.add("%s::%s" % (rel, qual))
    return bad


def _product_files():
    for p in sorted(BACKEND.rglob("*.py")):
        rel = p.relative_to(BACKEND).as_posix()
        parts = rel.split("/")
        if "tests" in parts or rel.startswith(("tools/", "conftest")) or "/tests/" in rel or rel in SKIP_FILES or ".venv" in rel or "site-packages" in rel:
            continue
        yield rel, p


def scan_all() -> set:
    out = set()
    for rel, p in _product_files():
        try:
            src = p.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            continue
        out |= scan_source(rel, src)
    return out


# ── 真實樹 ───────────────────────────────────────────────────────────────
def test_every_spreadsheet_writer_goes_through_xl_safe():
    cur = scan_all()
    new = sorted(cur - set(KNOWN_PROGRAM_TEXT_ONLY))
    assert not new, ("這些函式把非字面值寫進 xlsx／csv 卻沒有經過 helpers.xlsx_out.xl_safe／set_row（公式注入）：%s\n"
                     "⇒ 自由文字欄位用 xl_safe(v) 包起來；真的只寫程式自己產生的文字才列 KNOWN_PROGRAM_TEXT_ONLY 並寫原因。" % new)


def test_known_program_text_only_entries_are_justified_and_still_exist():
    cur = scan_all()
    stale = sorted(set(KNOWN_PROGRAM_TEXT_ONLY) - cur)
    assert not stale, "清單裡的函式已經不需要豁免（已中和或已不存在），請移除：%s" % stale
    vague = [k for k, why in KNOWN_PROGRAM_TEXT_ONLY.items() if len((why or "").strip()) < 12]
    assert not vague, "豁免原因太空泛：%s" % vague


def test_the_scanner_actually_sees_the_real_export_files():
    """守門的對象沒被搬走：這幾個檔案必須仍含寫入敘述（否則掃描器是在掃空氣）。"""
    for rel in ("modules/case/api/quotations.py", "modules/subcontract/api/contractors.py", "modules/netplan/export.py",
                "modules/daily_tasks/api.py", "modules/analytics/api/reports.py"):
        src = (BACKEND / rel).read_text(encoding="utf-8")
        tree = ast.parse(src)
        n = sum(1 for node in ast.walk(tree) if _sink_values(node))
        assert n >= 1, "%s 掃不到任何寫入敘述（檔案改了結構？）" % rel


# ── 正對照：合成原始碼 ────────────────────────────────────────────────────
_UNSAFE = """
import openpyxl
def export(rows):
    wb = openpyxl.Workbook(); ws = wb.active
    for r in rows:
        ws.append([r["customer"], r["total"]])
"""
_SAFE = _UNSAFE.replace('r["customer"]', 'xl_safe(r["customer"])').replace("import openpyxl", "import openpyxl\nfrom helpers.xlsx_out import xl_safe")


@pytest.mark.parametrize("src,expect_bad", [
    (_UNSAFE, True),
    (_SAFE, False),
    ("import openpyxl\ndef f(ws, name):\n    ws.cell(row=1, column=1, value=name)\n", True),
    ("import openpyxl\ndef f(ws, name):\n    ws.cell(row=1, column=1, value=xl_safe(name))\n", False),
    ("import openpyxl\ndef f(ws):\n    ws.cell(row=1, column=1, value='標題')\n    ws.append(['a', 1, 2.5])\n", False),     # 純字面值
    ("import csv\ndef f(buf, title):\n    w = csv.writer(buf)\n    w.writerow([title, 1])\n", True),
    ("import csv\ndef f(buf, title):\n    w = csv.writer(buf)\n    w.writerow([xl_safe(title), 1])\n", False),
    ("import openpyxl\ndef f(ws, name):\n    c = ws['A1']\n    c.value = name\n", True),
    ("import openpyxl\ndef f(ws, y):\n    c = ws['A1']\n    c.value = f'報表 {y}'\n", False),                                    # 開頭是字面文字
    ("import openpyxl\ndef f(ws, y):\n    ws['A1'] = y\n", True),
    ("import openpyxl\ndef outer(ws, rows):\n    def inner(r):\n        ws.append([r])\n    for r in rows:\n        inner(xl_safe(r))\n", False),   # 外層有中和 ⇒ 內層算數
    ("import openpyxl\nclass K:\n    def m(self, ws, v):\n        ws.append([v])\n", True),
    ("def f(lst, v):\n    lst.append(v)\n", False),                                                                           # 不是試算表：沒有 openpyxl／csv 線索
    ("import openpyxl\ndef f(items, v):\n    items.append(v)\n", False),                                                      # list.append 不是儲存格
])
def test_scanner_positive_and_negative_controls(src, expect_bad):
    assert bool(scan_source("synthetic.py", src)) is expect_bad


# ── 正對照：真實檔案拿掉 xl_safe 必須被抓到 ───────────────────────────────
@pytest.mark.parametrize("rel", [
    "modules/case/api/quotations.py",
    "modules/subcontract/api/contractors.py",
    "modules/netplan/export.py",
    "modules/daily_tasks/api.py",
])
def test_deliberate_break_removing_xl_safe_from_a_real_file_is_caught(rel):
    src = (BACKEND / rel).read_text(encoding="utf-8")
    assert "xl_safe" in src, "%s 沒有用 xl_safe（這個切片應該已經接上）" % rel
    assert not scan_source(rel, src) or all(k in KNOWN_PROGRAM_TEXT_ONLY for k in scan_source(rel, src))
    broken = re.sub(r"\bxl_safe\b", "str", src)             # 蓄意破壞：中和函式換成恆等轉換
    assert scan_source(rel, broken), "拿掉 xl_safe 後掃描器沒有抓到 ⇒ 守門失效：%s" % rel


# ── xl_safe 行為向量 ─────────────────────────────────────────────────────
def test_xl_safe_vectors_strings_prefixed_numbers_untouched():
    from helpers.xlsx_out import XL_FORMULA_TRIGGERS, xl_safe
    assert set(XL_FORMULA_TRIGGERS) == {"=", "+", "-", "@", "\t", "\r"}
    for bad in ("=HYPERLINK(\"http://x\",\"y\")", "+1+1", "-2+3", "@SUM(1)", "\t=1", "\r=1", "=", "-"):
        assert xl_safe(bad) == "'" + bad, bad
    for ok in ("", "abc", "客戶 A", " =1", "a=b", "'=1", "1+1"):
        assert xl_safe(ok) == ok, ok
    for num in (0, 5, -5, 3.14, -0.5, True, None):
        assert xl_safe(num) is num or xl_safe(num) == num, num
    assert isinstance(xl_safe(-5), int) and xl_safe(-5) == -5          # 真正的負數保持數字


def test_openpyxl_round_trip_free_text_is_not_a_formula(tmp_path):
    import openpyxl
    from helpers.xlsx_out import xl_safe
    wb = openpyxl.Workbook()
    ws = wb.active
    vals = ['=HYPERLINK("http://evil","x")', "+cmd|' /C calc'!A0", "-1+1", "@SUM(1,1)", "\t=1", "\r=1", "正常文字", -5, 12.5]
    ws.append([xl_safe(v) for v in vals])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    ws2 = openpyxl.load_workbook(p).active
    row = [c for c in ws2[1]]
    assert all(c.data_type != "f" for c in row)
    assert row[7].value == -5 and row[7].data_type == "n"             # 數字仍是數字
    assert row[0].value.startswith("'=")
    # 對照：沒中和的 '=' 字串會被存成真公式（證明上面的斷言有意義）
    wb3 = openpyxl.Workbook()
    wb3.active.append(['=1+1'])
    p3 = tmp_path / "raw.xlsx"
    wb3.save(p3)
    assert openpyxl.load_workbook(p3).active["A1"].data_type == "f"


# ── 前端 CSV：同一組向量以 node 互驗 ──────────────────────────────────────
_NODE = shutil.which("node")


@pytest.mark.parametrize("fn", ["ledger-reports.js", "ledger-statements.js"])
def test_frontend_csv_exports_use_the_safe_cell(fn):
    src = (FRONTEND_JS / fn).read_text(encoding="utf-8")
    assert "function csvSafeCell" in src, fn
    assert "csvSafeCell(c)" in src, "%s 的 CSV 組字沒有走 csvSafeCell" % fn
    assert "c === null || c === undefined ? '' : c" not in src, "%s 還留著未中和的舊組字" % fn


@pytest.mark.skipif(not _NODE, reason="沒有 node")
@pytest.mark.parametrize("fn", ["ledger-reports.js", "ledger-statements.js"])
def test_frontend_csv_safe_cell_vectors_match_backend(fn, tmp_path):
    src = (FRONTEND_JS / fn).read_text(encoding="utf-8")
    m = re.search(r"function csvSafeCell\(v\) \{.*?\n\}\n", src, re.S)
    assert m, fn
    js = m.group(0) + """
const bad = ['=HYPERLINK("x")', '+1+1', '-2+3', '@SUM(1)', '\\t=1', '\\r=1'];
const ok = ['', 'abc', '客戶', ' =1', 'a=b'];
const out = {bad: bad.map(csvSafeCell), ok: ok.map(csvSafeCell), nums: [0, 5, -5, 3.5].map(csvSafeCell), nul: [csvSafeCell(null), csvSafeCell(undefined)]};
console.log(JSON.stringify(out));
"""
    p = tmp_path / "v.js"
    p.write_text(js, encoding="utf-8")
    r = subprocess.run([_NODE, str(p)], capture_output=True, text=True, timeout=30, encoding="utf-8")
    assert r.returncode == 0, r.stderr
    import json
    o = json.loads(r.stdout)
    assert o["bad"] == ["'" + x for x in ['=HYPERLINK("x")', "+1+1", "-2+3", "@SUM(1)", "\t=1", "\r=1"]]
    assert o["ok"] == ["", "abc", "客戶", " =1", "a=b"]
    assert o["nums"] == [0, 5, -5, 3.5]                                # 數字（含負數）原樣
    assert o["nul"] == ["", ""]

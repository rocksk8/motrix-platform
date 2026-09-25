"""外包人員名冊 CRUD — superadmin only."""
import base64
import io
import os
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Body, File, HTTPException, Header, UploadFile
from fastapi.responses import Response
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel

from db import get_db
from helpers import _require_user, _tok, _audit, notify_module_activity

# 字體路徑（Windows 微軟正黑體，找不到退回預設）
_FONT_PATH = r"C:\Windows\Fonts\msjhbd.ttc"
if not os.path.exists(_FONT_PATH):
    _FONT_PATH = r"C:\Windows\Fonts\msjh.ttc"


def _stamp_id_card(data_uri: str) -> str:
    """在身分證影本中央加浮水印後回傳 data URI（JPEG）。"""
    if not data_uri or not data_uri.startswith("data:image/"):
        return data_uri
    _, b64 = data_uri.split(",", 1)
    img = Image.open(io.BytesIO(base64.b64decode(b64))).convert("RGBA")
    w, h = img.size
    MAX_W = 1800
    if w > MAX_W:
        ratio = MAX_W / w
        img = img.resize((MAX_W, int(h * ratio)), Image.LANCZOS)
        w, h = img.size

    overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))

    # ── 中央浮水印：在與原圖等大的透明層畫居中文字，整層旋轉後 alpha_composite ──
    wm_text = "僅供申報扣繳使用"
    wm_size = max(20, int(w * 0.60) // max(len(wm_text), 1))
    try:
        wm_font = ImageFont.truetype(_FONT_PATH, wm_size)
    except Exception:
        wm_font = ImageFont.load_default()

    # 量測文字實際尺寸
    _probe = ImageDraw.Draw(overlay)
    bbox   = _probe.textbbox((0, 0), wm_text, font=wm_font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]

    # 在與原圖同尺寸的透明層畫居中文字（anchor="mm" 保證正中央），整層旋轉後 alpha_composite
    txt_lay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    ImageDraw.Draw(txt_lay).text(
        (w // 2, h // 2), wm_text,
        font=wm_font, fill=(180, 0, 0, 150), anchor="mm"
    )
    rotated = txt_lay.rotate(30, resample=Image.BICUBIC)
    overlay = Image.alpha_composite(overlay, rotated)

    # ── 底部紅底白字橫幅 ──────────────────────────────────────────
    draw     = ImageDraw.Draw(overlay)
    banner_h = max(28, int(h * 0.065))
    draw.rectangle([0, h - banner_h, w, h], fill=(155, 0, 0, 214))
    b_size = max(10, int(banner_h * 0.50))
    try:
        b_font = ImageFont.truetype(_FONT_PATH, b_size)
    except Exception:
        b_font = wm_font
    draw.text((w // 2, h - banner_h // 2),
              "本影本依法留存，僅供扣繳憑單申報使用，不得挪作其他用途",
              font=b_font, fill=(255, 255, 255, 255), anchor="mm")

    result = Image.alpha_composite(img, overlay).convert("RGB")
    buf = io.BytesIO()
    result.save(buf, format="JPEG", quality=88)
    encoded = base64.b64encode(buf.getvalue()).decode()
    return f"data:image/jpeg;base64,{encoded}"


def _stamp_passbook(data_uri: str) -> str:
    """在銀行存簿影本底部加用途說明橫幅後回傳 data URI（JPEG）。"""
    if not data_uri or not data_uri.startswith("data:image/"):
        return data_uri
    _, b64 = data_uri.split(",", 1)
    img = Image.open(io.BytesIO(base64.b64decode(b64))).convert("RGBA")
    w, h = img.size
    MAX_W = 1800
    if w > MAX_W:
        ratio = MAX_W / w
        img = img.resize((MAX_W, int(h * ratio)), Image.LANCZOS)
        w, h = img.size

    overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    banner_h = max(28, int(h * 0.055))
    draw.rectangle([0, h - banner_h, w, h], fill=(0, 80, 160, 210))
    b_size = max(10, int(banner_h * 0.50))
    try:
        b_font = ImageFont.truetype(_FONT_PATH, b_size)
    except Exception:
        b_font = ImageFont.load_default()
    draw.text((w // 2, h - banner_h // 2),
              "本影本依法留存，僅供勞務報酬匯款核對使用，不得挪作其他用途",
              font=b_font, fill=(255, 255, 255, 255), anchor="mm")

    result = Image.alpha_composite(img, overlay).convert("RGB")
    buf = io.BytesIO()
    result.save(buf, format="JPEG", quality=88)
    encoded = base64.b64encode(buf.getvalue()).decode()
    return f"data:image/jpeg;base64,{encoded}"


router = APIRouter()

_LIST_COLS = (
    "id, name, id_number, nationality, has_union_insurance, "
    "phone, email, address, line_id, "
    "bank_code, bank_name, bank_branch, bank_account_name, bank_account_number, "
    "notes, active, created_at, updated_at, "
    "(CASE WHEN (id_card_image != '' AND id_card_image IS NOT NULL) OR "
    "(id_card_image_back != '' AND id_card_image_back IS NOT NULL) THEN 1 ELSE 0 END) AS has_id_card, "
    "(CASE WHEN bank_passbook_image != '' AND bank_passbook_image IS NOT NULL THEN 1 ELSE 0 END) AS has_passbook"
)


class ContractorIn(BaseModel):
    name:                str
    id_number:           Optional[str] = ''
    nationality:         Optional[str] = '本國籍'
    has_union_insurance: Optional[bool] = False
    phone:               Optional[str] = ''
    email:               Optional[str] = ''
    address:             Optional[str] = ''
    line_id:             Optional[str] = ''
    bank_code:           Optional[str] = ''
    bank_name:           Optional[str] = ''
    bank_branch:         Optional[str] = ''
    bank_account_name:   Optional[str] = ''
    bank_account_number: Optional[str] = ''
    notes:               Optional[str] = ''


def _row_to_dict(row) -> dict:
    d = dict(row)
    d['has_union_insurance'] = bool(d.get('has_union_insurance', 0))
    d['active'] = bool(d.get('active', 1))
    d['has_id_card']   = bool(d.get('has_id_card', 0))
    d['has_passbook']  = bool(d.get('has_passbook', 0))
    return d


@router.get("/api/contractors/selectable")
def list_contractors_selectable(authorization: str = Header(None)):
    """輕量列表供案件管理承攬商派發的「外包名單人員」下拉使用（所有登入者皆可讀，
    比照 vendor-contractors/selectable 的慣例——不含銀行/身分證等敏感欄位）。"""
    _require_user(authorization)
    conn = get_db()
    rows = conn.execute(
        "SELECT id, name, phone FROM contractors WHERE active=1 ORDER BY name"
    ).fetchall()
    conn.close()
    return [{"id": r["id"], "name": r["name"], "phone": r["phone"] or ""} for r in rows]


@router.get("/api/contractors")
def list_contractors(q: Optional[str] = None, active_only: bool = True,
                     authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    conn = get_db()
    sql = f"SELECT {_LIST_COLS} FROM contractors"
    params = []
    clauses = []
    if active_only:
        clauses.append("active=1")
    if q:
        clauses.append("name LIKE ?")
        params.append(f"%{q}%")
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY name"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [_row_to_dict(r) for r in rows]


@router.post("/api/contractors", status_code=201)
def create_contractor(body: ContractorIn, authorization: str = Header(None)):
    user = _require_user(authorization, require_superadmin=True, module='contractor_list')
    now = datetime.now().isoformat()
    conn = get_db()
    cur = conn.execute("""
        INSERT INTO contractors
          (name, id_number, nationality, has_union_insurance,
           phone, email, address, line_id,
           bank_code, bank_name, bank_branch, bank_account_name, bank_account_number,
           notes, active, created_at, updated_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,?)
    """, (
        body.name, body.id_number, body.nationality, 1 if body.has_union_insurance else 0,
        body.phone, body.email, body.address, body.line_id,
        body.bank_code, body.bank_name, body.bank_branch,
        body.bank_account_name, body.bank_account_number,
        body.notes, now, now
    ))
    cid = cur.lastrowid
    conn.commit()
    conn.close()
    _audit(_tok(authorization), 'contractor.create', 'contractor', str(cid), body.name)
    notify_module_activity("外包名冊", "建立", user.get("display_name") or user["username"],
                            body.name, "vendor-contractors.html")
    return {"id": cid, "created_at": now}


# ── CT1（2026-09-24 使用者裁示 D4）：名冊 Excel 匯入／匯出 ─────────────────────────────
# 欄位順序即匯出順序；匯入依表頭文字對應（缺的欄＝不動那個欄位，舊檔沒有「分行」照樣可匯入）。
_XLSX_COLS = [
    ("姓名", "name"), ("證件號碼", "id_number"), ("國籍", "nationality"), ("職業工會投保", "has_union_insurance"),
    ("電話", "phone"), ("Email", "email"), ("地址", "address"), ("LINE", "line_id"),
    ("銀行代碼", "bank_code"), ("銀行名稱", "bank_name"), ("分行", "bank_branch"),
    ("戶名", "bank_account_name"), ("帳號", "bank_account_number"), ("備註", "notes"), ("狀態", "active"),
]
_MASK = "******"


def _mask_id(v):
    """證件號碼只顯示末 4 碼（使用者裁示）。"""
    v = (v or "").strip()
    return (_MASK + v[-4:]) if len(v) > 4 else (_MASK if v else "")


@router.get("/api/contractors/export")
def export_contractors(authorization: str = Header(None)):
    """匯出外包名冊（含停用的人）。證件號碼遮成末 4 碼；分行與帳號照實（出納匯款要用）。寫稽核。"""
    import openpyxl
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    conn = get_db()
    try:
        rows = conn.execute("SELECT * FROM contractors ORDER BY name").fetchall()
    finally:
        conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "外包名冊"
    ws.append([h for h, _ in _XLSX_COLS])
    for r in rows:
        out = []
        for _h, col in _XLSX_COLS:
            v = r[col]
            if col == "id_number":
                v = _mask_id(v)
            elif col == "has_union_insurance":
                v = "是" if v else "否"
            elif col == "active":
                v = "往來中" if v else "停用"
            # ⚠ 空字串寫進 openpyxl 會產生不合法的 inlineStr（Excel 開檔報修復）⇒ 空值寫 None
            out.append(v if v not in ("", None) else None)
        ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    _audit(_tok(authorization), 'contractors.export', 'contractor', '', f"匯出外包名冊（{len(rows)} 人）")
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''contractors.xlsx"})


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                      # Excel 把純數字的電話／帳號讀成數字
    return str(v).strip()


@router.post("/api/contractors/import")
async def import_contractors(file: UploadFile = File(...), authorization: str = Header(None)):
    """匯入外包名冊。比對規則（使用者裁示「以姓名＋證件號碼比對」，**對不到唯一一人就不猜**）：
    - 證件號碼完整 ⇒ 姓名＋證件號碼完全相同者更新；沒有 ⇒ 新增（證件號碼已被別的姓名使用 ⇒ 報錯）
    - 證件號碼是匯出時的遮罩（******1234）⇒ 姓名＋末 4 碼剛好一人 ⇒ 更新；否則報錯；遮罩值不寫回
    - 證件號碼空白 ⇒ 同姓名剛好一人 ⇒ 更新；沒有 ⇒ 新增；多人 ⇒ 報錯
    檔案沒有的欄位不動（舊檔沒有「分行」照樣匯入、不會清掉既有分行）。
    """
    import openpyxl
    user = _require_user(authorization, require_superadmin=True, module='contractor_list')
    try:
        wb = openpyxl.load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(400, "讀不到這個檔案，請上傳 .xlsx")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "檔案是空的")
    head = [_cell(h) for h in rows[0]]
    known = dict(_XLSX_COLS)
    idx = {known[h]: i for i, h in enumerate(head) if h in known}
    if "name" not in idx:
        raise HTTPException(400, "找不到「姓名」欄")
    now = datetime.now().isoformat()
    created = updated = 0
    errors = []
    conn = get_db()
    try:
        existing = [dict(r) for r in conn.execute("SELECT id, name, id_number FROM contractors").fetchall()]
        for n, raw in enumerate(rows[1:], start=2):
            vals = {col: _cell(raw[i]) if i < len(raw) else "" for col, i in idx.items()}
            name = vals.get("name", "")
            if not name:
                if any(vals.values()):
                    errors.append({"row": n, "message": f"第 {n} 列沒有姓名，略過"})
                continue
            idn = vals.get("id_number", "")
            masked = idn.startswith("*")
            same_name = [e for e in existing if (e["name"] or "") == name]
            if masked:
                hits = [e for e in same_name if (e["id_number"] or "").endswith(idn.lstrip("*")) and idn.lstrip("*")]
            elif idn:
                hits = [e for e in same_name if (e["id_number"] or "") == idn]
                if not hits and any((e["id_number"] or "") == idn for e in existing):
                    errors.append({"row": n, "message": f"第 {n} 列：證件號碼已登記在另一個姓名下，未匯入"})
                    continue
            else:
                hits = same_name
            if len(hits) > 1:
                errors.append({"row": n, "message": f"第 {n} 列：「{name}」對到 {len(hits)} 位，無法判斷是哪一位，未匯入"})
                continue
            fields = {}
            for col, v in vals.items():
                if col in ("name",) or (col == "id_number" and masked):
                    continue
                if col == "has_union_insurance":
                    v = 1 if v in ("是", "1", "Y", "y", "true", "True") else 0
                elif col == "active":
                    v = 0 if v == "停用" else 1
                fields[col] = v
            if hits:
                if fields:
                    sets = ", ".join(f"{c}=?" for c in fields)
                    conn.execute(f"UPDATE contractors SET {sets}, updated_at=? WHERE id=?",
                                 list(fields.values()) + [now, hits[0]["id"]])
                updated += 1
            elif masked:
                errors.append({"row": n, "message": f"第 {n} 列：找不到「{name}」（證件號碼末 4 碼 {idn.lstrip('*')}），"
                                                    "遮罩的證件號碼不能用來新增人員，未匯入"})
                continue
            else:
                cols = ["name"] + list(fields)
                cur = conn.execute(
                    f"INSERT INTO contractors ({', '.join(cols)}, created_at, updated_at) VALUES "
                    f"({', '.join('?' * len(cols))}, ?, ?)", [name] + list(fields.values()) + [now, now])
                existing.append({"id": cur.lastrowid, "name": name, "id_number": fields.get("id_number", "")})
                created += 1
        conn.commit()
    finally:
        conn.close()
    _audit(_tok(authorization), 'contractors.import', 'contractor', '',
           f"匯入外包名冊（新增 {created}、更新 {updated}、未匯入 {len(errors)}）")
    notify_module_activity("外包名冊", "匯入", user.get("display_name") or user["username"],
                           f"新增 {created}、更新 {updated}", "contractors.html")
    return {"created": created, "updated": updated, "errors": errors}


@router.get("/api/contractors/{cid}")
def get_contractor(cid: int, authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    conn = get_db()
    row = conn.execute("SELECT * FROM contractors WHERE id=?", (cid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "找不到此外包人員")
    return _row_to_dict(row)


@router.put("/api/contractors/{cid}")
def update_contractor(cid: int, body: ContractorIn, authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    now = datetime.now().isoformat()
    conn = get_db()
    res = conn.execute("""
        UPDATE contractors SET
          name=?, id_number=?, nationality=?, has_union_insurance=?,
          phone=?, email=?, address=?, line_id=?,
          bank_code=?, bank_name=?, bank_branch=?, bank_account_name=?, bank_account_number=?,
          notes=?, updated_at=?
        WHERE id=?
    """, (
        body.name, body.id_number, body.nationality, 1 if body.has_union_insurance else 0,
        body.phone, body.email, body.address, body.line_id,
        body.bank_code, body.bank_name, body.bank_branch,
        body.bank_account_name, body.bank_account_number,
        body.notes, now, cid
    ))
    conn.commit()
    conn.close()
    if res.rowcount == 0:
        raise HTTPException(404, "找不到此外包人員")
    _audit(_tok(authorization), 'contractor.update', 'contractor', str(cid), body.name)
    return {"updated_at": now}


@router.get("/api/contractors/{cid}/privacy-notice")
def get_contractor_privacy_ack(cid: int, authorization: str = Header(None)):
    """R3（個資法 §8）：這位人員的「已告知」紀錄（沒有 ⇒ ack=None）。"""
    from helpers import privacy_notice as _pn
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    return {"ack": _pn.get_ack("contractor", cid)}


@router.post("/api/contractors/{cid}/privacy-notice/ack")
def ack_contractor_privacy_notice(cid: int, authorization: str = Header(None)):
    """記錄「已告知當事人」：時間與人員由伺服器決定；已記錄的不覆蓋。"""
    from helpers import privacy_notice as _pn
    user = _require_user(authorization, require_superadmin=True, module='contractor_list')
    conn = get_db()
    row = conn.execute("SELECT name FROM contractors WHERE id=?", (cid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "找不到此外包人員")
    rec, created = _pn.record_ack("contractor", cid, user)
    if created:
        _audit(_tok(authorization), 'contractor.privacy_notice_ack', 'contractor', str(cid), row["name"],
               {"noticeHash": rec.get("noticeHash")})
    return {"ack": rec, "created": created}


@router.patch("/api/contractors/{cid}/active")
def toggle_contractor_active(cid: int, authorization: str = Header(None)):
    user = _require_user(authorization, require_superadmin=True)
    conn = get_db()
    row = conn.execute("SELECT active, name FROM contractors WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "找不到此外包人員")
    new_active = 0 if row["active"] else 1
    conn.execute("UPDATE contractors SET active=?, updated_at=? WHERE id=?",
                 (new_active, datetime.now().isoformat(), cid))
    conn.commit()
    conn.close()
    action = 'contractor.activate' if new_active else 'contractor.deactivate'
    _audit(_tok(authorization), action, 'contractor', str(cid), row["name"])
    notify_module_activity("外包名冊", "啟用" if new_active else "停用",
                            user.get("display_name") or user["username"], row["name"], "vendor-contractors.html")
    return {"active": bool(new_active)}


@router.get("/api/contractors/{cid}/id-card")
def get_id_card(cid: int, authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True, module='contractor_list')
    conn = get_db()
    row = conn.execute(
        "SELECT id_card_image, id_card_image_back, bank_passbook_image FROM contractors WHERE id=?", (cid,)
    ).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "找不到此外包人員")
    front    = row["id_card_image"] or ""
    back     = row["id_card_image_back"] or ""
    passbook = row["bank_passbook_image"] or ""
    return {"image_front": front, "image_back": back, "image_data": front, "bank_passbook": passbook}


@router.put("/api/contractors/{cid}/id-card")
def upload_id_card(cid: int, body: dict = Body(...), authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True)
    image_front = body.get("image_front", body.get("image_data", ""))
    image_back  = body.get("image_back", "")
    passbook    = body.get("bank_passbook", "")
    for img in (image_front, image_back, passbook):
        if img and not img.startswith("data:image/"):
            raise HTTPException(400, "無效的圖片格式，需為 data URI")
    # 身分證加法律浮水印；存簿加輕量用途說明浮水印
    if image_front:
        try: image_front = _stamp_id_card(image_front)
        except Exception: pass
    if image_back:
        try: image_back = _stamp_id_card(image_back)
        except Exception: pass
    if passbook:
        try: passbook = _stamp_passbook(passbook)
        except Exception: pass
    now = datetime.now().isoformat()
    conn = get_db()
    res = conn.execute(
        "UPDATE contractors SET id_card_image=?, id_card_image_back=?, bank_passbook_image=?, updated_at=? WHERE id=?",
        (image_front, image_back, passbook, now, cid)
    )
    conn.commit()
    conn.close()
    if res.rowcount == 0:
        raise HTTPException(404, "找不到此外包人員")
    parts = []
    if image_front: parts.append("身分證正面")
    if image_back:  parts.append("身分證反面")
    if passbook:    parts.append("銀行存簿")
    label = f"上傳影本（{'、'.join(parts)}）" if parts else "移除影本"
    _audit(_tok(authorization), 'contractor.id_card.update', 'contractor', str(cid), label)
    return {"ok": True, "updated_at": now}

"""T100（鼎新）傳票批次匯出（2026-09-01 新增）。

背景：使用者要求把系統既有的財務資料對接鼎新 T100，且明確選擇「先做批次匯出成
T100 標準傳票匯入格式，不做即時 API 對接」——T100 API 需要貴公司自行向鼎新申請
存取權限，這裡沒有真實憑證可以測試，貿然做 API 串接無法驗證正確性；批次匯出成
Excel、財務人員在 T100 用既有匯入功能手動核對匯入，風險小很多，且可以立刻測試。

設計採**現金基礎**（cash basis）：只匯出「錢真的有進出」的事件，不匯出應計項目：
  - 收款事件：案件款項明細已填發票號碼且已收款（沿用 reports.py::_collect_tax_invoices()
    同一份資料源，跟稅務匯出數字保證一致）→ 借 銀行存款(含稅) / 貸 銷貨收入(未稅)
    ＋ 貸 銷項稅額(稅額，若有)
  - 付款事件：承攬商匯款申請已標記已匯款（沿用 cashier.py 出納模組同一份資料源）
    → 借 承攬商費用(含稅) / 貸 銀行存款(含稅)
  - 付款事件：料件/設備進貨批次已標記已付款（2026-09-01 同輪新增，見 DB v70
    `stock_batches`）→ 借 料件設備成本 / 貸 銀行存款

刻意排除 payment_requests（請款單）——那是對客戶要款的文件，沒有「已收款」狀態，
不是真的金流事件，比照 reports.py::_compute_cash_position() 既有的排除理由（同一
份資料在系統裡任何金流类彙總都不該把它算進去，避免各處各自決定要不要排除造成
不一致）。

每筆事件產生的傳票天生借貸平衡（同一 voucherNo 底下借方合計＝貸方合計），銷項
發票／承攬商應付／料件設備應付／銀行對帳四個面向用同一份匯出涵蓋，避免各自
獨立報表、資料源不同、數字對不上。

科目代號（借貸方會計科目）由 superadmin 在 GET/PUT /api/settings/t100-export-config
設定，預設全部留白——這是刻意的，貴公司財務團隊需要先確認實際使用的科目代號
才具備直接匯入 T100 的意義；金額/日期/摘要/來源單號/交易對象等其餘欄位在科目
代號填入前就已經正確可用，財務可以先核對數字正確性。

**科目代號分維度設定（2026-09-01 同輪擴充）**：
  - **依銀行帳戶**：`bankAccounts`（設定頁維護的清單 `[{name, acctCode}]`，供標記
    已付款/已收款時選擇）——但匯出計算本身**不查這份清單**，而是直接讀「標記
    當下」寫進各筆交易自己身上的 `bankAccountName`/`bankAccountCode`（比照
    `paidBankAccountName`/`paidBankAccountCode` 快照模式，見 §各表 docstring）。
    這樣設計的理由：銀行帳戶是「這一筆錢實際走哪個戶頭」的一次性事實，跟後續
    設定頁清單怎麼改都無關，不應該被之後的設定變動追溯影響。
  - **依料件分類**：`inventoryExpenseAccounts`（`{分類名稱: 科目代號}`，鍵對應
    `helpers/part_catalog.py::PART_CATEGORIES`）——這個**是**即時查表（不快照），因為分類本身
    不會變，之後財務更正某分類的科目代號，應該連未確認的舊事件都一起套用新值，
    跟 `salesRevenueAccount` 等其餘固定欄位是同一種「即時解析」邏輯。
  - 其餘科目（銷貨收入/銷項稅額/承攬商費用/部門別/傳票別）維持全公司單一設定，
    未要求分維度。

**標記已付款/已收款時的銀行帳戶預設值（2026-09-02 新增）**：使用者要求「須帶入
當時填寫或是預設的匯款帳戶」——三個標記畫面打開時，銀行帳戶下拉不再一律空白，
依序嘗試：①查「這個對象（承攬商/供應商/客戶）上一次標記時用的帳戶」（見
`contractor_vouchers.py::get_last_paid_bank_account()`／
`inventory.py::get_last_paid_bank_account()`／
`quotations.py::get_last_received_bank_account()` 三支各自查詢自己資料表最近
一筆已標記記錄）②查無上次紀錄則退回 `defaultBankAccountCode`（系統預設）
③兩者都沒有才維持空白。三支查詢都只讀不寫，前端仍可手動改選，這只是省下
「大多數情況根本不用選」的那次點擊，不是強制值。

**已匯入確認追蹤（2026-09-01 同輪新增，DB v69 `t100_export_confirmations`）**：
匯出 Excel 本身不代表財務真的把這批傳票匯入了 T100（匯出後可能發現資料有誤、
或財務決定分批匯入）——匯出跟「標記已匯入」是兩個獨立動作，只有明確標記過的
事件才會在之後的匯出/預覽自動排除，避免同一筆事件被財務重複匯入 T100 造成
金額灌水。流程：財務 `GET preview` 預覽本期未確認事件 → 實際到 T100 匯入 →
回來 `POST confirm` 標記整批已匯入 → 該批事件之後永久不再出現在任何日期區間
的匯出/預覽中（除非用 `POST unconfirm` 撤銷）。
"""
import io
from datetime import date, datetime
from typing import Dict, List
from urllib.parse import quote as _url_quote

import openpyxl
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Header, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from db import get_db
from core import registry as _registry
from helpers import _require_user, _tok, _audit, _get_setting, _set_setting
#: 收款事件（銷項）的資料屬 M05 應收應付（ROADMAP A8b 已收回模組）：經 provider 取用；M05 不在 ⇒ 沒有收款事件並明說
T100_RECEIVABLES_MISSING = "應收應付模組未安裝：T100 匯出不含收款事件（銷項）"


def _collect_tax_invoices(year=None, month=None):
    p = _registry.single_provider("receivables.tax_invoices")
    return [] if p is None else p(year, month)
from helpers.xlsx_out import check_export_rate, set_row, xl_style
from helpers.company_identity import company_heading
from helpers.part_catalog import PART_CATEGORIES
# X-VAT（2026-09-26）：金額一律四捨五入（內建 round() 是銀行家捨入：.5 取偶數）
from helpers.legal_params import round_half_up

router = APIRouter()

_DEFAULT_T100_CONFIG = {
    "bankAccounts":              [],  # [{"name": str, "acctCode": str}, ...]，設定頁維護，供標記已付款/收款時選擇（見下方 T100BankAccount）
    "defaultBankAccountCode":    "",  # 2026-09-02 新增：系統預設銀行帳戶（acctCode），標記已付款/收款時
                                       # 找不到「這個對象上次用哪個帳戶」才會退回用這個，見 §對象別最近一次使用
    "salesRevenueAccount":       "",  # 銷貨收入科目代號
    "outputTaxAccount":          "",  # 銷項稅額科目代號
    "contractorExpenseAccount":  "",  # 承攬商費用科目代號
    "inventoryExpenseAccounts":  {},  # {料件分類: 科目代號}，鍵對應 helpers/part_catalog.py::PART_CATEGORIES（2026-09-01 同輪新增）
    "departmentCode":            "",  # 部門別代號（選填，留空則傳票不分部門）
    "voucherCategory":           "轉", # 傳票別（T100 常見：現／轉／記，預設「轉」）
}


# ── 科目代號設定（superadmin 維護，admin+ 可查閱） ─────────────────────────────

class T100BankAccount(BaseModel):
    name: str = ""
    acctCode: str = ""


class T100ExportConfigBody(BaseModel):
    bankAccounts: List[T100BankAccount] = []
    defaultBankAccountCode: str = ""
    salesRevenueAccount: str = ""
    outputTaxAccount: str = ""
    contractorExpenseAccount: str = ""
    inventoryExpenseAccounts: Dict[str, str] = {}
    departmentCode: str = ""
    voucherCategory: str = "轉"


def _t100_config() -> dict:
    cfg = {**_DEFAULT_T100_CONFIG, **(_get_setting("t100_export_config", {}) or {})}
    # 確保目前所有料件分類（helpers/part_catalog.py::PART_CATEGORIES）都有一個鍵可填，即使
    # 使用者還沒存過任何值；分類名稱之後若新增，重新 GET 一次就會自動補上
    # 空白鍵，不需要額外 migration 或手動同步。
    filled = dict(cfg.get("inventoryExpenseAccounts") or {})
    for c in PART_CATEGORIES:
        filled.setdefault(c["name"], "")
    cfg["inventoryExpenseAccounts"] = filled
    return cfg


def validate_account_code(conn, code):
    """T100 設定裡的科目代號必須指得到 `account_items` 的一列。回 `(ok, err)`。

    ## ☠️ 打錯的代號**不會報錯**

    T100 匯出照樣產生，到**會計師匯入那一刻**才發現 ——
    而那時傳票已經開出去了。
    ⇒ 擋在**儲存設定**那一刻，不是匯出那一刻。

    ## ⚠️ 不擋「非 statutory」，而要擋「已停用」

    ```
    法定      可選
    自訂      **也可選**   <= 使用者自訂的科目也可能是正確的對應
    已停用    **不可選**   <= 停用的科目不該被新設定引用
    ```
    🔑 兩種拒絕的**訊息要分得出來**：「找不到」與「已停用」的下一步不同 ——
       前者是打錯字，後者是那個科目還在、只是不該再用。

    ## 📌 而已經設定好的**不因停用而失效**

    這一支只擋**新的設定值**。已存的值若指向一個被停用的科目，
    正確處置是**明著顯示「這個科目已停用」並要求重選** ——
    ☠️ 靜默失效的症狀是「匯出的科目代號突然變空」，**而沒有人會知道為什麼**。
    """
    value = (code or "").strip()
    if not value:
        # 空字串交給「完整性」那一關處理，不是這一支的事。
        return True, ""
    row = conn.execute(
        "SELECT code, is_active FROM account_items WHERE code = ?",
        (value,)).fetchone()
    if row is None:
        return False, ("科目代號「%s」不存在於會計科目表，請確認後重新輸入。" % value)
    if not row["is_active"]:
        return False, ("科目代號「%s」已停用，請改選一個仍在使用中的科目。" % value)
    return True, ""


def config_code_issues(conn, cfg):
    """已存設定裡**指不到東西**的科目代號。回 `[{where, code, reason}]`。

    ## 🔴 `validate_account_code()` 只擋**新值**，而舊值會安靜地失效

    ```
    設定存好了 => 那個科目後來被停用（或被改掉代號）
    => 匯出時 T100 那一欄用的是一個**已停用的科目**
    => ☠️ 沒有錯誤訊息；症狀是會計師匯入時才退件
    ```
    🔑 〈降級之後它還是會動〉：**壞掉會被報修，而「還能跑但不對」不會。**
    ⇒ 所以要在設定畫面上**明著顯示並要求重選**，不是靜默。

    ## ⚠️ 只回報，不修改

    這一支**不清空**任何值 ——
    ☠️ 自動清空的話，使用者下次打開會看到一個空欄位，
       而他不知道那裡**本來有值、是誰清掉的**。
    """
    issues = []

    def _check(where, code):
        ok, err = validate_account_code(conn, code)
        if not ok:
            issues.append({"where": where, "code": (code or "").strip(), "reason": err})

    _check("銷貨收入", cfg.get("salesRevenueAccount"))
    _check("銷項稅額", cfg.get("outputTaxAccount"))
    _check("承攬商費用", cfg.get("contractorExpenseAccount"))
    for name, code in (cfg.get("inventoryExpenseAccounts") or {}).items():
        _check("料件分類：%s" % name, code)
    for b in (cfg.get("bankAccounts") or []):
        acct = b.get("acctCode") if isinstance(b, dict) else getattr(b, "acctCode", "")
        _check("銀行帳戶：%s" % ((b.get("name") if isinstance(b, dict) else "") or "未命名"),
               acct)
    return issues


@router.get("/api/settings/t100-export-config")
def get_t100_export_config(authorization: str = Header(None)):
    u = _require_user(authorization)
    if u["role"] not in ("superadmin", "admin"):
        raise HTTPException(403, "僅管理員以上可查閱")
    cfg = _t100_config()
    # 🔑 把「已存的值現在還指不指得到東西」一起回去，讓畫面說得出來。
    #    ⚠️ 這個鍵**不進 `T100ExportConfigBody`** ⇒ 前端整包 PUT 回來時
    #       pydantic 會把它丟掉，不會被存進設定。
    conn = get_db()
    try:
        cfg["codeIssues"] = config_code_issues(conn, cfg)
    finally:
        conn.close()
    return cfg


@router.put("/api/settings/t100-export-config")
def set_t100_export_config(body: T100ExportConfigBody, authorization: str = Header(None)):
    _require_user(authorization, require_superadmin=True)
    _set_setting("t100_export_config", body.model_dump())
    _audit(_tok(authorization), "settings.t100_export_config.update", "settings",
           "t100_export_config", "T100 傳票匯出科目代號設定")
    return {"ok": True}


# ── 傳票資料組裝 ────────────────────────────────────────────────────────────────

#: IP-14 對方不在時（M04 外包工班）：T100 預覽的說明
T100_CONTRACTOR_MISSING = "外包工班模組未安裝：本次匯出不含承攬商費用的付款傳票"


def _collect_paid_contractor_vouchers(start: str, end: str) -> list:
    pub = _registry.single_provider("contractor_voucher.public")       # IP-14（M04）
    if pub is None:
        return []                                                     # M04 不在 ⇒ 沒有承攬付款傳票可匯
    conn = get_db()
    try:
        rows = conn.execute("""
            SELECT * FROM contractor_payment_vouchers
            WHERE is_paid=1 AND paid_at BETWEEN ? AND ?
            ORDER BY paid_at
        """, (start + "T00:00:00", end + "T23:59:59")).fetchall()
        return [pub(r, include_snapshot=False) for r in rows]
    finally:
        conn.close()


def _voucher_line(d, category, summary, acct_code, acct_name, debit, credit, dept, source_no, counterparty):
    return {
        "date": d, "category": category, "summary": summary,
        "acctCode": acct_code, "acctName": acct_name,
        "debit": round_half_up(debit) if debit else 0, "credit": round_half_up(credit) if credit else 0,
        "dept": dept, "sourceNo": source_no, "counterparty": counterparty,
    }


def _collect_paid_stock_batches(start: str, end: str) -> list:
    """料件/設備進貨已付款批次（2026-09-01 同輪新增，見 db.py::_m070_stock_batches()）。
    qty/total_cost 即時從 stock_items 群組加總（不信任任何快取值），比照
    routers/inventory.py::list_batches() 同一套「即時算，不信任快取」原則。
    額外 JOIN parts 取得料件分類，供 inventoryExpenseAccounts 依分類查科目代號。"""
    conn = get_db()
    try:
        rows = conn.execute("""
            SELECT sb.batch_no, sb.part_no, sb.supplier_name, sb.invoice_no, sb.paid_at,
                   sb.paid_bank_account_name, sb.paid_bank_account_code,
                   COALESCE(p.category, '') AS category,
                   SUM(si.cost) AS total_cost
            FROM stock_batches sb
            JOIN stock_items si ON si.batch_no = sb.batch_no
            LEFT JOIN parts p ON p.part_no = sb.part_no
            WHERE sb.is_paid=1 AND sb.paid_at BETWEEN ? AND ?
            GROUP BY sb.batch_no
            ORDER BY sb.paid_at
        """, (start, end)).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def _confirmed_keys(conn) -> set:
    rows = conn.execute("SELECT source_type, source_key FROM t100_export_confirmations").fetchall()
    return {(r["source_type"], r["source_key"]) for r in rows}


def _collect_t100_events(start: str, end: str, exclude_confirmed: bool = True) -> list:
    """回傳**事件層級**清單（一個事件＝一張傳票，含其借貸分錄 lines），供
    Excel 攤平輸出、預覽 JSON、標記已匯入共用同一份組裝邏輯，避免三處各自
    重寫一遍篩選條件而彼此不一致。"""
    conn = get_db()
    try:
        confirmed = _confirmed_keys(conn) if exclude_confirmed else set()
    finally:
        conn.close()

    cfg = _t100_config()
    events = []

    # 收款事件（銷項）：借 銀行存款(含稅) / 貸 銷貨收入(未稅) ＋ 貸 銷項稅額(稅額)
    for inv in _collect_tax_invoices():
        d = (inv.get("date") or "")[:10]
        if not d or not (start <= d <= end):
            continue
        key = f"{inv['quoteNo']}::{inv['invoiceNo']}"
        if ("quotation_payment", key) in confirmed:
            continue
        summary = f"{inv['customer']} {inv['quoteNo']} 發票{inv['invoiceNo']} 收款"[:60]
        bank_name = inv.get("bankAccountName") or "銀行存款"
        bank_code = inv.get("bankAccountCode") or ""
        lines = [
            _voucher_line(d, cfg["voucherCategory"], summary, bank_code, bank_name,
                          inv["amountTotal"], 0, cfg["departmentCode"], inv["quoteNo"], inv["customer"]),
            _voucher_line(d, cfg["voucherCategory"], summary, cfg["salesRevenueAccount"], "銷貨收入",
                          0, inv["amountPretax"], cfg["departmentCode"], inv["quoteNo"], inv["customer"]),
        ]
        if inv["taxAmount"]:
            lines.append(_voucher_line(d, cfg["voucherCategory"], summary, cfg["outputTaxAccount"], "銷項稅額",
                                        0, inv["taxAmount"], cfg["departmentCode"], inv["quoteNo"], inv["customer"]))
        events.append({
            "sourceType": "quotation_payment", "sourceKey": key,
            "date": d, "amount": inv["amountTotal"], "summary": summary, "lines": lines,
        })

    # 付款事件（承攬商費用）：借 承攬商費用(含稅) / 貸 銀行存款(含稅)
    for v in _collect_paid_contractor_vouchers(start, end):
        key = v["voucherNo"]
        if ("contractor_voucher", key) in confirmed:
            continue
        vendor = v["vendorName"] or "外包人員點工"
        summary = f"{vendor} {v['quoteNo']} 匯款申請{v['voucherNo']}"[:60]
        paid_d = (v["paidAt"] or "")[:10]
        bank_name = v.get("paidBankAccountName") or "銀行存款"
        bank_code = v.get("paidBankAccountCode") or ""
        lines = [
            _voucher_line(paid_d, cfg["voucherCategory"], summary, cfg["contractorExpenseAccount"], "承攬商費用",
                          v["grandTotal"], 0, cfg["departmentCode"], v["voucherNo"], vendor),
            _voucher_line(paid_d, cfg["voucherCategory"], summary, bank_code, bank_name,
                          0, v["grandTotal"], cfg["departmentCode"], v["voucherNo"], vendor),
        ]
        events.append({
            "sourceType": "contractor_voucher", "sourceKey": key,
            "date": paid_d, "amount": v["grandTotal"], "summary": summary, "lines": lines,
        })

    # 付款事件（料件/設備進貨）：借 料件設備成本 / 貸 銀行存款
    for b in _collect_paid_stock_batches(start, end):
        key = b["batch_no"]
        if ("stock_batch", key) in confirmed:
            continue
        supplier = b["supplier_name"] or "（未登記供應商）"
        summary = f"{supplier} {b['part_no']} 進貨批次{b['batch_no']}"[:60]
        total_cost = b["total_cost"] or 0
        category = b.get("category") or ""
        expense_code = (cfg["inventoryExpenseAccounts"] or {}).get(category, "")
        expense_name = f"料件設備成本（{category}）" if category else "料件設備成本"
        bank_name = b.get("paid_bank_account_name") or "銀行存款"
        bank_code = b.get("paid_bank_account_code") or ""
        lines = [
            _voucher_line(b["paid_at"], cfg["voucherCategory"], summary,
                          expense_code, expense_name,
                          total_cost, 0, cfg["departmentCode"], b["batch_no"], supplier),
            _voucher_line(b["paid_at"], cfg["voucherCategory"], summary,
                          bank_code, bank_name,
                          0, total_cost, cfg["departmentCode"], b["batch_no"], supplier),
        ]
        events.append({
            "sourceType": "stock_batch", "sourceKey": key,
            "date": b["paid_at"], "amount": total_cost, "summary": summary, "lines": lines,
        })

    events.sort(key=lambda e: (e["date"], e["sourceKey"]))
    return events


def _flatten_events_for_excel(events: list) -> list:
    """把事件層級清單攤平成傳票分錄列，補上每張傳票的流水傳票號
    （AR0001/AP0002，純顯示用，每次匯出重新編號，不是穩定識別碼——
    真正用來判斷「是否已匯入過」的是 sourceType/sourceKey，見 _collect_t100_events()）。"""
    rows = []
    ar_seq = ap_seq = pc_seq = 0
    for ev in events:
        if ev["sourceType"] == "quotation_payment":
            ar_seq += 1
            vno = f"AR{ar_seq:04d}"
        elif ev["sourceType"] == "contractor_voucher":
            ap_seq += 1
            vno = f"AP{ap_seq:04d}"
        else:
            pc_seq += 1
            vno = f"PC{pc_seq:04d}"
        for line in ev["lines"]:
            row = dict(line)
            row["voucherNo"] = vno
            rows.append(row)
    return rows


def _build_t100_voucher_excel(rows: list, start: str, end: str, cfg: dict, gen_at: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T100傳票匯出"
    ws.sheet_view.showGridLines = False
    mk, fill, mk_border, al = xl_style(wb)
    BD = mk_border()

    missing_codes = [k for k in ("salesRevenueAccount", "outputTaxAccount", "contractorExpenseAccount")
                      if not cfg.get(k)]
    if not cfg.get("bankAccounts"):
        missing_codes.append("bankAccounts（尚未設定任何銀行帳戶）")
    if not any((cfg.get("inventoryExpenseAccounts") or {}).values()):
        missing_codes.append("inventoryExpenseAccounts（料件分類科目代號皆未設定）")

    widths = [10, 12, 8, 30, 10, 12, 12, 12, 8, 14, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = company_heading(f"T100 傳票批次匯出（{start} ~ {end}）")
    c.font = mk(bold=True, size=13, color="FFFFFF")
    c.fill = fill("111827")
    c.alignment = al("center")
    ws.row_dimensions[1].height = 28

    note = f"產製時間：{gen_at}　現金基礎（僅含實際已收款/已匯款事件）　傳票別預設「{cfg['voucherCategory']}」"
    if missing_codes:
        note += "　⚠️ 尚未設定科目代號：" + "、".join(missing_codes) + "（請至出納頁「T100匯出」子頁籤按「展開科目代號設定」填入後再匯入 T100）"
    ws.merge_cells("A2:K2")
    c = ws["A2"]
    c.value = note
    c.font = mk(size=9, color="DC2626" if missing_codes else "6B7280")
    c.alignment = al("center", wrap=True)
    ws.row_dimensions[2].height = 18

    headers = ["傳票號", "傳票日期", "傳票別", "摘要", "科目代號", "科目名稱",
               "借方金額", "貸方金額", "部門別", "來源單號", "交易對象"]
    set_row(ws, 3, headers, font=mk(bold=True, color="FFFFFF"), fill=fill("2563EB"), border=BD, aligns=[al("center")])
    ws.row_dimensions[3].height = 22

    r = 4
    total_debit = total_credit = 0
    body_aligns = [al("center"), al("center"), al("center"), al("left"), al("center"),
                   al("center"), al("right"), al("right"), al("center"), al("center"), al("left")]
    for row in rows:
        set_row(ws, r, [
            row["voucherNo"], row["date"], row["category"], row["summary"],
            row["acctCode"], row["acctName"], row["debit"] or "", row["credit"] or "",
            row["dept"], row["sourceNo"], row["counterparty"],
        ], font=mk(), border=BD, aligns=body_aligns)
        total_debit += row["debit"]
        total_credit += row["credit"]
        r += 1

    set_row(ws, r, ["合計", "", "", "", "", "", total_debit, total_credit, "", "", ""],
             font=mk(bold=True), fill=fill("F9FAFB"), border=BD, aligns=body_aligns)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _require_t100_admin(authorization: str) -> dict:
    u = _require_user(authorization)
    if u["role"] not in ("superadmin", "admin"):
        raise HTTPException(403, "財務報告僅管理員以上可查閱")
    return u


def _validate_range(start: str, end: str) -> None:
    if not start or not end or start > end:
        raise HTTPException(400, "start/end 日期區間無效")


@router.get("/api/reports/t100-export/vouchers")
def t100_export_vouchers(
    start: str = Query(...),
    end: str = Query(...),
    authorization: str = Header(None),
):
    """T100 傳票批次匯出（Excel），現金基礎，涵蓋已收款發票／已匯款承攬商費用／
    已付款料件設備進貨；已標記「已匯入」的事件自動排除，不會重複出現在匯出檔裡。"""
    u = _require_t100_admin(authorization)
    _validate_range(start, end)
    check_export_rate(u["id"], "excel")
    cfg = _t100_config()
    events = _collect_t100_events(start, end)
    rows = _flatten_events_for_excel(events)
    gen_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    xlsx = _build_t100_voucher_excel(rows, start, end, cfg, gen_at)
    fname = f"MOTRIX_T100傳票匯出_{start}_{end}.xlsx"
    return StreamingResponse(
        io.BytesIO(xlsx),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_url_quote(fname)}"},
    )


@router.get("/api/reports/t100-export/preview")
def t100_export_preview(
    start: str = Query(...),
    end: str = Query(...),
    authorization: str = Header(None),
):
    """預覽本期尚未標記「已匯入」的事件（JSON，非 Excel），供財務在正式標記
    已匯入前先核對筆數/金額。"""
    _require_t100_admin(authorization)
    _validate_range(start, end)
    events = _collect_t100_events(start, end)
    return {
        # 對方不在時：預覽明說少了哪一類（匯出的 Excel 是 T100 匯入檔，不在裡面加說明列）
        #   IP-14（M04）⇒ 承攬商付款；receivables.tax_invoices（M05）⇒ 收款事件。兩者都缺 ⇒ 兩句都列
        "notice": "；".join(x for x in (
            "" if _registry.single_provider("contractor_voucher.public") else T100_CONTRACTOR_MISSING,
            "" if _registry.single_provider("receivables.tax_invoices") else T100_RECEIVABLES_MISSING,
        ) if x),
        "count": len(events),
        "totalAmount": sum(e["amount"] for e in events),
        "events": [
            {"sourceType": e["sourceType"], "sourceKey": e["sourceKey"],
             "date": e["date"], "amount": e["amount"], "summary": e["summary"]}
            for e in events
        ],
    }


class T100ConfirmBody(BaseModel):
    start: str
    end: str


@router.post("/api/reports/t100-export/confirm")
def t100_export_confirm(body: T100ConfirmBody, authorization: str = Header(None)):
    """財務確認「這個區間內尚未標記的事件已經實際匯入 T100」——標記後這些
    事件會從之後所有匯出/預覽自動排除，避免重複匯入。冪等：已標記過的事件
    這次呼叫不會出現在候選清單裡（_collect_t100_events 預設排除已確認）。"""
    u = _require_t100_admin(authorization)
    _validate_range(body.start, body.end)
    events = _collect_t100_events(body.start, body.end)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_db()
    try:
        for ev in events:
            conn.execute(
                "INSERT OR IGNORE INTO t100_export_confirmations "
                "(source_type, source_key, event_date, amount, summary, confirmed_by, confirmed_at) "
                "VALUES (?,?,?,?,?,?,?)",
                (ev["sourceType"], ev["sourceKey"], ev["date"], ev["amount"], ev["summary"],
                 u["username"], now),
            )
        conn.commit()
    finally:
        conn.close()
    _audit(_tok(authorization), "reports.t100_export.confirm", "t100_export",
           f"{body.start}~{body.end}", f"確認 {len(events)} 筆事件已匯入 T100")
    return {"confirmedCount": len(events)}


@router.get("/api/reports/t100-export/confirmed")
def t100_export_confirmed_list(
    start: str = Query(None),
    end: str = Query(None),
    authorization: str = Header(None),
):
    """已標記「已匯入」的事件清單（稽核／複核用），可選日期區間篩選。"""
    _require_t100_admin(authorization)
    conn = get_db()
    try:
        if start and end:
            rows = conn.execute(
                "SELECT * FROM t100_export_confirmations WHERE event_date BETWEEN ? AND ? "
                "ORDER BY confirmed_at DESC", (start, end),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM t100_export_confirmations ORDER BY confirmed_at DESC LIMIT 500"
            ).fetchall()
        return [
            {"sourceType": r["source_type"], "sourceKey": r["source_key"],
             "date": r["event_date"], "amount": r["amount"], "summary": r["summary"],
             "confirmedBy": r["confirmed_by"], "confirmedAt": r["confirmed_at"]}
            for r in rows
        ]
    finally:
        conn.close()


class T100UnconfirmBody(BaseModel):
    sourceType: str
    sourceKey: str


@router.post("/api/reports/t100-export/unconfirm")
def t100_export_unconfirm(body: T100UnconfirmBody, authorization: str = Header(None)):
    """撤銷單一事件的「已匯入」標記（標記錯誤時的救援手段），撤銷後該事件
    會在下次涵蓋其日期的匯出/預覽重新出現。"""
    u = _require_t100_admin(authorization)
    conn = get_db()
    try:
        cur = conn.execute(
            "DELETE FROM t100_export_confirmations WHERE source_type=? AND source_key=?",
            (body.sourceType, body.sourceKey),
        )
        conn.commit()
        if cur.rowcount == 0:
            raise HTTPException(404, "找不到對應的已匯入標記")
    finally:
        conn.close()
    _audit(_tok(authorization), "reports.t100_export.unconfirm", "t100_export",
           f"{body.sourceType}:{body.sourceKey}", "撤銷 T100 已匯入標記")
    return {"ok": True}


# ── 連接器（docs/platform/INTEGRATION-POINTS.md，契約版本 1）──────────────────────────────
# IP-2 voucher.account_check：科目代號有效性（與設定頁同一條規則）。回 (ok, err)。
_registry.provide("voucher.account_check", "accounting", validate_account_code)


# IP-3 accounting.settings：只公開別組需要的那一小塊（付款銀行清單與預設），不給整份 T100 設定。
def _provide_accounting_settings() -> dict:
    cfg = _t100_config()
    return {"bankAccounts": [{"name": b.get("name") or "", "acctCode": b.get("acctCode") or ""}
                             for b in (cfg.get("bankAccounts") or []) if b.get("acctCode")],
            "defaultBankAccountCode": cfg.get("defaultBankAccountCode") or ""}


_registry.provide("accounting.settings", "accounting", _provide_accounting_settings)

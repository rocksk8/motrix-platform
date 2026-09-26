"""CT1 外包名冊「分行」（2026-09-24 使用者：「外包名冊增加分行的欄位」；裁示見 HANDOFF「CT1」）。

contractors 表與編輯表單早有 bank_branch（08-01），這裡補使用者要的四處：
① 名冊清單單獨一欄「分行」（D1：不與銀行合成一欄）
② 簽核佇列的匯款申請：每位外包人員一列（姓名、銀行代碼＋名稱＋分行、戶名、帳號）（D2）
③ 勞報單 PDF「乙方銀行存簿影本」頁的對照表加「銀行／分行」（D3）
④ 名冊 Excel 匯入／匯出（D4）：匯出含分行與帳號、證件號碼只顯示末 4 碼、寫稽核；匯入以姓名＋證件號碼
   比對，舊檔沒有分行欄照樣可匯入；權限比照名冊編輯（superadmin 或 contractor_list 模組）。
"""
import io
import json

import pytest

PASSBOOK = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGA"
            "hKmMIQAAAABJRU5ErkJggg==")


def _login(client, username, password):
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _add_contractor(name, id_number, branch="台中分局", **kw):
    import db
    conn = db.get_db()
    try:
        cur = conn.execute(
            "INSERT INTO contractors (name, id_number, nationality, phone, bank_code, bank_name, bank_branch,"
            " bank_account_name, bank_account_number, notes, active, created_at, updated_at)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,1,?,?)",
            (name, id_number, "本國籍", kw.get("phone", "0912"), "700", "中華郵政", branch, name,
             kw.get("acct", "00012345678901"), "", "2026-01-01", "2026-01-01"))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def _contractor(cid):
    import db
    conn = db.get_db()
    try:
        return dict(conn.execute("SELECT * FROM contractors WHERE id=?", (cid,)).fetchone())
    finally:
        conn.close()


# ── ③ 勞報單 PDF 存簿影本頁 ──────────────────────────────────────────────────

def test_payslip_passbook_page_shows_bank_and_branch():
    from pdf_gen import _build_payslip_html
    html = _build_payslip_html({
        "slipNo": "PS-T1", "contractorName": "王小明", "bankCode": "700", "bankName": "中華郵政",
        "bankBranch": "台中分局", "bankAccountName": "王小明", "bankAccountNumber": "00012345678901",
        "_bank_passbook": PASSBOOK,
    })
    assert "附件：乙方銀行存簿影本" in html, "量尺：有存簿影本才會出這一頁"
    page = html.split("附件：乙方銀行存簿影本", 1)[1]
    assert "台中分局" in page and "中華郵政" in page, "存簿影本頁的對照表要有銀行／分行"


# ── ② 簽核佇列的匯款申請 ──────────────────────────────────────────────────────

def _seed_voucher():
    import db
    snap = {"vendorName": "某承攬商", "bankName": "臺灣銀行", "bankBranch": "總行", "bankAccountNumber": "111",
            "personnel": [
                {"id": 1, "name": "外包甲", "amount": 1000, "bankCode": "700", "bankName": "中華郵政",
                 "bankBranch": "台中分局", "bankAccountName": "外包甲", "bankAccountNumber": "00012345678901"},
                {"id": 2, "name": "外包乙", "amount": 800, "bankCode": "812", "bankName": "台新銀行",
                 "bankBranch": "", "bankAccountName": "外包乙", "bankAccountNumber": "22222"}]}
    conn = db.get_db()
    try:
        conn.execute(
            "INSERT INTO quotations (quote_no, status, customer_name, project_name, data_json, created_at, updated_at)"
            " VALUES (?,?,?,?,?,?,?)", ("MQ-CT1-001", "已送出", "客", "案", "{}", "2026-01-01", "2026-01-01"))
        conn.execute(   # 匯款申請的 dispatch_id 有外鍵，要先有一張派工單
            "INSERT INTO contractor_dispatches (id, quote_no, vendor_id, dispatch_date, scope, items_json,"
            " personnel_json, total_amount, tax_rate, status, notes, created_by, created_at, updated_at,"
            " accepted_at, accepted_by, files_json, invoice_files_json) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (99101, "MQ-CT1-001", None, "2026-09-01", "測試", "[]", "[]", 1800, 0, "已完成", "", "someone",
             "2026-01-01T00:00:00", "2026-01-01T00:00:00", "", "", "[]", "[]"))
        conn.execute(
            "INSERT INTO contractor_payment_vouchers (voucher_no, quote_no, dispatch_id, status, snapshot_json,"
            " data_json, created_by, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?)",
            ("CV-CT1-001", "MQ-CT1-001", 99101, "待審核", json.dumps(snap, ensure_ascii=False),
             json.dumps({"approval": {"tiers": [{"approvers": [{"username": "ct1_sa", "status": "pending"}]}],
                                      "currentTier": 0}}, ensure_ascii=False),
             "someone", "2026-01-01T00:00:00", "2026-01-01T00:00:00"))
        conn.commit()
    finally:
        conn.close()


def test_approval_queue_voucher_lists_each_personnel_bank_with_branch(client, make_user):
    h = _login(client, *make_user(username="ct1_sa", role="superadmin"))
    _seed_voucher()
    body = client.get("/api/approval-queue", headers=h).json()
    items = [it for g in body["queue"] for it in g["items"] if it.get("quoteNo") == "CV-CT1-001"
             or it.get("voucherNo") == "CV-CT1-001" or "CV-CT1-001" in json.dumps(it, ensure_ascii=False)]
    assert items, "量尺：佇列裡要找得到這張匯款申請"
    people = items[0].get("personnelBanks")
    assert people == [
        {"name": "外包甲", "bankCode": "700", "bankName": "中華郵政", "bankBranch": "台中分局",
         "bankAccountName": "外包甲", "bankAccountNumber": "00012345678901"},
        {"name": "外包乙", "bankCode": "812", "bankName": "台新銀行", "bankBranch": "",
         "bankAccountName": "外包乙", "bankAccountNumber": "22222"},
    ], people


def test_approval_queue_page_renders_personnel_bank_rows():
    """頁面要把 personnelBanks 畫出來（每人一列，含分行）。"""
    from pathlib import Path
    html = (Path(__file__).resolve().parents[4] / "frontend" / "pages" / "approval-queue.html").read_text(encoding="utf-8")
    assert 'x-for="pb in (selected?.personnelBanks' in html
    assert "pb.bankBranch" in html and "pb.bankAccountNumber" in html


# ── ① 名冊清單的「分行」欄 ────────────────────────────────────────────────────

def test_contractor_list_has_a_branch_column():
    from pathlib import Path
    import re
    html = (Path(__file__).resolve().parents[4] / "frontend" / "pages" / "contractors.html").read_text(encoding="utf-8")
    heads = re.findall(r"<th[^>]*>([^<]*)</th>", html)
    assert "分行" in heads, heads
    assert "c.bank_branch" in html, "清單列要顯示每個人的分行"


# ── ④ Excel 匯入／匯出 ───────────────────────────────────────────────────────

def _xlsx_rows(content):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    return rows[0], rows[1:]


def test_export_has_branch_account_and_masked_id_and_is_audited(client, make_user):
    import db
    h = _login(client, *make_user(username="ct1_ex", role="superadmin"))
    _add_contractor("王小明", "A123456789")
    r = client.get("/api/contractors/export", headers=h)
    assert r.status_code == 200, r.text
    head, rows = _xlsx_rows(r.content)
    assert "分行" in head and "帳號" in head and "證件號碼" in head, head
    row = dict(zip(head, rows[0]))
    assert row["分行"] == "台中分局" and row["帳號"] == "00012345678901"
    assert row["證件號碼"] == "******6789", "證件號碼只可以顯示末 4 碼"
    assert "A123456789" not in r.content.decode("latin-1", errors="ignore")
    conn = db.get_db()
    try:
        n = conn.execute("SELECT COUNT(*) FROM audit_log WHERE action='contractors.export'").fetchone()[0]
    finally:
        conn.close()
    assert n == 1, "匯出要寫稽核"


def _xlsx(head, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, h, content):
    return client.post("/api/contractors/import", headers=h,
                       files={"file": ("名冊.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_import_old_file_without_branch_column_updates_and_keeps_branch(client, make_user):
    h = _login(client, *make_user(username="ct1_im", role="superadmin"))
    cid = _add_contractor("王小明", "A123456789", branch="台中分局")
    r = _import(client, h, _xlsx(["姓名", "證件號碼", "電話"], [["王小明", "A123456789", "0987654321"],
                                                          ["新來的", "B223456789", "0911"]]))
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["updated"] == 1 and res["created"] == 1 and res["errors"] == [], res
    c = _contractor(cid)
    assert c["phone"] == "0987654321"
    assert c["bank_branch"] == "台中分局", "舊檔沒有分行欄：不可以把既有分行清掉"


def test_import_re_imports_an_export_by_name_and_last4(client, make_user):
    h = _login(client, *make_user(username="ct1_rt", role="superadmin"))
    cid = _add_contractor("王小明", "A123456789", branch="台中分局")
    content = client.get("/api/contractors/export", headers=h).content
    head, rows = _xlsx_rows(content)
    rows[0][head.index("分行")] = "北屯分局"
    r = _import(client, h, _xlsx(head, rows))
    assert r.status_code == 200, r.text
    assert r.json()["updated"] == 1, r.json()
    c = _contractor(cid)
    assert c["bank_branch"] == "北屯分局"
    assert c["id_number"] == "A123456789", "遮罩的證件號碼不可以寫回去"


def test_import_ambiguous_or_masked_unknown_rows_are_reported_not_guessed(client, make_user):
    h = _login(client, *make_user(username="ct1_amb", role="superadmin"))
    _add_contractor("同名", "A100006789")
    _add_contractor("同名", "B200006789")
    r = _import(client, h, _xlsx(["姓名", "證件號碼", "分行"], [["同名", "******6789", "X"],
                                                          ["查無", "******1111", "Y"]]))
    res = r.json()
    assert res["updated"] == 0 and res["created"] == 0, res
    assert len(res["errors"]) == 2 and all("第" in e["message"] for e in res["errors"]), res


def test_import_export_require_roster_edit_permission(client, make_user):
    h = _login(client, *make_user(username="ct1_no", role="admin", modules=["dashboard"]))
    assert client.get("/api/contractors/export", headers=h).status_code == 403
    assert _import(client, h, _xlsx(["姓名"], [["x"]])).status_code == 403

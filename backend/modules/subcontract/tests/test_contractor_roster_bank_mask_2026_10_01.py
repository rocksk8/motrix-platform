# -*- coding: utf-8 -*-
"""外包名冊（/api/contractors*）收款帳號遮蔽（使用者裁示 2026-10-01：只有最高管理者看得到完整帳號）。

缺口（主持 2026-10-01 確認）：這些端點的守門是 `_require_user(require_superadmin=True, module='contractor_list')`，
helpers/auth.py 會放行「持 contractor_list 模組的非最高管理者」⇒ 列表（_LIST_COLS）與詳情（SELECT *）原本回完整帳號。
先前回報「名冊本來就只有最高管理者」是錯的。
涵蓋：列表／詳情／存簿影本／匯出／匯入（遮蔽值不寫回）／編輯（遮蔽值送回＝保留原帳號）。
反向控制：把 `can_see_full` 改成永遠 True ⇒ 持模組者看到全碼。
"""
import io
import json

import pytest

FULL = "00012345678901"
MASKED = "****8901"
PASSBOOK = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGA"


def _login(client, u, p):
    r = client.post("/api/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}


@pytest.fixture()
def world(client, make_user):
    sa, sapw = make_user(username="rb_sa", role="superadmin")[:2]
    ad, adpw = make_user(username="rb_admin", role="admin", modules=["contractor_list"])[:2]
    h_sa, h_ad = _login(client, sa, sapw), _login(client, ad, adpw)
    r = client.post("/api/contractors", headers=h_sa, json={
        "name": "測試外包甲", "id_number": "A123456789", "bank_code": "700", "bank_name": "中華郵政", "bank_branch": "台中",
        "bank_account_name": "測試外包甲", "bank_account_number": FULL})
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    assert client.put("/api/contractors/%d/id-card" % cid, headers=h_sa, json={"bank_passbook": PASSBOOK}).status_code == 200
    return client, h_sa, h_ad, cid


def _stored(cid):
    import db
    conn = db.get_db()
    try:
        return conn.execute("SELECT bank_account_number FROM contractors WHERE id=?", (cid,)).fetchone()[0]
    finally:
        conn.close()


def test_list_and_detail_masked_for_module_holder_full_for_superadmin(world):
    client, h_sa, h_ad, cid = world
    for hdr, want in ((h_sa, FULL), (h_ad, MASKED)):                          # 正向控制：最高管理者看全碼
        lst = client.get("/api/contractors", headers=hdr)
        assert lst.status_code == 200, lst.text
        assert [x["bank_account_number"] for x in lst.json() if x["id"] == cid] == [want]
        d = client.get("/api/contractors/%d" % cid, headers=hdr).json()
        assert d["bank_account_number"] == want and d["has_passbook"] is True
    r = client.get("/api/contractors/%d" % cid, headers=h_ad)
    assert FULL not in r.text and PASSBOOK not in r.text and r.json().get("bank_passbook_image", "") == ""


def test_passbook_image_only_for_superadmin(world):
    client, h_sa, h_ad, cid = world
    assert client.get("/api/contractors/%d/id-card" % cid, headers=h_sa).json()["bank_passbook"].startswith("data:image/")
    r = client.get("/api/contractors/%d/id-card" % cid, headers=h_ad).json()
    assert r["bank_passbook"] == "" and r["masked"] is True


def test_put_with_masked_value_keeps_the_stored_number(world):
    client, _h_sa, h_ad, cid = world
    body = {"name": "測試外包甲", "id_number": "A123456789", "bank_code": "700", "bank_name": "中華郵政", "bank_branch": "改了分行",
            "bank_account_name": "測試外包甲", "bank_account_number": MASKED}
    assert client.put("/api/contractors/%d" % cid, headers=h_ad, json=body).status_code == 200
    assert _stored(cid) == FULL, "遮蔽值被原樣存回，真帳號被覆蓋"
    body["bank_account_number"] = "11112222333344"                         # 輸入新帳號才會更換
    assert client.put("/api/contractors/%d" % cid, headers=h_ad, json=body).status_code == 200
    assert _stored(cid) == "11112222333344"


def _xlsx_rows(resp):
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    col = rows[0].index("帳號")
    return [r[col] for r in rows[1:]]


def test_export_masks_for_module_holder_full_for_superadmin(world):
    client, h_sa, h_ad, _cid = world
    r_sa = client.get("/api/contractors/export", headers=h_sa)
    assert r_sa.status_code == 200 and _xlsx_rows(r_sa) == [FULL]            # 正向控制
    r_ad = client.get("/api/contractors/export", headers=h_ad)
    assert r_ad.status_code == 200 and _xlsx_rows(r_ad) == [MASKED]


def test_import_of_a_masked_account_does_not_overwrite_the_number(world):
    """匯入檔裡的帳號是遮蔽值（持模組者匯出的檔、再由人補上別的欄位後匯入）⇒ 不寫回。直接造檔，不依賴匯出的遮蔽。"""
    import openpyxl
    client, h_sa, _h_ad, cid = world
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "證件號碼", "分行", "帳號"])
    ws.append(["測試外包甲", "A123456789", "匯入改的分行", MASKED])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/contractors/import", headers=h_sa,
                    files={"file": ("roster.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200 and r.json()["updated"] == 1, r.text
    assert _stored(cid) == FULL, "匯入遮蔽值把真帳號覆蓋掉"
    import db
    conn = db.get_db()
    try:
        assert conn.execute("SELECT bank_branch FROM contractors WHERE id=?", (cid,)).fetchone()[0] == "匯入改的分行"   # 其他欄位照常更新
    finally:
        conn.close()


def test_reverse_control_without_the_mask_module_holder_would_see_the_full_number(world, monkeypatch):
    client, _h_sa, h_ad, cid = world
    from modules.subcontract import bank_mask
    monkeypatch.setattr(bank_mask, "can_see_full", lambda user: True)
    assert client.get("/api/contractors/%d" % cid, headers=h_ad).json()["bank_account_number"] == FULL


def test_payslip_form_never_copies_a_masked_account_into_the_payslip():
    """勞報單頁從外包名冊挑人會帶入帳號；非最高管理者拿到的是遮蔽值，不可存進勞報單（會變成假帳號）。"""
    import pathlib
    html = (pathlib.Path(__file__).resolve().parents[4] / "frontend" / "pages" / "payslip-form.html").read_text(encoding="utf-8")
    assert "startsWith('****') ? ''" in html

"""需要外包工班（M04）的題：刪掉 modules/subcontract 時隨模組消失（PLAYBOOK §B-11，稽核 D M04-M1）。

（2026-09-26 自 tests/test_cashier_execution_history_2026_08_31.py 拆出：這幾題需要本模組在，隨模組搬走。原檔的說明：）
2026-08-31 v2：出納模組「執行歷史」（已匯款／已收款彙整）＋ Excel 匯出。
`/api/cashier/execution-history` 依 paid_at／receivedAt 落在區間內篩選，
`/api/cashier/export` 沿用 reports.py 既有 Excel 樣式 helper 產出兩個
sheet（已匯款明細／已收款明細）。
"""
import io
import json

import openpyxl


def _login(client, username, password):
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return r.json()["token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def _make_approved_voucher(client, token, quote_no, payable_date=None):
    r = client.post(
        "/api/vendor-contractors", headers=_auth(token),
        json={"name": f"廠商{quote_no}", "data": {}},
    )
    assert r.status_code == 201, r.text
    vendor_id = r.json()["id"]

    body = {
        "quote_no": quote_no, "vendor_id": vendor_id,
        "items_json": [{"description": "測試品項", "qty": 1, "unit": "式", "unitPrice": 10000, "amount": 10000}],
        "status": "completed",
    }
    if payable_date:
        body["payable_date"] = payable_date
    r = client.post("/api/contractor-dispatches", headers=_auth(token), json=body)
    assert r.status_code == 201, r.text
    did = r.json()["id"]

    cv = client.post("/api/contractor-vouchers", headers=_auth(token), json={"dispatch_id": did})
    assert cv.status_code == 201, cv.text
    voucher_no = cv.json()["voucher_no"]

    import db
    conn = db.get_db()
    try:
        conn.execute("UPDATE contractor_payment_vouchers SET status='已核准' WHERE voucher_no=?", (voucher_no,))
        conn.commit()
    finally:
        conn.close()
    return voucher_no


def _make_quotation_with_unreceived_item(quote_no, expected_receipt_date=""):
    import db
    conn = db.get_db()
    try:
        data_json = json.dumps({
            "dealTag": "已成案",
            "caseRecord": {
                "payment": {"items": [
                    {"type": "訂金款", "pct": 30, "amount": 30000, "received": False,
                     "expectedReceiptDate": expected_receipt_date},
                ]},
            },
        })
        conn.execute(
            "INSERT INTO quotations (quote_no, status, customer_name, project_name, total, pretax, "
            "data_json, created_at, updated_at, deal_tag, quote_date) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (quote_no, "已送出", "測試客戶", "測試專案", 30000, 28571, data_json,
             "2026-01-01T00:00:00", "2026-01-01T00:00:00", "已成案", "2026-01-01"),
        )
        conn.commit()
    finally:
        conn.close()


def test_execution_history_date_range_filters_outgoing_and_incoming(client, make_user):
    username, password = make_user(username="hist_admin1", role="superadmin")
    token = _login(client, username, password)

    v_in_range = _make_approved_voucher(client, token, "MQ-HIST-001")
    pay = client.post(
        f"/api/contractor-vouchers/{v_in_range}/paid-toggle", headers=_auth(token),
        json={"action": "pay", "paid_at": "2026-08-15"},
    )
    assert pay.status_code == 200, pay.text

    v_out_of_range = _make_approved_voucher(client, token, "MQ-HIST-002")
    pay2 = client.post(
        f"/api/contractor-vouchers/{v_out_of_range}/paid-toggle", headers=_auth(token),
        json={"action": "pay", "paid_at": "2026-01-05"},
    )
    assert pay2.status_code == 200, pay2.text

    _make_quotation_with_unreceived_item("MQ-HIST-010")
    mark = client.patch(
        "/api/quotations/MQ-HIST-010/payment/0", headers=_auth(token),
        json={"received": True, "receivedAt": "2026-08-20", "actualAmount": 30000, "feeAmount": 0},
    )
    assert mark.status_code == 200, mark.text

    _make_quotation_with_unreceived_item("MQ-HIST-011")
    mark2 = client.patch(
        "/api/quotations/MQ-HIST-011/payment/0", headers=_auth(token),
        json={"received": True, "receivedAt": "2026-02-01", "actualAmount": 30000, "feeAmount": 0},
    )
    assert mark2.status_code == 200, mark2.text

    r = client.get("/api/cashier/execution-history?start=2026-08-01&end=2026-08-31", headers=_auth(token))
    assert r.status_code == 200, r.text
    body = r.json()

    outgoing_nos = [v["voucherNo"] for v in body["outgoing"]]
    assert v_in_range in outgoing_nos
    assert v_out_of_range not in outgoing_nos

    incoming_quote_nos = [i["quoteNo"] for i in body["incoming"]]
    assert "MQ-HIST-010" in incoming_quote_nos
    assert "MQ-HIST-011" not in incoming_quote_nos

    assert body["outgoingTotal"] == sum(v["grandTotal"] for v in body["outgoing"])
    assert body["incomingTotal"] == sum(
        (i["actualAmount"] if i["actualAmount"] is not None else i["amount"]) for i in body["incoming"]
    )


def test_export_excel_has_both_sheets_with_data(client, make_user):
    username, password = make_user(username="hist_admin2", role="superadmin")
    token = _login(client, username, password)

    v = _make_approved_voucher(client, token, "MQ-HIST-020")
    pay = client.post(
        f"/api/contractor-vouchers/{v}/paid-toggle", headers=_auth(token),
        json={"action": "pay", "paid_at": "2026-08-10"},
    )
    assert pay.status_code == 200, pay.text

    _make_quotation_with_unreceived_item("MQ-HIST-021")
    mark = client.patch(
        "/api/quotations/MQ-HIST-021/payment/0", headers=_auth(token),
        json={"received": True, "receivedAt": "2026-08-12", "actualAmount": 29800, "feeAmount": 200},
    )
    assert mark.status_code == 200, mark.text

    r = client.get("/api/cashier/export?start=2026-08-01&end=2026-08-31", headers=_auth(token))
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    # 2026-09-25（CORE-SPEC 獎金分潤：送交出納）：最高管理者／出納多一張獎金發放明細（IP-8）
    assert wb.sheetnames == ["已匯款明細", "已收款明細", "獎金發放明細"]

    ws1 = wb["已匯款明細"]
    voucher_nos_in_sheet = [row[0].value for row in ws1.iter_rows(min_row=3, max_col=1)]
    assert v in voucher_nos_in_sheet

    ws2 = wb["已收款明細"]
    quote_nos_in_sheet = [row[0].value for row in ws2.iter_rows(min_row=3, max_col=1)]
    assert "MQ-HIST-021" in quote_nos_in_sheet

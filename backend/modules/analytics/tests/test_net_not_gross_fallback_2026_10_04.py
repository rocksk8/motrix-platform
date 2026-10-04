# -*- coding: utf-8 -*-
"""T38 F9：營運報表／儀表板的「實際毛利」不可在淨利為 0 時退回毛利；儀表板只算 finalized、比的是淨利率。

`settle.get("netProfit") or settle.get("grossProfit")` 把「淨利剛好 0」當成「沒有這個欄位」⇒ 顯示毛利（payroll/bonus.py 的〈null 不等於 0〉）。
舊 finalized（沒有 netProfit／netMarginPct 鍵）仍走毛利，值不改寫。
"""
import json

import pytest

from modules.analytics.tests.test_reports_logic_fixes_2026_08_28 import _insert_case


def _case(quote_no, summary, status="finalized"):
    _insert_case(quote_no, deal_tag="已結案", settlement={"status": status, "summary": summary})


def _mc(quote_no):
    from modules.analytics.api.reports import _collect
    data = _collect("2026-01-01", "2026-12-31")
    return {c["quoteNo"]: c for c in data["marginCases"]}[quote_no]


# ── reports._collect ──────────────────────────────────────────────────────────

def test_net_profit_zero_does_not_fall_back_to_gross(client):
    _case("T38-R1", {"netProfit": 0, "netMarginPct": 0, "grossProfit": 30000, "grossMarginPct": 30.0})
    c = _mc("T38-R1")
    assert c["grossProfit"] == 0
    assert c["actualMarginPct"] == 0.0


def test_net_negative_is_kept(client):
    _case("T38-R2", {"netProfit": -500, "netMarginPct": -0.5, "grossProfit": 30000, "grossMarginPct": 30.0})
    c = _mc("T38-R2")
    assert c["grossProfit"] == -500 and c["actualMarginPct"] == -0.5


def test_net_profit_is_rounded_not_truncated(client):
    _case("T38-R3", {"netProfit": 1234.6, "netMarginPct": 12.3, "grossProfit": 1, "grossMarginPct": 1.0})
    assert _mc("T38-R3")["grossProfit"] == 1235


def test_net_positive_unchanged(client):
    _case("T38-R4", {"netProfit": 5000, "netMarginPct": 5.0, "grossProfit": 30000, "grossMarginPct": 30.0})
    c = _mc("T38-R4")
    assert c["grossProfit"] == 5000 and c["actualMarginPct"] == 5.0


def test_legacy_finalized_without_net_keys_still_uses_gross(client):
    _case("T38-R5", {"grossProfit": 30000, "grossMarginPct": 30.0})
    c = _mc("T38-R5")
    assert c["grossProfit"] == 30000 and c["actualMarginPct"] == 30.0


def test_ytd_and_sales_aggregates_follow_net_zero(client):
    """下游讀者（業務績效 amProfitSum、年度毛利）不得因 0 被吃成毛利。"""
    from modules.analytics.api.reports import _collect
    _case("T38-R6", {"netProfit": 0, "netMarginPct": 0, "grossProfit": 30000, "grossMarginPct": 30.0})
    s = _collect("2026-01-01", "2026-12-31")["summary"]
    assert s["totalActualGrossProfit"] == 0


# ── dashboard.marginComparison ────────────────────────────────────────────────

def _admin_headers(client, make_user):
    username, password = make_user(username="t38_sa", role="superadmin")
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _set_net_margin(quote_no, pct):
    import db
    conn = db.get_db()
    try:
        conn.execute("UPDATE quotations SET net_margin_pct=? WHERE quote_no=?", (pct, quote_no))
        conn.commit()
    finally:
        conn.close()


def _comparison(client, make_user):
    h = _admin_headers(client, make_user)
    r = client.get("/api/dashboard/stats", headers=h)
    assert r.status_code == 200, r.text
    j = r.json()
    return {x["quoteNo"]: x for x in j["marginComparison"]}, j["settledSummary"]


def test_dashboard_draft_settlement_is_not_counted(client, make_user):
    _case("T38-D1", {"netProfit": 100, "netMarginPct": 10.0, "grossMarginPct": 40.0}, status="draft")
    _set_net_margin("T38-D1", 12.0)
    cmp, summ = _comparison(client, make_user)
    assert "T38-D1" not in cmp and not summ


def test_dashboard_compares_net_margin_not_gross(client, make_user):
    _case("T38-D2", {"netProfit": 1000, "netMarginPct": 10.0, "grossProfit": 4000, "grossMarginPct": 40.0, "quotedPretax": 10000})
    _set_net_margin("T38-D2", 12.0)
    cmp, summ = _comparison(client, make_user)
    assert cmp["T38-D2"]["actualMarginPct"] == 10.0
    assert cmp["T38-D2"]["grossProfit"] == 1000
    assert summ["totalGrossProfit"] == 1000 and summ["avgActualMarginPct"] == 10.0


def test_dashboard_net_zero_is_kept(client, make_user):
    _case("T38-D3", {"netProfit": 0, "netMarginPct": 0, "grossProfit": 4000, "grossMarginPct": 40.0, "quotedPretax": 10000})
    _set_net_margin("T38-D3", 12.0)
    cmp, _ = _comparison(client, make_user)
    assert cmp["T38-D3"]["actualMarginPct"] == 0.0 and cmp["T38-D3"]["grossProfit"] == 0


def test_dashboard_legacy_finalized_without_net_keys_uses_gross(client, make_user):
    _case("T38-D4", {"grossProfit": 4000, "grossMarginPct": 40.0, "quotedPretax": 10000})
    _set_net_margin("T38-D4", 12.0)
    cmp, _ = _comparison(client, make_user)
    assert cmp["T38-D4"]["actualMarginPct"] == 40.0 and cmp["T38-D4"]["grossProfit"] == 4000


# ── 字樣：預估／實際都是淨利口徑 ⇒ 標「淨利」；舊精算（沒有 netProfit 鍵）的實際欄加註 ────────────────────

LEGACY_NOTE = "（舊精算為毛利）"


def _data():
    from modules.analytics.api.reports import _augment_with_targets, _build_income_expense_scopes, _collect, _parse_period
    lab, d0, d1 = _parse_period("2026")
    data = _augment_with_targets(_collect(d0, d1, None), d0)
    data["arAging"] = []
    data.update(_build_income_expense_scopes(2026, "2026-03", None, basis="accrual"))
    return lab, data


def _seed_new_and_legacy():
    _case("T38-L-NEW", {"netProfit": 5000, "netMarginPct": 5.0, "grossProfit": 30000, "grossMarginPct": 30.0})
    _case("T38-L-OLD", {"grossProfit": 30000, "grossMarginPct": 30.0})


def test_legacy_flag_only_on_summaries_without_the_net_profit_key(client):
    from modules.analytics.api.reports import _collect
    _seed_new_and_legacy()
    _case("T38-L-ZERO", {"netProfit": 0, "netMarginPct": 0, "grossProfit": 30000, "grossMarginPct": 30.0})
    cases = {c["quoteNo"]: c for c in _collect("2026-01-01", "2026-12-31")["casesAll"]}
    assert cases["T38-L-NEW"]["actualIsGross"] is False
    assert cases["T38-L-ZERO"]["actualIsGross"] is False            # 淨利 0 是真的 0，不是舊精算
    assert cases["T38-L-OLD"]["actualIsGross"] is True


def test_excel_headers_say_net_note_only_on_legacy_rows_and_sheet_name_is_unchanged(client):
    import io
    import openpyxl
    from modules.analytics.api.reports import _build_excel
    _seed_new_and_legacy()
    lab, data = _data()
    wb = openpyxl.load_workbook(io.BytesIO(_build_excel(data, lab, "t")))
    assert "毛利分析" in wb.sheetnames and "利潤分析" not in wb.sheetnames        # 使用者裁：工作表名不改
    ws = wb["案件清單"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert {"預估淨利率", "實際淨利率", "實際淨利"} <= set(cells)
    assert not {"預估毛利率", "實際毛利率", "實際毛利"} & set(cells)
    rows = {r[0].value: r for r in ws.iter_rows() if r[0].value in ("T38-L-NEW", "T38-L-OLD")}
    assert LEGACY_NOTE in str(rows["T38-L-OLD"][12].value)
    assert rows["T38-L-NEW"][12].value == "5.0%"
    sales = [str(c.value) for row in wb["業務員績效"].iter_rows() for c in row if c.value is not None] if "業務員績效" in wb.sheetnames else []
    assert "實際毛利率" not in sales


def test_pdf_html_labels_are_net_and_note_only_on_legacy_rows(client):
    from modules.analytics.api.reports import _build_report_html
    _seed_new_and_legacy()
    lab, data = _data()
    html = _build_report_html(data, lab, "t")
    for gone in ("預估毛利率", "實際毛利率", "實際毛利<", "精算實際毛利合計", "平均淨毛利率", "年度實際毛利"):
        assert gone not in html, gone
    for there in ("預估淨利率", "實際淨利率", "精算實際淨利合計", "利潤分析"):
        assert there in html, there
    assert html.count(LEGACY_NOTE) >= 1
    # 新格式案件的列不帶註記：整份只有舊案那兩處（案件清單＋利潤分析）才出現
    rows = [seg.split("</tr>")[0] for seg in html.split("<tr")]
    old_rows = [seg for seg in rows if "<td>T38-L-OLD</td>" in seg]
    new_rows = [seg for seg in rows if "<td>T38-L-NEW</td>" in seg]
    assert old_rows and all(LEGACY_NOTE in seg for seg in old_rows)
    assert new_rows and not any(LEGACY_NOTE in seg for seg in new_rows)


def test_gross_wording_is_kept_where_it_really_is_gross(client):
    from modules.analytics.api.reports import _build_report_html
    _seed_new_and_legacy()
    lab, data = _data()
    html = _build_report_html(data, lab, "t")
    assert "真實毛利率" in html and "原始直接毛利" in html and "原始毛利率" in html


def test_screen_page_and_script_use_net_wording_and_the_same_legacy_note():
    from pathlib import Path
    from core import source_tree
    from modules.analytics.api.reports import _LEGACY_GROSS_NOTE
    assert _LEGACY_GROSS_NOTE == LEGACY_NOTE
    page = source_tree.page_file("reports.html").read_text(encoding="utf-8")
    js = (Path(source_tree.FRONTEND_PAGES).parent / "js" / "reports.js").read_text(encoding="utf-8")
    for gone in ("預估毛利率", "實際毛利率", "預估毛利<", "實際毛利<", "精算實際毛利", "平均淨毛利率", "年度實際毛利", "毛利率比較"):
        assert gone not in page, gone
    for gone in ("預估毛利率", "實際毛利率", "平均毛利率", "平均淨毛利率", "年度實際毛利"):
        assert gone not in js, gone
    for there in ("預估淨利率", "實際淨利率", "預估淨利", "實際淨利", "精算實際淨利", "利潤分析", "淨利率比較"):
        assert there in page, there
    assert page.count(LEGACY_NOTE) == 4                              # 兩張案件表＋利潤分析表的實際率與實際金額
    assert page.count("actualIsGross") == 4
    for keep in ("真實毛利", "原始直接毛利", "原始毛利率"):
        assert keep in page, keep

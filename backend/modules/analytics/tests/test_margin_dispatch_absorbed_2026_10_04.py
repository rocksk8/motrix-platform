# -*- coding: utf-8 -*-
"""T38：利潤分析（Excel「毛利分析」表、PDF 明細）的成本分項要加得回「實際總成本」。

精算摘要新鍵 `dispatchAbsorbedTotal`（M01 寫入；已併入品項實際成本的承攬金額）。
未併入＝dispatchTotal − dispatchAbsorbedTotal ⇒ 併進 Excel 的「額外支出」欄；PDF 另列一行並在 absorbed>0 時註明。
舊精算沒有這個鍵 ⇒ 未併入＝0、輸出與以前完全相同。
"""
import io

import openpyxl

from modules.analytics.tests.test_reports_logic_fixes_2026_08_28 import _insert_case

NEW = {"netProfit": 5000, "netMarginPct": 5.0, "grossProfit": 9000, "grossMarginPct": 9.0,
       "itemActualTotal": 1000, "itemPoUnadopted": 200, "extraTotal": 300,
       "dispatchTotal": 500, "dispatchAbsorbedTotal": 120, "totalActualCost": 1880}
LEGACY = {k: v for k, v in NEW.items() if k not in ("dispatchTotal", "dispatchAbsorbedTotal")} | {"totalActualCost": 1500}


def _seed():
    _insert_case("T38-X-NEW", deal_tag="已結案", settlement={"status": "finalized", "summary": NEW})
    _insert_case("T38-X-OLD", deal_tag="已結案", settlement={"status": "finalized", "summary": LEGACY})


def _data():
    from modules.analytics.api.reports import _augment_with_targets, _build_income_expense_scopes, _collect, _parse_period
    lab, d0, d1 = _parse_period("2026")
    data = _augment_with_targets(_collect(d0, d1, None), d0)
    data["arAging"] = []
    data.update(_build_income_expense_scopes(2026, "2026-03", None, basis="accrual"))
    return lab, data


def _excel_rows():
    from modules.analytics.api.reports import _build_excel
    lab, data = _data()
    ws = openpyxl.load_workbook(io.BytesIO(_build_excel(data, lab, "t")))["毛利分析"]
    return {r[0].value: r for r in ws.iter_rows() if r[0].value in ("T38-X-NEW", "T38-X-OLD", "T38-X-PAGE")}


def test_excel_item_plus_extra_equals_total_when_dispatch_is_partly_absorbed(client):
    _seed()
    r = _excel_rows()["T38-X-NEW"]
    item, extra, total = r[10].value, r[11].value, r[12].value          # 品項成本／額外支出／實際總成本
    assert (item, extra, total) == (1200, 680, 1880)                     # 300 + (500 − 120)
    assert item + extra == total


def test_excel_legacy_summary_is_unchanged(client):
    _seed()
    r = _excel_rows()["T38-X-OLD"]
    assert (r[10].value, r[11].value, r[12].value) == (1200, 300, 1500)


def test_excel_total_row_follows_the_folded_extra(client):
    from modules.analytics.api.reports import _build_excel
    _seed()
    lab, data = _data()
    ws = openpyxl.load_workbook(io.BytesIO(_build_excel(data, lab, "t")))["毛利分析"]
    tot = [r for r in ws.iter_rows() if r[0].value == "合計"][-1]
    assert tot[11].value == 680 + 300


def test_pdf_detail_lists_unabsorbed_dispatch_and_the_absorbed_note(client):
    from modules.analytics.api.reports import _build_report_html
    _seed()
    lab, data = _data()
    html = _build_report_html(data, lab, "t")
    blocks = html.split("各案件利潤分析明細")[1].split("page-break-inside:avoid")
    new = [b for b in blocks if "T38-X-NEW" in b][0]
    old = [b for b in blocks if "T38-X-OLD" in b][0]
    assert "承攬商（未併入品項成本）" in new and "NT$ 380" in new
    assert "其中 NT$ 120 已併入品項實際成本" in new
    assert "承攬商（未併入品項成本）" not in old and "已併入品項實際成本" not in old


# ── t36／t37 完結的案件只帶頁面鍵 dispatchAbsorbed（稽核 S-2）──────────────────

PAGE_KEY_ONLY = {k: v for k, v in NEW.items() if k != "dispatchAbsorbedTotal"} | {"dispatchAbsorbed": 120}


def test_page_key_only_summary_still_sums_to_total_cost(client):
    _insert_case("T38-X-PAGE", deal_tag="已結案", settlement={"status": "finalized", "summary": PAGE_KEY_ONLY})
    r = _excel_rows()["T38-X-PAGE"]
    assert (r[10].value, r[11].value, r[12].value) == (1200, 680, 1880)


def test_total_key_wins_over_page_key_and_both_absent_is_unchanged():
    from modules.analytics.api.reports import _dispatch_split
    assert _dispatch_split({"dispatchTotal": 500, "dispatchAbsorbedTotal": 120, "dispatchAbsorbed": 999}) == (380, 120)
    assert _dispatch_split({"dispatchTotal": 500, "dispatchAbsorbed": 120}) == (380, 120)
    assert _dispatch_split({"dispatchTotal": 500}) == (0, 0)
    assert _dispatch_split({"dispatchTotal": 500, "dispatchAbsorbedTotal": 0, "dispatchAbsorbed": 120}) == (500, 0)   # 有總計鍵（含 0）就以它為準

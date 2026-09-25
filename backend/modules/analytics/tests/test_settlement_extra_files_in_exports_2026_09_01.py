"""自 `tests/test_settlement_extra_files_in_exports_2026_09_01.py` 拆出（M08 搬遷反向控制：這幾題需要營運分析模組，拿掉模組時一起消失）。"""
import json
from tests.test_settlement_extra_files_in_exports_2026_09_01 import (  # noqa: E402,F401  含 fixture
    _auth,
    _login,
    _make_closed_quotation,
)


def test_monthly_expense_report_details_include_extra_expense_files(client, make_user):
    username, password = make_user(role="admin")
    token = _login(client, username, password)
    _make_closed_quotation("MQ-EXPFILE-002")

    r = client.get("/api/reports/expenses-monthly?year=2026", headers=_auth(token))
    assert r.status_code == 200, r.text
    body = r.json()
    detail = next(d for d in body["expenses"]["details"]["other"] if d["quoteNo"] == "MQ-EXPFILE-002")
    assert detail["files"][0]["filename"] == "receipt.pdf"
    assert detail["date"] == "2026-07-15"


def test_financial_excel_export_includes_extra_expense_filename(client, make_user):
    """走真實的《營運報表》Excel 匯出端點（_build_excel()，純 openpyxl 組裝，
    不像 PDF 需要 Edge headless），確認「當月收支」分頁的支出明細真的帶出
    發票/收據附件檔名，不是只有 API JSON 層有資料。"""
    import io as _io
    from openpyxl import load_workbook
    username, password = make_user(role="admin")
    token = _login(client, username, password)
    _make_closed_quotation("MQ-EXPFILE-003")

    r = client.get(
        "/api/reports/financial/excel?period=2026&expense_month=2026-07",
        headers=_auth(token),
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(_io.BytesIO(r.content))
    ws = wb["當月收支"]
    found = any(
        cell.value == "receipt.pdf"
        for row in ws.iter_rows()
        for cell in row
    )
    assert found

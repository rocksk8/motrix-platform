# -*- coding: utf-8 -*-
"""第 48 班 Q7（使用者裁示）：PDF 只在百分比『和全域預設不同』時才印；畫面標籤仍一律顯示；Excel「毛利分析」最右多一欄「管銷比率」。

守門：default 25 且 summary.overheadPct=25 ⇒ PDF／報表 PDF 的管銷列只寫「管銷分攤（直接毛利）」；=30 ⇒ 「（直接毛利 30%）」；舊口徑 ⇒ 「（報價稅前 10%）」；
預設被改成 30 時 30 就不印、25 反而要印。Excel 既有欄號不動，新欄附在最右（第 27 欄）。
"""
import io

import openpyxl
import pytest

import pdf_gen
from modules.analytics.tests.test_original_indirect_reserve_2026_10_05 import _data
from modules.analytics.tests.test_reports_logic_fixes_2026_08_28 import _insert_case

V2 = {"formulaVer": 2, "overheadPct": 25}


def _set_default(value):
    import db
    import json
    c = db.get_db()
    try:
        c.execute("INSERT INTO system_settings(key, value_json, updated_at) VALUES('overhead_default_pct', ?, 'x') "
                  "ON CONFLICT(key) DO UPDATE SET value_json=excluded.value_json", (json.dumps(value),))
        c.commit()
    finally:
        c.close()


def test_label_hides_pct_only_when_equal_to_default_for_pdf_but_not_on_screen(client):
    _set_default(25)
    assert pdf_gen.admin_cost_label(V2, hide_default=True) == "管銷分攤（直接毛利）"
    assert pdf_gen.admin_cost_label({"formulaVer": 2, "overheadPct": 30}, hide_default=True) == "管銷分攤（直接毛利 30%）"
    assert pdf_gen.admin_cost_label({"formulaVer": 2, "overheadPct": 7.5}, hide_default=True) == "管銷分攤（直接毛利 7.5%）"
    assert pdf_gen.admin_cost_label({}, hide_default=True) == "管銷分攤（報價稅前 10%）"          # 舊口徑一定要說清楚基數
    assert pdf_gen.admin_cost_label(V2) == "管銷分攤（直接毛利 25%）", "沒有 hide_default（畫面用）一律顯示"
    _set_default(30)                                                                      # 預設改了：以『現在的預設』為準
    assert pdf_gen.admin_cost_label({"formulaVer": 2, "overheadPct": 30}, hide_default=True) == "管銷分攤（直接毛利）"
    assert pdf_gen.admin_cost_label(V2, hide_default=True) == "管銷分攤（直接毛利 25%）"
    _set_default("壞掉")                                                                  # 壞值 ⇒ 25
    assert pdf_gen.admin_cost_label(V2, hide_default=True) == "管銷分攤（直接毛利）"


def test_closing_pdf_and_ops_report_pdf_use_the_hide_default_rule(client):
    _set_default(25)
    src = open(pdf_gen.__file__, encoding="utf-8").read()
    assert src.count("admin_cost_label(summary, True, hide_default=True)") == 1 and src.count("admin_cost_label(summary, hide_default=True)") == 1
    rep = open(__import__("modules.analytics.api.reports", fromlist=["x"]).__file__, encoding="utf-8").read()
    assert rep.count("admin_cost_label(ss, True, hide_default=True)") == 1 and rep.count("admin_cost_label(ss, hide_default=True)") == 1


BASE = {"netProfit": 26630, "netMarginPct": 26.6, "grossProfit": 37000, "grossMarginPct": 37.0, "itemActualTotal": 60000, "extraTotal": 0,
        "totalActualCost": 63000, "origTotalCost": 60000, "origNetProfit": 21630, "origNetMarginPct": 21.6, "profitDiff": 5000}


def _sheet():
    from modules.analytics.api.reports import _build_excel
    lab, data = _data()
    wb = openpyxl.load_workbook(io.BytesIO(_build_excel(data, lab, "t")))
    return wb["毛利分析"]


def test_excel_has_the_overhead_rate_column_at_the_far_right(client):
    _insert_case("Q7-NEW", deal_tag="已結案", settlement={"status": "finalized", "summary": dict(BASE, formulaVer=2, overheadPct=25)})
    _insert_case("Q7-CUSTOM", deal_tag="已結案", settlement={"status": "finalized", "summary": dict(BASE, formulaVer=2, overheadPct=7.5)})
    _insert_case("Q7-OLD", deal_tag="已結案", settlement={"status": "finalized", "summary": BASE})
    ws = _sheet()
    hdr = [c.value for c in ws[3]]
    assert hdr[-1] == "管銷比率" and len(hdr) == 27, hdr
    assert hdr[:26][-2:] == ["已出貨數量合計", "報價品項數量合計"], "既有欄號不可動，新欄附在最右"
    rows = {r[0].value: r[26].value for r in ws.iter_rows(min_row=4) if r[0].value and str(r[0].value).startswith("Q7-")}
    assert rows == {"Q7-NEW": "25%", "Q7-CUSTOM": "7.5%", "Q7-OLD": "稅前 10%"}
    assert ws.cell(row=2, column=27).value == "管銷"


def test_excel_total_row_still_aligns_with_the_new_column(client):
    _insert_case("Q7-T", deal_tag="已結案", settlement={"status": "finalized", "summary": BASE})
    ws = _sheet()
    last = [r for r in ws.iter_rows(min_row=4) if r[0].value == "合計"]
    assert last and len(last[0]) >= 27 and last[0][26].value in ("", None)

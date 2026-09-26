"""獎金分潤三項（CORE-SPEC「使用者裁示」獎金分潤：通知／送交出納／財務報表）＋ U4 扣繳與補充保費。

INTEGRATION-POINTS：IP-8 `bonus.payouts`（M07 → M05 出納）、IP-9 `expense.entries`（M07 → M08 報表）、
U4 使用 IP-7 L1 `helpers.legal_params`（撥付日選版；題目以 monkeypatch `lp.load_versions` 餵合成的版本）。

① 通知：送審 ⇒ 輪到的簽核人＋代理人；換人 ⇒ 下一位；核准進待發放 ⇒ 出納；名單成員不因在名單上收到
② 出納頁：待發放清單、標記已發放＝獎金那一支 API、執行紀錄與 Excel；財務看不到獎金
③ 報表：以發放日列支出（其他支出／獎金分潤）；案件頁相關傳票看得到獎金傳票
④ U4：純函式邊界；撥付時自動計算並寫進傳票；沒有投保金額 ⇒ 拒絕、狀態不變
⑤ 反向控制：拿掉 M07 的提供者（出納、報表照常，只少獎金並明說）；拿掉會計模組（照算、照發、明說沒傳票）；
   撥付日沒有適用的法規參數版本、版本缺倍數欄位（拒絕撥付並說明，不以 0 或預設值代替）
"""
import io

import pytest

from core import registry
from modules.payroll import bonus_deductions as bd
from helpers import legal_params as lp
from modules.payroll.tests._bonus_insure import insure_all
from modules.payroll.tests.test_bonus_case_api_2026_09_24 import (  # noqa: F401
    people, _seed_case, _create, _members_spec, _auth, _set_flow, _delegate)

YEAR_PARAMS_VERSION = {   # R1 的一版的形狀（helpers.legal_params，IP-7）——合成，不讀產品的預設值
    "version": "2026", "effectiveFrom": "2000-01-01",
    "resident": {"50": {"tax_rate": 0.05, "tax_threshold": 90501}},
    "nhi": {"rate": 0.0211, "max_single_payment": 10000000, "bonus_insured_multiple": 4},
}
PARAMS = bd.params_from_legal_version(YEAR_PARAMS_VERSION)


def _db():
    import db
    return db.get_db()


def _q(sql, args=()):
    conn = _db()
    try:
        return [dict(r) for r in conn.execute(sql, args)]
    finally:
        conn.close()


@pytest.fixture(autouse=True)
def _legal(monkeypatch):
    """每一題都用合成的版本（IP-7：lp.load_versions() → rules_for_date）；要測別的情境時再覆寫。"""
    monkeypatch.setattr(lp, "load_versions", lambda: [YEAR_PARAMS_VERSION])


def _with_legal(monkeypatch, version=YEAR_PARAMS_VERSION):
    monkeypatch.setattr(lp, "load_versions", lambda: [version])


def _drop(monkeypatch, *caps):
    """拿掉提供者：legacy 登記與已載入模組的 ModuleSpec.providers 兩處都要處理（M07 搬進 modules/ 後在後者）。"""
    monkeypatch.setattr(registry, "_LEGACY_PROVIDERS",
                        {k: v for k, v in registry._LEGACY_PROVIDERS.items() if k[0] not in caps})
    orig = registry.providers
    monkeypatch.setattr(registry, "providers", lambda cap: {} if cap in caps else orig(cap))
    for c in caps:
        assert not registry.providers(c), c


def _to_payout(client, people, no, net=2000000):
    _seed_case(no, net=net)
    assert _create(client, people["bc_sa"], no, members=_members_spec()).status_code == 200
    assert client.post("/api/bonus/cases/%s/submit" % no, headers=_auth(people["bc_sa"])).status_code == 200
    r = client.post("/api/bonus/cases/%s/approve" % no, headers=_auth(people["bc_sa2"]))
    assert r.status_code == 200 and r.json()["status"] == "待發放", r.text
    return r


def _insure_all(client, people, amount=45800, **override):
    for u in ("bc_s1", "bc_p1", "bc_p2", "bc_a1", "bc_a2", "bc_a3"):
        r = client.put("/api/bonus/insurance/" + u, headers=_auth(people["bc_sa"]),
                       json={"insuredAmount": override.get(u, amount)})
        assert r.status_code == 200, r.text


@pytest.fixture
def mails(monkeypatch):
    from helpers import email_notify
    sent = []
    monkeypatch.setattr(email_notify, "notify_bonus_submitted",
                        lambda no, cust, who: sent.append(("submitted", no, list(who))))
    monkeypatch.setattr(email_notify, "notify_bonus_payout_ready",
                        lambda no, cust, who: sent.append(("payout_ready", no, list(who))))
    return sent


MEMBERS = {"bc_s1", "bc_p1", "bc_p2", "bc_a1", "bc_a2", "bc_a3"}


# ── ① 通知 ───────────────────────────────────────────────────────────────────

def test_notify_approver_with_delegate_then_cashier_never_members(client, people, mails):
    _set_flow(["bc_sa2"])
    _delegate("bc_sa2", "bc_sa")        # 代理人（W1：代理人也必須是最高管理者）
    _seed_case("MQ-BP-N1", net=100000)
    assert _create(client, people["bc_sa"], "MQ-BP-N1", members=_members_spec()).status_code == 200
    assert client.post("/api/bonus/cases/MQ-BP-N1/submit", headers=_auth(people["bc_sa"])).status_code == 200
    assert mails == [("submitted", "MQ-BP-N1", ["bc_sa2", "bc_sa"])]
    r = client.post("/api/bonus/cases/MQ-BP-N1/approve", headers=_auth(people["bc_sa2"]))
    assert r.json()["status"] == "待發放", r.text
    kind, no, who = mails[-1]
    assert (kind, no) == ("payout_ready", "MQ-BP-N1") and "bc_cash" in who
    for _k, _n, who in mails:
        assert not MEMBERS & set(who), who          # 名單成員不收到


def test_notify_next_approver_in_same_tier(client, people, mails):
    _set_flow(["bc_sa", "bc_sa2"])       # 同一層兩位，依序簽
    _seed_case("MQ-BP-N2", net=100000)
    assert _create(client, people["bc_sa"], "MQ-BP-N2", members=_members_spec()).status_code == 200
    client.post("/api/bonus/cases/MQ-BP-N2/submit", headers=_auth(people["bc_sa"]))
    assert mails[-1] == ("submitted", "MQ-BP-N2", ["bc_sa"])
    r = client.post("/api/bonus/cases/MQ-BP-N2/approve", headers=_auth(people["bc_sa"]))
    assert r.json()["status"] == "待審核"
    assert mails[-1] == ("submitted", "MQ-BP-N2", ["bc_sa2"])


def test_no_chain_notifies_superadmins_except_requester(client, people, mails):
    _seed_case("MQ-BP-N3", net=100000)
    assert _create(client, people["bc_sa"], "MQ-BP-N3", members=_members_spec()).status_code == 200
    client.post("/api/bonus/cases/MQ-BP-N3/submit", headers=_auth(people["bc_sa"]))
    who = mails[-1][2]
    assert "bc_sa2" in who and "bc_sa" not in who


def test_notification_events_are_individually_mutable():
    """「可在通知設定個別關閉」：後端事件清單＋使用者頁的偏好清單（那一份是寫死的副本）都要有。"""
    import pathlib
    from helpers.notification_prefs import EVENT_KEYS
    assert {"bonus_submitted", "bonus_payout_ready"} <= set(EVENT_KEYS)
    # 2026-09-26：使用者頁的退訂清單改由信件類型登記表 API 產生（不再寫死副本）
    from helpers import mail_types
    from core import source_tree
    for k in ("bonus_submitted", "bonus_payout_ready"):
        assert mail_types.get(k) is not None, k
    users = source_tree.page_file("users.html").read_text(encoding="utf-8")   # 頁面位置一律經 page_file（階段 C）
    assert "/api/mail-types/receivable?user_id=" in users


def test_mail_body_has_no_amount(monkeypatch):
    from helpers import email_notify
    got = []
    monkeypatch.setattr(email_notify, "_lookup_emails", lambda users, key: ["x@example.com"])
    monkeypatch.setattr(email_notify, "_async_send", lambda to, subj, html: got.append(subj + html))
    email_notify.notify_bonus_submitted("MQ-X", "客戶", ["u"])
    email_notify.notify_bonus_payout_ready("MQ-X", "客戶", ["u"])
    assert len(got) == 2 and not any("NT$" in g or "金額：" in g for g in got)


# ── ② 出納頁（IP-8）─────────────────────────────────────────────────────────

def test_cashier_queue_mark_paid_is_the_same_action_and_history(client, people, make_user):
    insure_all()
    from modules.payroll.tests.test_bonus_case_api_2026_09_24 import _login
    _to_payout(client, people, "MQ-BP-C1")
    q = client.get("/api/cashier/bonus-queue", headers=_auth(people["bc_cash"])).json()
    row = next(i for i in q["items"] if i["quoteNo"] == "MQ-BP-C1")
    assert q["available"] and q["canMarkPaid"] and row["total"] > 0 and row["people"] == 6
    # 出納頁的按鈕打的是獎金那一支（前端綁定見 test_cashier_page_binds_bonus_mark_paid）
    r = client.post("/api/bonus/cases/MQ-BP-C1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 200, r.text
    q = client.get("/api/cashier/bonus-queue", headers=_auth(people["bc_cash"])).json()
    assert "MQ-BP-C1" not in [i["quoteNo"] for i in q["items"]]
    today = _q("SELECT substr(paid_at,1,10) AS d FROM bonus_case_awards WHERE quote_no='MQ-BP-C1'")[0]["d"]
    h = client.get("/api/cashier/execution-history?start=%s&end=%s" % (today, today),
                   headers=_auth(people["bc_cash"])).json()
    b = next(x for x in h["bonusPaid"] if x["quoteNo"] == "MQ-BP-C1")
    assert b["total"] == row["total"] and h["bonusPaidTotal"] >= row["total"]
    assert b["withholding"] is not None and b["nhiPremium"] is not None and b["net"] is not None
    x = client.get("/api/cashier/export?start=%s&end=%s" % (today, today), headers=_auth(people["bc_cash"]))
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["獎金發放明細"]
    assert "MQ-BP-C1" in [c.value for c in ws["A"]]
    # 財務（finance）：出納頁其他頁籤看得到，獎金看不到
    name, pw = make_user(username="bp_fin", role="engineer", modules=["finance"])
    fin = _login(client, name, pw)
    q = client.get("/api/cashier/bonus-queue", headers=_auth(fin)).json()
    assert q == {"available": False, "visible": False, "notice": "", "items": []}
    h = client.get("/api/cashier/execution-history?start=%s&end=%s" % (today, today), headers=_auth(fin)).json()
    assert h["bonusPaid"] == [] and h["bonusVisible"] is False
    wb = openpyxl.load_workbook(io.BytesIO(client.get(
        "/api/cashier/export?start=%s&end=%s" % (today, today), headers=_auth(fin)).content))
    assert "獎金發放明細" not in wb.sheetnames


def test_cashier_page_binds_bonus_mark_paid():
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[4] / "frontend"
    js = (root / "js" / "cashier.js").read_text(encoding="utf-8")
    from core import source_tree
    html = source_tree.page_file("cashier.html").read_text(encoding="utf-8")   # 頁面位置一律經 page_file（階段 C）
    assert "'/api/bonus/cases/' + encodeURIComponent(this.bonusPay.quoteNo) + '/mark-paid'" in js
    assert "/api/cashier/bonus-queue" in js
    assert 'data-testid="cashier-bonus-tab"' in html and "bonusQueue.notice" in html


def test_reverse_without_payroll_cashier_still_works_and_says_so(client, people, monkeypatch):
    insure_all()
    _to_payout(client, people, "MQ-BP-C2")
    _drop(monkeypatch, "bonus.payouts")
    q = client.get("/api/cashier/bonus-queue", headers=_auth(people["bc_cash"]))
    assert q.status_code == 200
    assert q.json() == {"available": False, "visible": True, "notice": "薪資獎金模組未安裝：出納頁不顯示獎金分潤",
                        "items": []}
    h = client.get("/api/cashier/execution-history", headers=_auth(people["bc_cash"]))
    assert h.status_code == 200 and h.json()["bonusNotice"] and h.json()["bonusPaid"] == []
    # 待付款是外包工班的承攬商匯款（IP-14），與薪資獎金無關；外包工班不在的包裡它回 404 ⇒ 這裡看與兩者都無關的應收
    assert client.get("/api/cashier/receivable-queue", headers=_auth(people["bc_cash"])).status_code == 200
    assert client.get("/api/cashier/export", headers=_auth(people["bc_cash"])).status_code == 200


# ── ③ 報表（IP-9）與案件頁相關傳票 ──────────────────────────────────────────

def _year_other(client, tok, year):
    r = client.get("/api/reports/expenses-monthly?year=%s" % year, headers=_auth(tok))
    assert r.status_code == 200, r.text
    return r.json()["expenses"]


def test_report_counts_bonus_on_paid_date(client, people):
    insure_all()
    _to_payout(client, people, "MQ-BP-R1")
    total = _q("SELECT SUM(l.amount) AS s FROM bonus_case_award_lines l JOIN bonus_case_awards a"
               " ON a.id=l.award_id WHERE a.quote_no='MQ-BP-R1'")[0]["s"]
    before = _year_other(client, people["bc_sa"], 2026)
    assert not [e for e in before["details"]["other"] if e["quoteNo"] == "MQ-BP-R1"]   # 待發放不算
    client.post("/api/bonus/cases/MQ-BP-R1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    paid = _q("SELECT paid_at FROM bonus_case_awards WHERE quote_no='MQ-BP-R1'")[0]["paid_at"][:10]
    exp = _year_other(client, people["bc_sa"], int(paid[:4]))
    rows = [e for e in exp["details"]["other"] if e["quoteNo"] == "MQ-BP-R1"]
    assert rows == [dict(rows[0], date=paid, amount=total, category="獎金分潤")]
    month = next(m for m in exp["monthly"] if m["month"] == paid[:7])
    month_before = next(m for m in before["monthly"] if m["month"] == paid[:7])
    assert month["other"] - month_before["other"] == total


def test_reverse_without_payroll_report_still_works(client, people, monkeypatch):
    insure_all()
    _to_payout(client, people, "MQ-BP-R2")
    client.post("/api/bonus/cases/MQ-BP-R2/mark-paid", headers=_auth(people["bc_cash"]), json={})
    paid = _q("SELECT paid_at FROM bonus_case_awards WHERE quote_no='MQ-BP-R2'")[0]["paid_at"][:4]
    assert [e for e in _year_other(client, people["bc_sa"], paid)["details"]["other"] if e["quoteNo"] == "MQ-BP-R2"]
    _drop(monkeypatch, "expense.entries")
    exp = _year_other(client, people["bc_sa"], paid)
    assert not [e for e in exp["details"]["other"] if e["quoteNo"] == "MQ-BP-R2"]


def test_case_page_related_vouchers_show_bonus_vouchers(client, people):
    insure_all()
    _to_payout(client, people, "MQ-BP-V1")
    client.post("/api/bonus/cases/MQ-BP-V1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    ids = {r["accrual_voucher_id"] for r in _q("SELECT * FROM bonus_case_awards WHERE quote_no='MQ-BP-V1'")} | \
        {r["payment_voucher_id"] for r in _q("SELECT * FROM bonus_case_awards WHERE quote_no='MQ-BP-V1'")}
    assert all(ids) and len(ids) == 2
    r = client.get("/api/vouchers/by-case/MQ-BP-V1", headers=_auth(people["bc_sa"]))
    assert r.status_code == 200, r.text
    assert {v["id"] for v in r.json()["vouchers"]} == ids


# ── ④ U4 純函式 ──────────────────────────────────────────────────────────────

def _one(amount, insured=45800, before=0, params=PARAMS):
    return bd.compute_bonus_deductions([{"username": "u", "amount": amount}], params=params,
                                       insured={"u": insured}, ytd_before={"u": before})["lines"][0]


def test_withholding_threshold_boundary_and_floor():
    assert _one(90500)["withholding"] == 0
    assert _one(90501)["withholding"] == 4525          # 4525.05 捨去
    assert _one(100019)["withholding"] == 5000         # 5000.95 捨去


def test_same_person_in_two_categories_is_merged_before_threshold():
    r = bd.compute_bonus_deductions([{"username": "u", "amount": 50000}, {"username": "u", "amount": 50000}],
                                    params=PARAMS, insured={"u": 45800}, ytd_before={})
    assert len(r["lines"]) == 1 and r["lines"][0]["gross"] == 100000 and r["lines"][0]["withholding"] == 5000


def test_nhi_only_the_part_over_four_times_insured():
    # 門檻 4 × 20,000 = 80,000
    assert _one(70000, insured=20000)["nhiBase"] == 0
    assert _one(100000, insured=20000)["nhiBase"] == 20000
    assert _one(100000, insured=20000)["nhiPremium"] == 422
    assert _one(30000, insured=20000, before=60000)["nhiBase"] == 10000     # 累計跨過門檻
    assert _one(30000, insured=20000, before=90000)["nhiBase"] == 30000     # 已超過 ⇒ 全額
    assert _one(20000000, insured=20000, before=90000)["nhiBase"] == 10000000   # 單次上限


def test_nhi_rounds_half_up():
    p = dict(PARAMS, nhi_rate=0.025)            # 20 × 0.025 = 0.5
    assert _one(80020, insured=20000, params=p)["nhiPremium"] == 1


def test_nhi_35000_is_739_not_banker_rounding():
    """D 稽核 R1：健保署規定四捨五入；Python round() 是銀行家捨入（738.5 → 738）。
    捨入集中在 bonus_deductions.nhi_premium_of 一處；L1 共用四捨五入函式合回後只換那一處。"""
    assert bd.nhi_premium_of(35000, PARAMS) == 739
    assert round(35000 * 0.0211) == 738                     # 對照：round() 會少 1 元


def test_missing_insured_is_not_zero():
    r = bd.compute_bonus_deductions([{"username": "u", "amount": 1000}], params=PARAMS,
                                    insured={}, ytd_before={})
    assert r["missing"] == ["u"] and r["lines"][0]["nhiPremium"] is None and r["lines"][0]["net"] is None


def test_params_are_not_guessed():
    v = {k: dict(x) if isinstance(x, dict) else x for k, x in YEAR_PARAMS_VERSION.items()}
    v["nhi"] = {"rate": 0.0211, "max_single_payment": 10000000}      # R1 目前沒有倍數欄位
    with pytest.raises(bd.DeductionParamsError, match="nhi_bonus_multiple"):
        bd.params_from_legal_version(v)


# ── ④ U4 撥付整合 ────────────────────────────────────────────────────────────

def _payment_lines(no):
    return _q("SELECT l.account_code, l.debit, l.credit, l.summary, l.source_type, l.source_key"
              " FROM voucher_lines l JOIN bonus_case_awards a ON a.payment_voucher_id = l.voucher_id"
              " WHERE a.quote_no = ? ORDER BY l.line_no", (no,))


def test_mark_paid_refuses_without_insured_amount(client, people, monkeypatch):
    _with_legal(monkeypatch)
    _to_payout(client, people, "MQ-BP-U1")
    d = client.get("/api/bonus/cases/MQ-BP-U1", headers=_auth(people["bc_cash"])).json()
    assert d["deductions"]["missing"] and d["deductionNotice"]
    r = client.post("/api/bonus/cases/MQ-BP-U1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 409 and "投保金額" in r.json()["detail"]
    assert _q("SELECT status FROM bonus_case_awards WHERE quote_no='MQ-BP-U1'")[0]["status"] == "待發放"


def test_mark_paid_computes_and_books_deductions(client, people, monkeypatch):
    _with_legal(monkeypatch)
    _insure_all(client, people, bc_s1=20000)
    _to_payout(client, people, "MQ-BP-U2")          # 淨利 2,000,000 ⇒ 業務 100,000
    d = client.get("/api/bonus/cases/MQ-BP-U2", headers=_auth(people["bc_cash"])).json()
    s1 = next(x for x in d["deductions"]["lines"] if x["username"] == "bc_s1")
    assert (s1["gross"], s1["withholding"], s1["nhiBase"], s1["nhiPremium"]) == (100000, 5000, 20000, 422)
    r = client.post("/api/bonus/cases/MQ-BP-U2/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 200, r.text
    tot = r.json()["deductions"]["totals"]
    assert tot["withholding"] == 5000 and tot["nhiPremium"] == 422
    lines = _payment_lines("MQ-BP-U2")
    credit = {l["summary"].split(" ")[-1]: l["credit"] for l in lines if l["credit"]}
    assert credit["實發"] == tot["gross"] - 5000 - 422
    assert credit["代扣稅款"] == 5000 and credit["代收二代健保補充保費"] == 422
    assert sum(l["debit"] for l in lines) == sum(l["credit"] for l in lines) == tot["gross"]
    assert {(l["source_type"], l["source_key"]) for l in lines} == {("case", "MQ-BP-U2")}
    # 快照：已發放後仍顯示發放當時的計算；全年累計進下一張的「累計前」
    d = client.get("/api/bonus/cases/MQ-BP-U2", headers=_auth(people["bc_sa"])).json()
    assert d["deductions"]["totals"] == tot
    assert d["deductions"]["version"] == "2026" and d["deductions"]["params"] == {
        k: PARAMS[k] for k in bd.PARAM_KEYS}                      # 版本＋參數快照存在單據上
    assert d["deductions"]["rules"] == YEAR_PARAMS_VERSION         # 整份法規參數快照（IP-7 單據凍結）
    ins = client.get("/api/bonus/insurance", headers=_auth(people["bc_sa"])).json()
    assert next(i for i in ins["items"] if i["username"] == "bc_s1")["ytdMotrix"] == 100000


def test_ytd_external_counts_toward_cap(client, people, monkeypatch):
    _with_legal(monkeypatch)
    _insure_all(client, people, bc_s1=20000)
    from datetime import date
    r = client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_sa"]),
                   json={"year": date.today().year, "ytdExternal": 80000})
    assert r.status_code == 200, r.text
    _to_payout(client, people, "MQ-BP-U3", net=200000)       # 業務 10,000，全部超過門檻
    d = client.get("/api/bonus/cases/MQ-BP-U3", headers=_auth(people["bc_sa"])).json()
    s1 = next(x for x in d["deductions"]["lines"] if x["username"] == "bc_s1")
    assert (s1["ytdBefore"], s1["nhiBase"]) == (80000, 10000)


def test_insurance_endpoints_validate_and_are_superadmin_only(client, people):
    for bad in (0, -1, "12.5", True):
        assert client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_sa"]),
                          json={"insuredAmount": bad}).status_code == 400, bad
    assert client.put("/api/bonus/insurance/nobody", headers=_auth(people["bc_sa"]),
                      json={"insuredAmount": 1}).status_code == 404
    assert client.get("/api/bonus/insurance", headers=_auth(people["bc_cash"])).status_code == 403
    assert client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_cash"]),
                      json={"insuredAmount": 1}).status_code == 403
    # 空＝清除（不是 0）
    client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_sa"]), json={"insuredAmount": 30000})
    client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_sa"]), json={"insuredAmount": ""})
    it = next(i for i in client.get("/api/bonus/insurance", headers=_auth(people["bc_sa"])).json()["items"]
              if i["username"] == "bc_s1")
    assert it["insuredAmount"] is None


def test_bonus_page_binds_insurance_and_deductions():
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[4] / "frontend"
    js = (root / "js" / "bonus.js").read_text(encoding="utf-8")
    from core import source_tree
    html = source_tree.page_file("bonus.html").read_text(encoding="utf-8")   # 頁面位置一律經 page_file（階段 C）
    assert "'/api/bonus/insurance/' + encodeURIComponent(it.username)" in js
    assert 'data-testid="bn-deductions"' in html and "detail.deductionNotice" in html
    assert "'nhi'" in html


# ── ⑤ 反向控制：法規參數不在、會計模組不在 ───────────────────────────────────

def test_no_applicable_version_refuses_payout(client, people, monkeypatch):
    """撥付日沒有適用的法規參數版本 ⇒ 拒絕撥付並說明，狀態不變、沒有支出傳票（不以 0 或預設值代替）。"""
    insure_all()
    _to_payout(client, people, "MQ-BP-L1")
    _with_legal(monkeypatch, dict(YEAR_PARAMS_VERSION, effectiveFrom="2099-01-01"))
    d = client.get("/api/bonus/cases/MQ-BP-L1", headers=_auth(people["bc_cash"])).json()
    assert d["deductions"] is None and "無法計算扣繳與補充保費" in d["deductionNotice"]
    r = client.post("/api/bonus/cases/MQ-BP-L1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 409 and "無法計算扣繳與補充保費" in r.json()["detail"], r.text
    a = _q("SELECT status, payment_voucher_id FROM bonus_case_awards WHERE quote_no='MQ-BP-L1'")[0]
    assert a == {"status": "待發放", "payment_voucher_id": 0}


def test_version_without_multiple_refuses_payout(client, people, monkeypatch):
    insure_all()
    _to_payout(client, people, "MQ-BP-L2")
    v = dict(YEAR_PARAMS_VERSION, nhi={"rate": 0.0211, "max_single_payment": 10000000})
    _with_legal(monkeypatch, v)
    r = client.post("/api/bonus/cases/MQ-BP-L2/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 409 and "nhi_bonus_multiple" in r.json()["detail"], r.text


def test_real_legal_params_default_version_is_usable(client, monkeypatch):
    """正對照（不 monkeypatch）：產品內建的版本接得上本項——有倍數欄位、轉得成 params。"""
    monkeypatch.undo()
    p, rules, why = bd.legal_params_for(lp.today().isoformat())
    assert why == "" and p["nhi_bonus_multiple"] >= 1 and rules["version"], (p, why)


def test_reverse_without_accounting_still_computes_and_pays(client, people, monkeypatch):
    _with_legal(monkeypatch)
    _insure_all(client, people, bc_s1=20000)
    _to_payout(client, people, "MQ-BP-A1")
    _drop(monkeypatch, "voucher.draft", "voucher.account_check", "accounting.settings",
          "voucher.void_draft", "voucher.status")
    r = client.post("/api/bonus/cases/MQ-BP-A1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 200, r.text
    assert "會計模組未安裝" in r.json()["notice"]
    assert r.json()["deductions"]["totals"]["nhiPremium"] == 422
    assert _q("SELECT payment_voucher_id FROM bonus_case_awards WHERE quote_no='MQ-BP-A1'")[0][
        "payment_voucher_id"] == 0
    # 出納頁、報表照常
    assert client.get("/api/cashier/bonus-queue", headers=_auth(people["bc_cash"])).status_code == 200
    paid = _q("SELECT paid_at FROM bonus_case_awards WHERE quote_no='MQ-BP-A1'")[0]["paid_at"][:4]
    assert [e for e in _year_other(client, people["bc_sa"], paid)["details"]["other"] if e["quoteNo"] == "MQ-BP-A1"]


# ── 稽核 D（AUDIT-D-A-bonus-U4）建議 ─────────────────────────────────────────────

def test_ytd_counts_only_paid_bonuses(client, people):
    """A-S1：全年累計只算「已發放」的獎金；待發放（還沒付）的不算。"""
    insure_all()
    _to_payout(client, people, "MQ-BP-Y1")                                  # 待發放
    assert bd.ytd_in_motrix(bd_conn(), date_year()).get("bc_s1", 0) == 0
    client.post("/api/bonus/cases/MQ-BP-Y1/mark-paid", headers=_auth(people["bc_cash"]), json={})
    paid = _q("SELECT SUM(l.amount) AS s FROM bonus_case_award_lines l JOIN bonus_case_awards a"
              " ON a.id=l.award_id WHERE a.quote_no='MQ-BP-Y1' AND l.username='bc_s1'")[0]["s"]
    assert bd.ytd_in_motrix(bd_conn(), date_year()).get("bc_s1") == paid


def bd_conn():
    import db
    return db.get_db()


def date_year():
    from datetime import date
    return date.today().year


def test_corrupt_profiles_are_never_overwritten(client, people):
    """A-S2：投保金額設定壞掉 ⇒ 讀、寫、撥付都拒絕；原本的內容一個字都不動。"""
    import db
    conn = db.get_db()
    try:
        conn.execute("INSERT INTO system_settings (key, value_json, updated_at) VALUES (?,?,?) "
                     "ON CONFLICT(key) DO UPDATE SET value_json=excluded.value_json",
                     (bd.PROFILE_KEY, "{壞掉", "t"))
        conn.commit()
    finally:
        conn.close()
    before = _q("SELECT value_json FROM system_settings WHERE key=?", (bd.PROFILE_KEY,))[0]["value_json"]
    r = client.put("/api/bonus/insurance/bc_s1", headers=_auth(people["bc_sa"]), json={"insuredAmount": 30000})
    assert r.status_code == 409 and "損毀" in r.json()["detail"], r.text
    assert client.get("/api/bonus/insurance", headers=_auth(people["bc_sa"])).status_code == 409
    _to_payout(client, people, "MQ-BP-Y2")
    r = client.post("/api/bonus/cases/MQ-BP-Y2/mark-paid", headers=_auth(people["bc_cash"]), json={})
    assert r.status_code == 409 and "損毀" in r.json()["detail"], r.text
    assert _q("SELECT value_json FROM system_settings WHERE key=?", (bd.PROFILE_KEY,))[0]["value_json"] == before

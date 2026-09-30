# -*- coding: utf-8 -*-
"""總帳 C5 · 營業稅 401：由總帳分錄彙總（稅碼＋科目類別）、對帳、警示、稅額結轉草稿（E14）、API 與 Excel。
反向控制：未過帳的傳票不算；發票對不上 ⇒ 不可產生結轉；估計稅額／免稅銷售警示；已過帳的結轉傳票不可重做；留抵與用完留抵的分錄平衡。
"""
import io
import json

import pytest
from openpyxl import load_workbook

import db
from core import registry
from modules.accounting.ledger import engine as E
from modules.accounting.ledger import features as F
from modules.accounting.ledger import roles as ROLES
from modules.accounting.ledger import tax401 as T

_N = [0]


@pytest.fixture
def conn(client):
    c = db.get_db()
    ROLES.ensure_meta(c)
    ROLES.ensure_default_roles(c)
    c.commit()
    yield c
    c.close()


@pytest.fixture
def fake(monkeypatch):
    monkeypatch.setattr(registry, "_LEGACY_PROVIDERS", {})
    monkeypatch.setattr(registry, "_LOADED", {})
    state = {"events": []}
    registry._LEGACY_PROVIDERS[("gl.events", "fake")] = lambda s, e, changed_since="": {"events": list(state["events"])}
    return state


def _ev(kind, date, **kw):
    _N[0] += 1
    key = "T401-%d-%d" % (id(_N), _N[0])
    if kind == "sale":
        lines = [{"role": "AR", "side": "D", "amount": 10500}, {"role": "REV_SALES", "side": "C", "amount": 10000}, {"role": "OUTPUT_TAX", "side": "C", "amount": 500}]
        e = {"source_type": "t_sale", "event_code": "E01", "tax_code": "OUT-5"}
    elif kind == "zero":
        lines = [{"role": "AR", "side": "D", "amount": 3000}, {"role": "REV_SALES", "side": "C", "amount": 3000}]
        e = {"source_type": "t_zero", "event_code": "E01", "tax_code": "OUT-0"}
    elif kind == "exempt":
        lines = [{"role": "AR", "side": "D", "amount": 800}, {"role": "REV_SALES", "side": "C", "amount": 800}]
        e = {"source_type": "t_ex", "event_code": "E01", "tax_code": "OUT-EX"}
    elif kind == "buy":
        lines = [{"role": "COST_PROJECT", "side": "D", "amount": 2000}, {"role": "INPUT_TAX", "side": "D", "amount": 100}, {"role": "AP", "side": "C", "amount": 2100}]
        e = {"source_type": "t_buy", "event_code": "E04", "tax_code": "IN-5"}
    elif kind == "bigbuy":
        lines = [{"role": "COST_PROJECT", "side": "D", "amount": 4000}, {"role": "INPUT_TAX", "side": "D", "amount": 200}, {"role": "AP", "side": "C", "amount": 4200}]
        e = {"source_type": "t_bigbuy", "event_code": "E04", "tax_code": "IN-5"}
    e.update({"source_key": key, "event_date": date, "doc_no": key, "case_no": "", "party": {"key": "12345678", "name": "甲"}, "mode": "snapshot", "lines": lines})
    e.update(kw)
    return e


def _run_and_post(conn, fake, events, start, end, post=True):
    fake["events"] = events
    E.run(conn, start, end, "acc")
    if post:
        conn.execute("UPDATE vouchers_all SET status='已過帳' WHERE kind='auto' AND status='草稿' AND voided_at=''")
    conn.commit()


def test_bimonthly_periods():
    assert T.bimonthly(2178, 1) == ("2178-01-01", "2178-02-28")
    assert T.bimonthly(2176, 1) == ("2176-01-01", "2176-02-29")          # 閏年
    assert T.bimonthly(2178, 6) == ("2178-11-01", "2178-12-31")
    with pytest.raises(T.TaxError):
        T.bimonthly(2178, 7)


def test_summary_from_posted_lines_and_calc(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10"), _ev("zero", "2178-01-15"), _ev("buy", "2178-02-05")], "2178-01-01", "2178-02-28")
    s = T.summarize(conn, 2178, 1)
    by = {r["tax_code"]: r for r in s["rows"]}
    assert (by["OUT-5"]["amount"], by["OUT-5"]["tax"]) == (10000, 500) and by["OUT-0"]["zero_amount"] == 3000
    assert (by["IN-5"]["amount"], by["IN-5"]["tax"]) == (2000, 100)
    assert {k: s["calc"][k] for k in ("101", "106", "107", "108", "110", "111", "112", "25")} == {"101": 500, "106": 500, "107": 100, "108": 0, "110": 100, "111": 400, "112": 0, "25": 13000}
    assert (s["calc"]["21"], s["calc"]["22"], s["calc"]["23"], s["calc"]["44"], s["calc"]["45"]) == (10000, 500, 3000, 2000, 100)     # 官方合計列
    assert s["reconciled"] is True and [c["ok"] for c in s["checks"][:2]] == [True, True]
    assert by["OUT-5"]["field_amt"] == "5" and by["IN-5"]["field_tax"] == "29"


def test_unposted_vouchers_are_not_counted(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10")], "2178-01-01", "2178-02-28", post=False)
    s = T.summarize(conn, 2178, 1)
    assert s["rows"] == [] and s["calc"]["101"] == 0


def test_invoice_comparison_ok_and_mismatch(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10")], "2178-01-01", "2178-02-28")
    ok = T.summarize(conn, 2178, 1, [{"invoiceDate": "2178-01-10", "amountPretax": 10000, "taxAmount": 500}])
    assert ok["reconciled"] is True and [c["ok"] for c in ok["checks"] if c["key"].startswith("invoices")] == [True, True]
    bad = T.summarize(conn, 2178, 1, [{"invoiceDate": "2178-01-10", "amountPretax": 10000, "taxAmount": 501}])
    assert bad["reconciled"] is False
    none = T.summarize(conn, 2178, 1, None)
    assert [c for c in none["checks"] if c["key"] == "invoices"][0]["ok"] is None and none["reconciled"] is True


def test_warnings_for_exempt_and_estimated_tax(conn, fake):
    est = _ev("buy", "2178-01-20")
    est["meta"] = {"tax_estimated": True}
    _run_and_post(conn, fake, [_ev("exempt", "2178-01-12"), est], "2178-01-01", "2178-02-28")
    keys = {w["key"] for w in T.summarize(conn, 2178, 1)["warnings"]}
    assert {"exempt", "estimated"} <= keys


def _settle(conn, y=2178, n=1):
    r = T.generate_settlement(conn, y, n, "acc")
    conn.commit()
    return r


def _vlines(conn, vid):
    return sorted((r["account_code"], r["debit"], r["credit"]) for r in conn.execute("SELECT account_code, debit, credit FROM voucher_lines WHERE voucher_id=?", (vid,)))


def test_settlement_payable_case_and_regenerate(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10"), _ev("buy", "2178-02-05")], "2178-01-01", "2178-02-28")
    r = _settle(conn)
    assert (r["payable"], r["carry_new"]) == (400, 0)
    assert _vlines(conn, r["voucher_id"]) == sorted([("2204", 500, 0), ("1268", 0, 100), ("2194", 0, 400)])
    row = conn.execute("SELECT * FROM gl_tax_settlements WHERE period_start='2178-01-01'").fetchone()
    assert (row["output_tax"], row["input_tax"], row["payable"], row["voucher_id"]) == (500, 100, 400, r["voucher_id"])
    r2 = _settle(conn)                                                       # 重建：舊草稿作廢
    assert r2["voucher_id"] != r["voucher_id"] and conn.execute("SELECT voided_at FROM vouchers_all WHERE id=?", (r["voucher_id"],)).fetchone()[0] != ""
    conn.execute("UPDATE vouchers_all SET status='已過帳' WHERE id=?", (r2["voucher_id"],))
    conn.commit()
    with pytest.raises(T.TaxError):
        T.generate_settlement(conn, 2178, 1, "acc")                           # 已過帳 ⇒ 不重做


def test_settlement_blocked_when_not_reconciled_or_empty(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10")], "2178-01-01", "2178-02-28")
    with pytest.raises(T.TaxError):
        T.generate_settlement(conn, 2178, 1, "acc", [{"invoiceDate": "2178-01-10", "amountPretax": 1, "taxAmount": 1}])
    with pytest.raises(T.TaxError):
        T.generate_settlement(conn, 2178, 3, "acc")                           # 沒有稅額的期別


def test_carry_forward_then_used_up_next_period(conn, fake):
    _run_and_post(conn, fake, [_ev("bigbuy", "2178-03-10")], "2178-03-01", "2178-04-30")           # 只有進項 200
    r1 = _settle(conn, 2178, 2)
    assert (r1["payable"], r1["carry_new"]) == (0, 200)
    assert _vlines(conn, r1["voucher_id"]) == sorted([("1268", 0, 200), ("1269", 200, 0)])
    conn.execute("UPDATE vouchers_all SET status='已過帳' WHERE id=?", (r1["voucher_id"],))
    conn.commit()
    _run_and_post(conn, fake, [_ev("sale", "2178-05-10")], "2178-05-01", "2178-06-30")            # 下期銷項 500
    s = T.summarize(conn, 2178, 3)
    assert s["calc"]["108"] == 200 and s["calc"]["111"] == 300 and s["reconciled"] is True         # 上期留抵被承接（結轉傳票不重複算進 101/107）
    r2 = _settle(conn, 2178, 3)
    assert (r2["payable"], r2["carry_new"]) == (300, 0)
    assert _vlines(conn, r2["voucher_id"]) == sorted([("2204", 500, 0), ("1269", 0, 200), ("2194", 0, 300)])


def test_api_flag_gate_summary_export_and_settlement(client, conn, fake, make_user):
    u, p = make_user(username="t401_admin%d" % id(client), role="superadmin")
    tok = client.post("/api/auth/login", json={"username": u, "password": p}).json()["token"]
    h = {"Authorization": "Bearer " + tok}
    assert client.get("/api/ledger/tax401?year=2178&period=1", headers=h).status_code == 409           # 旗標預設關
    F.set_flag(conn, "tax401", True)
    conn.commit()
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10"), _ev("buy", "2178-02-05")], "2178-01-01", "2178-02-28")
    r = client.get("/api/ledger/tax401?year=2178&period=1", headers=h)
    assert r.status_code == 200 and r.json()["calc"]["111"] == 400
    assert client.get("/api/ledger/tax401?year=2178&period=9", headers=h).status_code == 400
    x = client.get("/api/ledger/tax401/export?year=2178&period=1", headers=h)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert any(c.value == "本期應實繳稅額" for row in ws.iter_rows() for c in row)
    s = client.post("/api/ledger/tax401/settlement", headers=h, json={"year": 2178, "period": 1})
    assert s.status_code == 200 and s.json()["payable"] == 400
    assert client.post("/api/ledger/tax401/settlement", headers=h, json={"year": 2178, "period": 4}).status_code == 400


# ── 官方欄位代號（財政部《營業稅電子資料申報繳稅作業要點》附件六，113/04/12 令）────────────────────────

OFFICIAL = {   # 代號 → 附件六的欄項名稱（逐字抄自官方檔案格式）
    "1": "三聯式發票", "2": "三聯式發票", "5": "收銀機發票(三聯式)及電子發票", "6": "收銀機發票(三聯式)及電子發票", "9": "二聯式收銀機(二聯式)發票", "10": "二聯式收銀機(二聯式)發票",
    "17": "退回及折讓", "18": "退回及折讓", "7": "非經海關出口應附證明文件者", "28": "進貨及費用", "29": "進貨及費用", "30": "固定資產", "31": "固定資產",
    "40": "進貨及費用", "41": "進貨及費用", "101": "本(期)月銷項稅額合計", "107": "得扣抵進項稅額合計", "108": "上期(月)累積留抵稅額", "110": "小計(7+8+9)",
    "111": "本期(月)應實繳稅額(6-10)", "112": "本期(月)申報留抵稅額(10-6)",
}


def test_default_map_uses_only_official_codes_in_the_right_columns():
    by = {(c, k): (a, t, z) for c, k, _side, a, t, z, _n in T.DEFAULT_MAP}
    assert by[("OUT-5", "einvoice")] == ("5", "6", "") and by[("OUT-5", "triplicate")] == ("1", "2", "") and by[("OUT-5", "duplicate")] == ("9", "10", "")
    assert by[("OUT-0", "")] == ("", "", "7") and by[("OUT-ADJ", "")] == ("17", "18", "")          # 7＝零稅率(非經海關)；19 是零稅率退回折讓，不可放進應稅退回折讓
    assert by[("IN-5", "")] == ("28", "29", "") and by[("IN-FA", "")] == ("30", "31", "") and by[("IN-ADJ", "")] == ("40", "41", "")
    used = {x for v in by.values() for x in v if x}
    assert used <= set(T.LINE_NAMES), sorted(used - set(T.LINE_NAMES))                             # 每個用到的代號都有官方名稱
    for code, name in OFFICIAL.items():
        assert name in T.LINE_NAMES[code], (code, T.LINE_NAMES[code])                              # 名稱與官方欄項名稱一致


def test_ensure_map_repairs_early_wrong_seeds_but_keeps_user_edits(conn):
    conn.execute("DELETE FROM gl_tax401_map")
    conn.execute("INSERT INTO gl_tax401_map(tax_code, invoice_kind, side, field_amt, field_tax, field_zero, note) VALUES ('OUT-5','einvoice','OUT','5','6','7','舊')")
    conn.execute("INSERT INTO gl_tax401_map(tax_code, invoice_kind, side, field_amt, field_tax, field_zero, note) VALUES ('OUT-0','einvoice','OUT','','','7','舊')")
    conn.execute("INSERT INTO gl_tax401_map(tax_code, invoice_kind, side, field_amt, field_tax, field_zero, note) VALUES ('OUT-ADJ','','OUT','17','18','19','舊')")
    conn.execute("INSERT INTO gl_tax401_map(tax_code, invoice_kind, side, field_amt, field_tax, field_zero, note) VALUES ('IN-5','','IN','88','89','','使用者自訂')")
    T.ensure_map(conn)
    conn.commit()
    got = {(r["tax_code"], r["invoice_kind"]): (r["field_amt"], r["field_tax"], r["field_zero"]) for r in conn.execute("SELECT * FROM gl_tax401_map")}
    assert got[("OUT-5", "einvoice")] == ("5", "6", "") and ("OUT-0", "einvoice") not in got and got[("OUT-0", "")] == ("", "", "7")
    assert got[("OUT-ADJ", "")] == ("17", "18", "") and got[("IN-5", "")] == ("88", "89", "")      # 使用者改過的列不動


def test_default_invoice_medium_selects_the_row(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10")], "2178-01-01", "2178-02-28")
    conn.execute("INSERT OR REPLACE INTO gl_settings(key, value) VALUES ('default_invoice_medium', 'triplicate')")
    conn.commit()
    row = [r for r in T.summarize(conn, 2178, 1)["rows"] if r["tax_code"] == "OUT-5"][0]
    assert (row["field_amt"], row["field_tax"], row["invoice_kind"]) == ("1", "2", "triplicate")
    conn.execute("DELETE FROM gl_settings WHERE key='default_invoice_medium'")
    conn.commit()
    assert [r for r in T.summarize(conn, 2178, 1)["rows"] if r["tax_code"] == "OUT-5"][0]["field_amt"] == "5"          # 預設＝電子發票


def test_adjustments_reduce_sales_and_output_tax_like_the_paper_form(conn, fake):
    adj = _ev("sale", "2178-01-20", source_type="t_adj")
    adj["tax_code"] = "OUT-ADJ"
    adj["lines"] = [{"role": "SALES_RETURN", "side": "D", "amount": 2000}, {"role": "OUTPUT_TAX", "side": "D", "amount": 100}, {"role": "AR", "side": "C", "amount": 2100}]
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10"), adj], "2178-01-01", "2178-02-28")
    s = T.summarize(conn, 2178, 1)
    assert s["calc"]["101"] == 400 and s["calc"]["21"] == 10000 - 2000 and s["calc"]["22"] == 500 - 100     # 退回及折讓為減項（代號17／18）


def test_summary_and_export_cite_the_official_source_and_list_only_unverified_items(conn, fake):
    _run_and_post(conn, fake, [_ev("sale", "2178-01-10")], "2178-01-01", "2178-02-28")
    s = T.summarize(conn, 2178, 1)
    assert "附件六" in s["official_source"] and "1130001073" in s["official_source"] and s["line_names"]["101"] == "本(期)月銷項稅額合計"
    assert len(s["unverified"]) == len(T.UNVERIFIED) and any("代號 115" in u for u in s["unverified"])
    from modules.accounting.ledger import export as EX
    import io as _io
    from openpyxl import load_workbook
    ws = load_workbook(_io.BytesIO(EX.tax401_workbook(s))).active
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "5：應稅銷售額—收銀機發票(三聯式)及電子發票" in text and "尚未核實" in text and "附件六" in text

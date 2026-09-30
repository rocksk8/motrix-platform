# -*- coding: utf-8 -*-
"""總帳 B5 · 年度結轉與決算（proposal-gl/03-periods-close.md §5）。

守門：結轉前後資產負債表不變、損益科目歸零；決算要「1～11 期已結帳＋結轉傳票已過帳＋四表對帳通過」才成立；決算後快照凍結（帳怎麼動匯出不變）；
年度重開把期間、結轉傳票、後續期間狀態都復原並留稽核軌跡。反向控制：每個前置缺一個就必須被擋下並說明原因。
"""
import io
import sqlite3

import pytest

import db
from modules.accounting.ledger import closing as CL
from modules.accounting.ledger import export as EX
from modules.accounting.ledger import periods as P
from modules.accounting.ledger import roles as ROLES
from modules.accounting.ledger import statements as ST

_SEQ = [0]
_YEAR = [1919]           # 從 1920 起往上用（create_year 的合理範圍 1911～2200；每題一個新年度）


@pytest.fixture
def conn(client):
    c = db.get_db()
    ROLES.ensure_meta(c)
    ROLES.ensure_default_roles(c)
    c.commit()
    yield c
    c.close()


def _login(client, make_user, username, role="superadmin", modules=("finance",)):
    u, p = make_user(username=username, role=role, modules=list(modules))
    r = client.post("/api/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["token"]}


def _v(conn, date, lines, status="已過帳", kind="manual"):
    _SEQ[0] += 1
    cur = conn.execute(
        "INSERT INTO vouchers_all(voucher_no, voucher_date, category, summary, status, created_by, created_at, updated_at, kind)"
        " VALUES (?,?, '轉','s',?, 't','n','n', ?)", ("%s-%03d" % (date.replace("-", ""), _SEQ[0]), date,
                                                        "已核准" if status == "已過帳" else status, kind))
    for i, (code, d, c) in enumerate(lines, 1):
        conn.execute("INSERT INTO voucher_lines(voucher_id,line_no,account_code,debit,credit) VALUES (?,?,?,?,?)", (cur.lastrowid, i, code, d, c))
    if status == "已過帳":
        conn.execute("UPDATE vouchers_all SET status='已過帳' WHERE id=?", (cur.lastrowid,))
    conn.commit()
    return cur.lastrowid


def _new_year(conn, book=True, loss=False):
    """新的一個會計年度（避免題間干擾）＋一月的帳：資本、進貨、銷售、成本、租金 ⇒ 淨利 5,000（loss=True ⇒ 租金 8,000 ⇒ 淨損 2,000）。"""
    _YEAR[0] += 1
    y = _YEAR[0]
    P.create_year(conn, y, "t")
    conn.commit()
    if book:
        _v(conn, "%d-01-05" % y, [("1113", 100000, 0), ("3111", 0, 100000)])
        _v(conn, "%d-01-08" % y, [("1231", 6000, 0), ("2171", 0, 6000)])
        _v(conn, "%d-01-10" % y, [("1191", 10500, 0), ("4111", 0, 10000), ("2204", 0, 500)])
        _v(conn, "%d-01-15" % y, [("5111", 4000, 0), ("1231", 0, 4000)])
        rent = 8000 if loss else 1000
        _v(conn, "%d-01-20" % y, [("6112", rent, 0), ("1113", 0, rent)])
    return y


def _close_periods(conn, y, upto=11):
    for n in range(1, upto + 1):
        pid = conn.execute("SELECT id FROM gl_periods WHERE year=? AND period_no=?", (y, n)).fetchone()[0]
        P.close_period(conn, pid, "acc", accept_warnings=True)
    conn.commit()


def _post_closing(conn, y):
    for v in CL.closing_vouchers(conn, y):
        conn.execute("UPDATE vouchers_all SET status='已過帳' WHERE id=?", (v["id"],))
    conn.commit()


def _full_year(conn, **kw):
    y = _new_year(conn, **kw)
    _close_periods(conn, y)
    CL.generate(conn, y, "acc")
    conn.commit()
    _post_closing(conn, y)
    return y


# ── 預覽與產生 ────────────────────────────────────────────────────────────

def test_preview_lines_are_balanced_and_state_prerequisites(conn):
    y = _new_year(conn)
    pv = CL.preview(conn, y)
    assert pv["net_income"] == 5000 and pv["pl_accounts"] == 3 and not pv["problems"]
    v1, v2 = pv["vouchers"]
    assert sum(l["debit"] for l in v1["lines"]) == sum(l["credit"] for l in v1["lines"])
    by = {l["account_code"]: l for l in v1["lines"]}
    assert by["4111"]["debit"] == 10000 and by["5111"]["credit"] == 4000 and by["6112"]["credit"] == 1000 and by["3353"]["credit"] == 5000
    assert [(l["account_code"], l["debit"], l["credit"]) for l in v2["lines"]] == [("3353", 5000, 0), ("3351", 0, 5000)]
    assert any("還沒結帳" in u for u in pv["prerequisites_unmet"])            # 1～11 期都還開著 ⇒ 前置未滿足要明說


def test_preview_for_a_loss_year_reverses_the_direction(conn):
    y = _new_year(conn, loss=True)
    pv = CL.preview(conn, y)
    assert pv["net_income"] == -2000
    by = {l["account_code"]: l for l in pv["vouchers"][0]["lines"]}
    assert by["3353"]["debit"] == 2000 and by["3353"]["credit"] == 0
    assert [(l["account_code"], l["debit"], l["credit"]) for l in pv["vouchers"][1]["lines"]] == [("3353", 0, 2000), ("3351", 2000, 0)]


def test_generate_refuses_until_periods_are_closed_then_makes_two_drafts_once(conn):
    y = _new_year(conn)
    with pytest.raises(CL.ClosingError) as ei:
        CL.generate(conn, y, "acc")
    assert "還沒結帳" in str(ei.value)
    _close_periods(conn, y)
    made = CL.generate(conn, y, "acc")
    conn.commit()
    assert len(made) == 2
    rows = conn.execute("SELECT kind, origin, status, voucher_date FROM vouchers_all WHERE id IN (?,?)", (made[0]["id"], made[1]["id"])).fetchall()
    assert all(tuple(r)[:3] == ("closing", "year_closing", "草稿") and r["voucher_date"] == "%d-12-31" % y for r in rows)
    with pytest.raises(CL.ClosingError) as ei:
        CL.generate(conn, y, "acc")
    assert "重新產生" in str(ei.value)
    again = CL.generate(conn, y, "acc", regenerate=True)                       # 只作廢草稿後重產
    conn.commit()
    assert len(again) == 2 and len(CL.closing_vouchers(conn, y)) == 2


def test_generate_refuses_after_posting_and_when_no_pl(conn):
    y = _full_year(conn)
    with pytest.raises(CL.ClosingError) as ei:
        CL.generate(conn, y, "acc", regenerate=True)
    assert "年度重開" in str(ei.value)
    y2 = _new_year(conn, book=False)
    _close_periods(conn, y2)
    with pytest.raises(CL.ClosingError) as ei:
        CL.generate(conn, y2, "acc")
    assert "不需要結轉" in str(ei.value)


def test_generate_reports_a_missing_role_instead_of_guessing(conn):
    y = _new_year(conn)
    _close_periods(conn, y)
    # ensure_default_roles 只補「完全沒有這個角色」的預設；把唯一一列改成未來才生效 ⇒ 年度末日解析不到科目 ⇒ 必須說明缺哪個角色
    conn.execute("DELETE FROM gl_account_roles WHERE role='PL_SUMMARY'")
    conn.execute("INSERT INTO gl_account_roles(role,scope_type,scope_key,account_code,effective_from) VALUES ('PL_SUMMARY','','','3353','2999-01-01')")
    conn.commit()
    with pytest.raises(CL.ClosingError) as ei:
        CL.generate(conn, y, "acc")
    assert "PL_SUMMARY" in str(ei.value)
    conn.execute("DELETE FROM gl_account_roles WHERE role='PL_SUMMARY'")
    conn.commit()
    ROLES.ensure_default_roles(conn)
    conn.commit()


# ── 結轉的效果與決算 ───────────────────────────────────────────────────────

def test_closing_zeroes_pl_keeps_balance_sheet_and_income_statement(conn):
    y = _new_year(conn)
    before = ST.balance_sheet(conn, "%d-12-31" % y)
    _close_periods(conn, y)
    CL.generate(conn, y, "acc")
    conn.commit()
    _post_closing(conn, y)
    after = ST.balance_sheet(conn, "%d-12-31" % y)
    assert after["totals"] == before["totals"] and after["current_pl"] == 0 and after["checks"]["balanced"]
    assert CL._pl_nets(conn, "%d-01-01" % y, "%d-12-31" % y, include_closing=True) == {}          # 含結轉後損益科目全部歸零
    inc = ST.income_statement(conn, "%d-01-01" % y, "%d-12-31" % y)
    assert inc["net_income"]["ytd"] == 5000 and inc["checks"]["balanced"]                            # 損益表不含結轉傳票，淨利照舊
    nxt = ST.balance_sheet(conn, "%d-01-31" % (y + 1))                                              # 隔年：已結轉，無以前年度損益列
    assert nxt["prior_pl"] == 0 and nxt["checks"]["balanced"]


def test_close_year_success_freezes_snapshot_and_closes_period_12(conn):
    y = _full_year(conn)
    res = CL.close_year(conn, y, "acc", accept_warnings=True)
    conn.commit()
    assert res == {"year": y, "net_income": 5000}
    assert conn.execute("SELECT status FROM gl_fiscal_years WHERE year=?", (y,)).fetchone()[0] == "closed"
    assert conn.execute("SELECT status FROM gl_periods WHERE year=? AND period_no=12", (y,)).fetchone()[0] == "closed"
    snap = conn.execute("SELECT frozen, payload_json FROM gl_statement_snapshots WHERE kind='year_close' AND fy=?", (y,)).fetchone()
    assert snap["frozen"] == 1 and '"balance_sheet"' in snap["payload_json"]
    with pytest.raises(sqlite3.IntegrityError):                                                     # 凍結快照連 DB 都改不動
        conn.execute("UPDATE gl_statement_snapshots SET payload_json='{}' WHERE fy=?", (y,))
    conn.rollback()
    assert P.read_log(conn, y)[0]["action"] == "year_close"
    with pytest.raises(CL.ClosingError):
        CL.close_year(conn, y, "acc")                                                               # 已決算不可重複決算


def test_reverse_controls_each_missing_prerequisite_blocks_year_close(conn):
    # ①期間沒結完
    y = _new_year(conn)
    with pytest.raises(CL.ClosingError) as ei:
        CL.close_year(conn, y, "acc")
    assert "還沒結帳" in str(ei.value)
    # ②損益科目還有餘額（沒結轉）
    _close_periods(conn, y)
    with pytest.raises(CL.ClosingError) as ei:
        CL.close_year(conn, y, "acc")
    assert "損益科目還有餘額" in str(ei.value)
    # ③結轉傳票還沒過帳
    CL.generate(conn, y, "acc")
    conn.commit()
    with pytest.raises(CL.ClosingError) as ei:
        CL.close_year(conn, y, "acc")
    assert "損益科目還有餘額" in str(ei.value) or "還沒過帳" in str(ei.value)
    # ④壞帳（借貸不平衡）連期間都結不了
    y2 = _new_year(conn)
    _v(conn, "%d-03-05" % y2, [("1113", 500, 0), ("3111", 0, 400)])
    with pytest.raises(P.PeriodError) as ei:
        _close_periods(conn, y2)
    assert "不平衡" in str(ei.value)
    conn.rollback()
    conn.execute("UPDATE vouchers_all SET voided_at='n', voided_by='t', void_reason='t' WHERE voucher_date=?", ("%d-03-05" % y2,))   # 收拾壞帳，免得殃及後面的題
    conn.commit()
    # ⑤四大表對帳沒過：結轉都做完了，但有餘額的科目沒有報表列 ⇒ 決算被擋
    y3 = _full_year(conn)
    conn.execute("UPDATE gl_account_meta SET fs_line='' WHERE code='1191'")
    conn.commit()
    try:
        with pytest.raises(CL.ClosingError) as ei:
            CL.close_year(conn, y3, "acc", accept_warnings=True)
        assert "對帳沒有全部通過" in str(ei.value) and "資產負債表" in str(ei.value)
    finally:
        conn.execute("UPDATE gl_account_meta SET fs_line='BS_CA_AR' WHERE code='1191'")
        conn.commit()


# ── 年度重開 ─────────────────────────────────────────────────────────────

def test_reopen_year_restores_everything_and_marks_later_periods_stale(conn):
    y = _full_year(conn)
    CL.close_year(conn, y, "acc", accept_warnings=True)
    conn.commit()
    P.create_year(conn, y + 1, "t")
    conn.commit()
    pid = conn.execute("SELECT id FROM gl_periods WHERE year=? AND period_no=1", (y + 1,)).fetchone()[0]
    P.close_period(conn, pid, "acc", accept_warnings=True)
    conn.commit()
    with pytest.raises(CL.ClosingError):
        CL.reopen_year(conn, y, "root", "  ")                                                        # 理由必填
    res = CL.reopen_year(conn, y, "root", "會計師要求調整")
    conn.commit()
    assert len(res["voided"]) == 2 and res["stale_later_periods"] >= 1
    assert conn.execute("SELECT status FROM gl_fiscal_years WHERE year=?", (y,)).fetchone()[0] == "open"
    assert {r[0] for r in conn.execute("SELECT status FROM gl_periods WHERE year=?", (y,))} == {"open"}
    assert CL.closing_vouchers(conn, y) == [] and conn.execute("SELECT stale FROM gl_periods WHERE id=?", (pid,)).fetchone()[0] == 1
    assert P.read_log(conn, y)[0]["action"] == "year_reopen" and P.read_log(conn, y)[0]["reason"] == "會計師要求調整"
    # 重開後損益科目恢復（結轉傳票已作廢），可以重新結轉
    assert CL.preview(conn, y)["net_income"] == 5000
    _close_periods(conn, y)
    CL.generate(conn, y, "acc")
    conn.commit()


def test_reopen_refuses_locked_periods_and_open_years_without_closing(conn):
    y = _full_year(conn)
    CL.close_year(conn, y, "acc", accept_warnings=True)
    conn.commit()
    pid = conn.execute("SELECT id FROM gl_periods WHERE year=? AND period_no=12", (y,)).fetchone()[0]
    P.lock_period(conn, pid, "root")
    conn.commit()
    with pytest.raises(CL.ClosingError) as ei:
        CL.reopen_year(conn, y, "root", "想改")
    assert "已鎖定" in str(ei.value)
    y2 = _new_year(conn, book=False)
    with pytest.raises(CL.ClosingError) as ei:
        CL.reopen_year(conn, y2, "root", "亂試")
    assert "不需要重開" in str(ei.value)


# ── 匯出 ─────────────────────────────────────────────────────────────────

def test_export_is_live_before_close_and_frozen_after(conn):
    y = _full_year(conn)
    stm, source = EX.year_statements(conn, y)
    assert source == "live"
    CL.close_year(conn, y, "acc", accept_warnings=True)
    conn.commit()
    _v(conn, "%d-12-30" % y, [("1113", 999, 0), ("3111", 0, 999)], status="草稿")                   # 之後的帳怎麼動都不影響凍結版
    stm2, source2 = EX.year_statements(conn, y)
    assert source2 == "frozen" and stm2["balance_sheet"]["totals"] == stm["balance_sheet"]["totals"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(EX.statements_workbook(stm2, y, source2)))
    assert wb.sheetnames == ["資產負債表", "綜合損益表", "權益變動表", "現金流量表"]
    assert "決算凍結版" in wb["資產負債表"]["A1"].value


def test_export_neutralises_formula_injection_in_labels(conn):
    y = _full_year(conn)
    stm, _ = EX.year_statements(conn, y)
    stm["balance_sheet"]["sections"]["current_assets"]["items"][0]["label"] = "=HYPERLINK(\"http://x\",\"a\")"
    stm["income_statement"]["lines"][0]["label"] = "=HYPERLINK(\"http://y\",\"b\")"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(EX.statements_workbook(stm, y)))
    for ws in wb.worksheets:
        assert not [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"], "%s 有儲存格被當成公式" % ws.title
    texts = [c.value for row in wb["綜合損益表"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("HYPERLINK" in t and not t.startswith("=") for t in texts)          # 內容還在，只是被中和成純文字


# ── API ─────────────────────────────────────────────────────────────────

def test_api_flow_permissions_and_export(client, make_user, conn):
    fin = _login(client, make_user, "gl_b5_fin", role="staff", modules=("finance",))
    cash = _login(client, make_user, "gl_b5_cash", role="staff", modules=("cashier",))
    sup = _login(client, make_user, "gl_b5_sup")
    none = _login(client, make_user, "gl_b5_none", role="staff", modules=())
    y = _new_year(conn)
    _close_periods(conn, y)
    base = "/api/ledger/years/%d" % y
    assert client.get(base + "/closing/preview", headers=none).status_code == 403
    assert client.get(base + "/closing/preview", headers=cash).status_code == 200
    assert client.post(base + "/closing/generate", headers=cash, json={}).status_code == 403          # cashier 只能看
    r = client.post(base + "/closing/generate", headers=fin, json={})
    assert r.status_code == 200 and len(r.json()["vouchers"]) == 2, r.text
    assert client.post(base + "/closing/generate", headers=fin, json={}).status_code == 400          # 已有草稿：要明說
    assert client.post(base + "/close", headers=fin, json={"accept_warnings": True}).status_code == 400   # 結轉傳票還沒過帳
    _post_closing(conn, y)
    r = client.post(base + "/close", headers=fin, json={"accept_warnings": True})
    assert r.status_code == 200 and r.json()["net_income"] == 5000, r.text
    assert client.post(base + "/reopen", headers=fin, json={"reason": "x"}).status_code == 403     # 重開只有 superadmin
    r = client.get(base + "/statements", headers=cash).json()
    assert r["source"] == "frozen" and r["statements"]["income_statement"]["net_income"]["ytd"] == 5000
    x = client.get(base + "/statements/export", headers=cash)
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats") and x.content[:2] == b"PK"
    assert client.post(base + "/reopen", headers=sup, json={"reason": ""}).status_code == 400
    assert client.post(base + "/reopen", headers=sup, json={"reason": "調整"}).status_code == 200
    assert client.get("/api/ledger/years/1999/closing/preview", headers=fin).status_code == 400   # 年度不存在：說明，不是 500
